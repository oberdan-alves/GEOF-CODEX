# -*- coding: utf-8 -*-
"""
GeradorCotacao.py
==================
MÓDULO PRINCIPAL — Tela de Entrada / Menu Principal
-----------------------------------------------------
Este módulo é a "base" do sistema. Contém tudo que é compartilhado pelos
demais módulos, para evitar duplicação de código e facilitar manutenção:

  - Constantes globais (mapeamento de colunas da planilha, cores, fontes)
  - Funções utilitárias (formatação, normalização de texto, cálculo de
    moeda, leitura da planilha "Controle", conversão de PDF para JPG,
    cálculo de dias úteis considerando feriados etc.)
  - Persistência de configuração (config_padrao.json) e de feriados
    (feriados.json)
  - Classe AppCotacao: monta a tela principal (menu) e herda os métodos de
    ação (botões) definidos nos demais módulos, através de "mixins".

Demais módulos do sistema (cada um cuida de um "Passo" do fluxo):
    Cotacao.py           -> 1º Passo: Nova cotação (CRIAR PASTA / PEDIDO / AJUSTE DO PEDIDO)
    Mapa.py               -> 2º Passo: Preparar Mapa Comparativo (propostas)
    AdjudicacaoOrdem.py  -> 3º Passo: Adjudicação e Ordem de Compra
    Emails.py             -> E-mails (Cotação / Ordem) e cadastro de FERIADOS

Para rodar o sistema, execute este arquivo:
    python GeradorCotacao.py

Dependências: customtkinter, openpyxl, pandas, reportlab, pillow, pymupdf
    pip install customtkinter openpyxl pandas reportlab pillow pymupdf
"""

import os
import re
import json
import glob
import datetime
import unicodedata
import subprocess
import sys
import tkinter as tk
from tkinter import filedialog, ttk
from Mensagens import messagebox

import customtkinter as ctk
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Table, TableStyle, Paragraph, Image as RLImage

# ----------------------------------------------------------------------
# MAPEAMENTO DE COLUNAS DA ABA "Controle" (arquivo Matrix_*.xlsx)
# ----------------------------------------------------------------------
# Cabeçalho na linha 2, dados a partir da linha 3.
COL_TR          = "A"
COL_ORIGEM      = "B"
COL_COTACAO     = "F"   # formato "código/ano", ex.: ".009/2025" ou "501/2026"
COL_SEI         = "D"
COL_ITEM        = "G"
COL_CODIGO      = "H"
COL_DESCRICAO   = "L"
COL_UN          = "M"
COL_QTD         = "N"

LINHA_INICIAL_DADOS = 3
QTD_CHAR_REMOVER_NOME_PLANILHA = 11

OPERADORES_PADRAO = ["BRUNO", "CLAYTON", "GUILHERME", "JEFFERSON", "OBERDAN"]
OPCAO_NOVO_OPERADOR = "NOVO OPERADOR"

QTD_MAX_COTACOES = 10

# Aceita ".009/2025" ou "501/2026"
REGEX_COTACAO = re.compile(r"\.?\s*(\d+)\s*/\s*(\d+)")

CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config_padrao.json")
FERIADOS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "feriados.json")

# Pasta de imagens fixas do sistema (brasão etc.). Fica ao lado deste .py.
CAMINHO_ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
CAMINHO_LOGO_GDF = os.path.join(CAMINHO_ASSETS, "brasao_gdf.png")

# ----------------------------------------------------------------------
# CORES GERAIS DA INTERFACE
# ----------------------------------------------------------------------
COR_FUNDO_APP      = "#1a1a1a"
COR_CARD           = "#242424"
COR_TEXTO          = "#f2f2f2"
COR_TEXTO_SEC      = "#c9c9c9"
COR_BORDA          = "#3a3a3a"
COR_DESTAQUE       = "#3b82f6"

# Ordem de exibição no dropdown "NOME PLANILHA"
SUFIXOS_PLANILHA_EXIBICAO = ["APS", "HRG", "APS INV", "HRG INV"]
# Ordem para detecção do sufixo a partir do texto completo (mais específico
# primeiro, para não confundir "APS" com "APS INV")
SUFIXOS_PLANILHA = ["APS INV", "HRG INV", "APS", "HRG"]

CORES_NOME_PLANILHA = {
    "APS":     {"fg": "#a9e3a9", "text": "#000000", "bold": False},
    "HRG":     {"fg": "#a9cdf0", "text": "#000000", "bold": False},
    "APS INV": {"fg": "#1f5c1f", "text": "#ffe14d", "bold": True},
    "HRG INV": {"fg": "#1a3a66", "text": "#ffe14d", "bold": True},
}

COR_COTACAO_ATIVA = "#ffd23f"
COR_COTACAO_INATIVA = "#565b5f"

# Cores dos grupos de botões (parte inferior da tela principal)
COR_GRUPO_1 = "#bfe6bf"   # 1º Passo - verde claro
COR_GRUPO_2 = "#f3edaa"   # 2º Passo - amarelo claro
COR_GRUPO_3 = "#f0b3b3"   # 3º Passo - vermelho claro
COR_GRUPO_4 = "#d6bdf0"   # E-mails - roxo claro

# Cores usadas nos documentos PDF (Pedido de Cotação, Mapa Comparativo,
# Ordem de Compra) — compartilhadas entre Cotacao.py e AdjudicacaoOrdem.py
COR_CINZA_ADJ = colors.HexColor("#A5A5A5")
COR_CIANO_ADJ = colors.HexColor("#03E7ED")     # HRG (azul claro)
COR_VERDE_ADJ = colors.HexColor("#00FF00")     # APS (verde claro)
COR_AMARELO_ADJ = colors.HexColor("#FFFF00")   # aparece junto quando é "...INV"

# ----------------------------------------------------------------------
# CONTATOS INSTITUCIONAIS USADOS NA ORDEM DE COMPRA (editáveis, por tipo)
# ----------------------------------------------------------------------
CONTATOS_ORDEM_PADRAO = {
    "GERENTE": "CREMILDA,  Fone: (61) 9 9523-1257",
    "GERENTE_HRG": "WANDSON, Fone: (61) 9173-3605",
    "NFL": "CLENILCE DIAS SILVA - Matr.1711246-X (61)9 9202-2307",
    "NUAL_APS": "CARLOS EDUARDO LIRA DOS SANTOS, matrícula 1.709.549-2",
    "NECFM": "RODRIGO RIBEIRO TENÓRIO, matrícula 1.691.497-X",
    "NPDOC_APSINV": "ANNA GARDENIA BASTOS MENDES - Matr.1442838-5 - 61 9283-6088",
    "NFH": "MIRELLE CARDARELLI (61) 3449-7070/7071/7073",
    "NECFM_HRG": "RODRIGO TENÓRIO (Mezanino Acima do PAI no HRG) (61) 9 9639-0162 ***MANUTENÇÃO EQUIPAMENTOS***",
    "NUAL_HRG": "WEILLER MACHADO NOBRE ALVES (61) 3449-7092/7093/7094",
    "NPDOC_HRGINV": "Anderson Carrascozza Rodrigues Pereira - Matr.1.442.917-9 - 61 994 343 504",
}

# ----------------------------------------------------------------------
# FONTES
# ----------------------------------------------------------------------
FONTE_LABEL   = ("Segoe UI", 13)
FONTE_ENTRY   = ("Segoe UI", 14)
FONTE_TITULO  = ("Segoe UI", 19, "bold")
FONTE_SECAO   = ("Segoe UI", 13, "bold")
FONTE_BOTAO   = ("Segoe UI", 12, "bold")
FONTE_TABELA  = ("Segoe UI", 11)


# ----------------------------------------------------------------------
# PERSISTÊNCIA DE VARIÁVEIS CONSTANTES (defaults salvos em disco)
# ----------------------------------------------------------------------
def carregar_config():
    padrao = {
        "pasta_geof": "",
        "ano": "",
        "bimestre": "1",
        "da_nome": "",
        "da_matricula": "",
        "super_nome": "",
        "super_matricula": "",
        "operadores": OPERADORES_PADRAO.copy(),
        "ultimo_operador": OPERADORES_PADRAO[0],
        "ultima_planilha_tipo": "APS",
        "contatos_ordem": CONTATOS_ORDEM_PADRAO.copy(),
    }
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                dados = json.load(f)
            padrao.update(dados)
            # garante que chaves novas (de versões futuras) sempre existam
            contatos_merge = CONTATOS_ORDEM_PADRAO.copy()
            contatos_merge.update(padrao.get("contatos_ordem") or {})
            padrao["contatos_ordem"] = contatos_merge
        except Exception:
            pass
    return padrao


def salvar_config(config: dict):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(config, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Erro ao salvar config: {e}")


# ----------------------------------------------------------------------
# PERSISTÊNCIA E CÁLCULO DE FERIADOS
# ----------------------------------------------------------------------
# Usados pelo botão FERIADOS (cadastro) e pelo E-mail Cotação (cálculo do
# prazo de recebimento das propostas em dias úteis, pulando sábados,
# domingos e os feriados cadastrados). Guardados num arquivo próprio
# (feriados.json), separado do config_padrao.json, por ser um cadastro
# de natureza diferente (datas do calendário, não preferências do usuário).
def carregar_feriados():
    """Lê feriados.json e devolve a lista de feriados cadastrados, cada um
    como {"data": "DD/MM/AAAA", "feriado": "texto curto"}. Devolve lista
    vazia se o arquivo ainda não existir ou não puder ser lido."""
    if os.path.exists(FERIADOS_FILE):
        try:
            with open(FERIADOS_FILE, "r", encoding="utf-8") as f:
                dados = json.load(f)
            return dados.get("feriados", [])
        except Exception:
            return []
    return []


def salvar_feriados(lista_feriados):
    try:
        with open(FERIADOS_FILE, "w", encoding="utf-8") as f:
            json.dump({"feriados": lista_feriados}, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Erro ao salvar feriados: {e}")


def obter_datas_feriados(lista_feriados):
    """Converte a lista de feriados salvos (strings 'DD/MM/AAAA') num set de
    objetos datetime.date, para consulta rápida em somar_dias_uteis().
    Entradas malformadas (data incompleta, texto não numérico) são
    ignoradas silenciosamente."""
    datas = set()
    for item in lista_feriados or []:
        texto = (item.get("data") or "").strip()
        partes = texto.split("/")
        if len(partes) != 3:
            continue
        try:
            d, m, a = partes
            datas.add(datetime.date(int(a), int(m), int(d)))
        except ValueError:
            continue
    return datas


def somar_dias_uteis(data_inicio, quantidade_dias_uteis, datas_feriados=None):
    """Soma 'quantidade_dias_uteis' dias úteis a partir do dia seguinte a
    'data_inicio' (a própria data_inicio nunca é contada), pulando sábados,
    domingos e as datas presentes em 'datas_feriados'.

    Exemplo: hoje é sexta-feira, segunda-feira é feriado e quer-se 3 dias
    úteis -> a contagem pula sábado/domingo, pula a segunda (feriado) e
    conta terça=1, quarta=2, quinta=3; o resultado é a quinta-feira."""
    datas_feriados = datas_feriados or set()
    data_atual = data_inicio
    contados = 0
    while contados < quantidade_dias_uteis:
        data_atual += datetime.timedelta(days=1)
        if data_atual.weekday() >= 5:  # 5 = sábado, 6 = domingo
            continue
        if data_atual in datas_feriados:
            continue
        contados += 1
    return data_atual


# ----------------------------------------------------------------------
# MÁSCARAS / FORMATAÇÃO (utilizadas por vários módulos)
# ----------------------------------------------------------------------
def formatar_matricula(texto: str) -> str:
    """Aplica a máscara 9.999.999-9 a partir dos dígitos digitados."""
    digitos = re.sub(r"\D", "", texto)[:8]
    if len(digitos) <= 1:
        return digitos
    parte_inicial = digitos[0]
    resto = digitos[1:]
    grupos = []
    for i in range(0, len(resto), 3):
        grupos.append(resto[i:i + 3])
    corpo = ".".join(grupos)
    if len(digitos) == 8:
        corpo_sem_verificador = ".".join([resto[0:3], resto[3:6]])
        verificador = resto[6:7]
        return f"{parte_inicial}.{corpo_sem_verificador}-{verificador}"
    return f"{parte_inicial}.{corpo}"


def formatar_apenas_digitos(texto: str, tamanho: int) -> str:
    return re.sub(r"\D", "", texto)[:tamanho]


def formatar_data(texto):
    """Aplica a máscara DD/MM/AAAA a partir dos dígitos digitados."""
    digitos = re.sub(r"\D", "", texto)[:8]
    partes = []
    if len(digitos) > 0:
        partes.append(digitos[0:2])
    if len(digitos) > 2:
        partes.append(digitos[2:4])
    if len(digitos) > 4:
        partes.append(digitos[4:8])
    return "/".join(partes)


def formatar_hora(texto):
    """Aplica a máscara HH:MM a partir dos dígitos digitados."""
    digitos = re.sub(r"\D", "", texto)[:4]
    partes = []
    if len(digitos) > 0:
        partes.append(digitos[0:2])
    if len(digitos) > 2:
        partes.append(digitos[2:4])
    return ":".join(partes)


def parse_numero(texto):
    """Converte texto digitado (aceita vírgula ou ponto decimal) em float."""
    if texto is None:
        return None
    texto = str(texto).strip().replace(" ", "")
    if not texto:
        return None
    texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def formatar_moeda(valor):
    try:
        return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "R$ 0,00"


def formatar_numero_inteiro(valor, separador_milhar=False):
    """Formata um número que deveria ser sempre inteiro (campos ITEM e QTD
    lidos da aba Controle, que o Excel/openpyxl às vezes devolve como
    float — 1.0, 2.0 — mesmo quando o valor digitado na planilha é um
    número inteiro) como texto sem casas decimais. Usada tanto na tabela
    da tela (ITENS DA COTAÇÃO) quanto no PDF do Pedido, para manter os
    dois lugares consistentes.

    Com separador_milhar=True, usa o ponto como separador de milhar
    (padrão brasileiro, ex.: 1500 -> "1.500") — usado só para QTD; ITEM
    nunca usa separador de milhar (são sempre números pequenos e
    sequenciais). Valores que não são numéricos voltam como texto puro,
    sem alteração."""
    if valor is None or valor == "":
        return ""
    try:
        inteiro = int(round(float(valor)))
    except (TypeError, ValueError):
        return str(valor)
    if separador_milhar:
        return f"{inteiro:,}".replace(",", ".")
    return str(inteiro)


def truncar_em_palavra(texto, limite):
    """Corta o texto no máximo em 'limite' caracteres, sem partir uma
    palavra no meio (corta no último espaço antes do limite)."""
    if len(texto) <= limite:
        return texto
    corte = texto[:limite]
    ultimo_espaco = corte.rfind(" ")
    if ultimo_espaco > 0:
        return corte[:ultimo_espaco].rstrip()
    return corte.rstrip()


# Remove QUALQUER sequência de caracteres que não sejam letras no início da
# descrição do item (números, pontos — inclusive vários seguidos, como em
# "1.234.344" —, traços, parênteses, espaços etc.), até chegar na primeira
# letra. Isso evita que gerar_objeto_automatico() confunda um ponto de
# numeração/código no começo da descrição com o ponto final da frase: sem
# essa limpeza, "1. Manutenção..." ou "1.234.344 - Manutenção..." cortavam
# no primeiro "." e o objeto automático saía só o número. A descrição
# considerada pela função sempre passa a começar com letra.
REGEX_PREFIXO_NAO_LETRA_ITEM = re.compile(r"^[^A-Za-zÀ-ÖØ-öø-ÿ]+")


def gerar_objeto_automatico(mapa_itens, limite_por_item=40):
    """Monta o título/objeto (usado no Mapa Comparativo e no E-mail
    Cotação) a partir do início da descrição de cada item (até a primeira
    vírgula ou ponto, ignorando um eventual prefixo de numeração de lista
    como "1. "), em maiúsculas, cortado em no máximo 'limite_por_item'
    caracteres (sem partir palavra), sem repetir nomes iguais, separados
    por vírgula e terminando em '... e ÚLTIMO ITEM.' — como nos modelos
    reais. Função compartilhada por todos os módulos (movida para cá para
    não duplicar a lógica em cada um)."""
    nomes = []
    vistos = set()
    for item in mapa_itens:
        desc = (item.get("descricao") or "").strip()
        desc = REGEX_PREFIXO_NAO_LETRA_ITEM.sub("", desc, count=1)
        corte = len(desc)
        for sep in (",", "."):
            pos = desc.find(sep)
            if pos != -1:
                corte = min(corte, pos)
        nome_curto = desc[:corte].strip().upper()
        nome_curto = truncar_em_palavra(nome_curto, limite_por_item)
        if not nome_curto:
            continue
        chave = normalizar_texto(nome_curto)
        if chave in vistos:
            continue
        vistos.add(chave)
        nomes.append(nome_curto)
    if not nomes:
        return ""
    if len(nomes) == 1:
        return nomes[0]
    return ", ".join(nomes[:-1]) + " e " + nomes[-1] + "."


def normalizar_texto(texto):
    """Normaliza texto (maiúsculas, sem acento, espaços colapsados) para
    comparações — usado ao casar nomes de empresas e códigos."""
    if texto is None:
        return ""
    texto = str(texto).strip().upper()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto)
    return texto


def corrigir_acentuacao(valor):
    """Corrige o texto lido do Excel para exibição em PDF (mantendo os
    acentos, ao contrário de normalizar_texto acima).

    Alguns arquivos .xlsx guardam letras acentuadas na forma "decomposta"
    do Unicode (ex.: a letra 'ã' gravada como 'a' + um sinal de til
    separado, em vez do único caractere composto 'ã'). Isso é invisível no
    Excel, mas o ReportLab (fontes Helvetica usadas nos PDFs) não tem um
    desenho para o sinal de acento sozinho e imprime um quadradinho no
    lugar da letra — é a causa das letras/acentos corrompidos no PEDIDO,
    na ADJUDICAÇÃO e na ORDEM. A normalização NFC junta de volta a letra
    com o acento em um único caractere, que a fonte sabe desenhar.
    Números e outros tipos de valor voltam sem alteração."""
    if valor is None:
        return valor
    if isinstance(valor, str):
        return unicodedata.normalize("NFC", valor)
    return valor


# Reconhece "NOME DA EMPRESA (CODIGO)" -> ("NOME DA EMPRESA", "CODIGO")
REGEX_FORNECEDOR = re.compile(r"^(.*?)\s*\(([^)]*)\)\s*$")


def extrair_nome_codigo_fornecedor(texto_coluna_a):
    texto = str(texto_coluna_a).strip()
    m = REGEX_FORNECEDOR.match(texto)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return texto, ""


def extrair_codigo_ano_cotacao(valor):
    """Recebe o valor bruto da coluna Cotação e devolve (codigo_3_digitos, ano) ou (None, None)."""
    if valor is None:
        return None, None
    texto = str(valor).strip()
    m = REGEX_COTACAO.match(texto)
    if not m:
        return None, None
    codigo = m.group(1).zfill(3)[-3:]
    ano = m.group(2)
    return codigo, ano


# ----------------------------------------------------------------------
# LEITURA DA PLANILHA "Controle" (busca principal de cotação, usada na tela)
# ----------------------------------------------------------------------
def buscar_cotacoes_na_planilha(caminho_arquivo: str, codigos_cotacao: list, ano: str = ""):
    """
    Lê a aba 'Controle' do arquivo Excel informado e retorna a lista de
    linhas (dicts) cujo código de cotação (3 dígitos, coluna F) bate com
    algum dos códigos informados. Se 'ano' for informado, também filtra
    pelo ano contido no valor da cotação (ex.: ".009/2025" -> ano 2025),
    evitando colisão entre códigos repetidos em anos diferentes.
    """
    if not os.path.exists(caminho_arquivo):
        raise FileNotFoundError(f"Arquivo não encontrado:\n{caminho_arquivo}")

    wb = load_workbook(caminho_arquivo, data_only=True, read_only=True)
    if "Controle" not in wb.sheetnames:
        raise ValueError("A aba 'Controle' não foi encontrada na planilha.")
    ws = wb["Controle"]

    idx = {
        "tr": column_index_from_string(COL_TR),
        "origem": column_index_from_string(COL_ORIGEM),
        "cotacao": column_index_from_string(COL_COTACAO),
        "sei": column_index_from_string(COL_SEI),
        "item": column_index_from_string(COL_ITEM),
        "codigo": column_index_from_string(COL_CODIGO),
        "descricao": column_index_from_string(COL_DESCRICAO),
        "un": column_index_from_string(COL_UN),
        "qtd": column_index_from_string(COL_QTD),
    }

    codigos_norm = {c.strip().zfill(3) for c in codigos_cotacao if c.strip()}
    ano_norm = ano.strip()
    linhas_encontradas = []

    for linha in ws.iter_rows(min_row=LINHA_INICIAL_DADOS, values_only=False):
        val_cotacao = linha[idx["cotacao"] - 1].value
        codigo, ano_linha = extrair_codigo_ano_cotacao(val_cotacao)
        if codigo is None or codigo not in codigos_norm:
            continue
        if ano_norm and ano_linha and ano_linha != ano_norm:
            continue
        linhas_encontradas.append({
            "TR": corrigir_acentuacao(linha[idx["tr"] - 1].value),
            "ORIGEM": corrigir_acentuacao(linha[idx["origem"] - 1].value),
            "COTACAO": codigo,
            "COTACAO_ANO": ano_linha,
            "COTACAO_BRUTA": val_cotacao,
            "SEI": corrigir_acentuacao(linha[idx["sei"] - 1].value),
            "ITEM": linha[idx["item"] - 1].value,
            "CODIGO": corrigir_acentuacao(linha[idx["codigo"] - 1].value),
            "DESCRICAO": corrigir_acentuacao(linha[idx["descricao"] - 1].value),
            "UN": corrigir_acentuacao(linha[idx["un"] - 1].value),
            "QTD": linha[idx["qtd"] - 1].value,
        })

    wb.close()
    return linhas_encontradas


# ----------------------------------------------------------------------
# UTILITÁRIOS DE PDF/IMAGEM (compartilhados por Cotacao.py e AdjudicacaoOrdem.py)
# ----------------------------------------------------------------------
def converter_pdf_para_jpg(caminho_pdf, caminho_jpg_sem_extensao):
    """Converte TODAS as páginas do PDF em UMA ÚNICA imagem JPG (páginas
    empilhadas verticalmente, com uma fina linha cinza separando cada uma),
    já pronta para anexar ao SEI. O arquivo final leva o MESMO NOME do PDF
    (apenas com extensão .jpg).

    Usa PyMuPDF (pacote "pymupdf", importado como "fitz") para renderizar as
    páginas — não depende de nenhum programa externo instalado à parte
    (como o Poppler/pdftoppm), o que evita o erro "pdftoppm não gerou
    nenhuma imagem de página" quando esse programa não está instalado ou
    não está no PATH do Windows.
    """
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz  # nome antigo do mesmo pacote, mantido por compatibilidade
        except ImportError:
            raise RuntimeError(
                "A biblioteca 'pymupdf' não está instalada neste Python.\n\n"
                "Abra o terminal (o mesmo que você usa para rodar o programa) e execute:\n"
                "    pip install pymupdf\n\n"
                "Se o computador tiver mais de uma versão de Python instalada, use:\n"
                "    py -m pip install pymupdf\n"
                "(ou o caminho completo do python.exe usado para abrir o programa)."
            ) from None
    from PIL import Image, ImageDraw

    # Limpa eventuais arquivos "_pagN.jpg" deixados por uma tentativa
    # anterior que tenha falhado no meio do caminho, para não acumular lixo
    # na pasta da cotação.
    for antigo in glob.glob(caminho_jpg_sem_extensao + "_pag*.jpg"):
        try:
            os.remove(antigo)
        except OSError:
            pass

    documento = fitz.open(caminho_pdf)
    try:
        if documento.page_count == 0:
            raise RuntimeError("O PDF gerado não contém nenhuma página.")

        zoom = 150 / 72  # ~150 DPI
        matriz = fitz.Matrix(zoom, zoom)
        imagens = []
        for pagina in documento:
            pix = pagina.get_pixmap(matrix=matriz, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            imagens.append(img)
    finally:
        documento.close()

    largura = max(img.width for img in imagens)
    espaco = 14  # px entre páginas
    altura_total = sum(img.height for img in imagens) + espaco * (len(imagens) - 1)

    combinada = Image.new("RGB", (largura, altura_total), "white")
    desenho = ImageDraw.Draw(combinada)
    y = 0
    for i, img in enumerate(imagens):
        x = (largura - img.width) // 2
        combinada.paste(img, (x, y))
        y += img.height
        if i < len(imagens) - 1:
            desenho.rectangle([0, y + espaco // 2 - 1, largura, y + espaco // 2 + 1],
                              fill=(170, 170, 170))
            y += espaco

    caminho_final = caminho_jpg_sem_extensao + ".jpg"
    combinada.save(caminho_final, quality=92)

    for img in imagens:
        img.close()

    return caminho_final


def montar_cabecalho_institucional(largura_util, linhas_texto, tamanho_logo=1.7 * cm):
    """Monta o cabeçalho padrão dos documentos PDF (Pedido de Cotação, Mapa
    Comparativo, Ordem de Compra): o brasão do GDF à esquerda e o texto
    institucional ALINHADO À ESQUERDA logo ao lado dele — replicando o
    layout oficial usado nos modelos reais da SES/PDPAS (o texto NÃO fica
    centralizado na página; fica encostado no brasão, como no papel
    timbrado oficial).

    'linhas_texto' é uma lista de strings (uma por linha, já prontas — ex.:
    o sufixo "/HRG" já concatenado na última linha). Se o arquivo do brasão
    (assets/brasao_gdf.png) não existir — ou não puder ser aberto — o
    cabeçalho é montado só com o texto, sem quebrar a geração do PDF; um
    aviso é impresso no console para facilitar o diagnóstico.
    """
    estilo_texto = ParagraphStyle(
        "CabecalhoInstitucional", fontName="Helvetica-Bold", fontSize=9,
        leading=11, alignment=TA_LEFT)
    paragrafo = Paragraph("<br/>".join(linhas_texto), estilo_texto)

    logo_flowable = ""
    if os.path.exists(CAMINHO_LOGO_GDF):
        try:
            logo_flowable = RLImage(CAMINHO_LOGO_GDF, width=tamanho_logo, height=tamanho_logo)
        except Exception as e:
            print(f"[Aviso] Não foi possível carregar a logo em {CAMINHO_LOGO_GDF}: {e}")
    else:
        print(f"[Aviso] Logo do GDF não encontrada em {CAMINHO_LOGO_GDF} — "
              "confira se a pasta 'assets' está ao lado de GeradorCotacao.py.")

    largura_logo_col = tamanho_logo + 0.3 * cm
    largura_texto_col = largura_util - largura_logo_col
    tabela = Table([[logo_flowable, paragrafo]],
                    colWidths=[largura_logo_col, largura_texto_col])
    tabela.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, 0), "CENTER"),
        ("LEFTPADDING", (0, 0), (0, 0), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 6),
        ("LEFTPADDING", (1, 0), (1, 0), 0),
    ]))
    return tabela


def maximizar_janela(janela):
    """Maximiza após a inicialização do gerenciador de janelas do CTk."""
    def maximizar():
        try:
            janela.state("zoomed")
        except tk.TclError:
            try:
                janela.attributes("-zoomed", True)
            except tk.TclError:
                escala = janela._get_window_scaling() if hasattr(janela, "_get_window_scaling") else 1
                largura = int(janela.winfo_screenwidth() / escala)
                altura = int(janela.winfo_screenheight() / escala) - 60
                janela.geometry(f"{largura}x{altura}+0+0")
    janela.after(250, maximizar)


def centralizar_janela(janela):
    """Centraliza após a montagem da tela, limitando o tamanho ao monitor."""
    def posicionar():
        janela.update_idletasks()
        largura, altura = map(int, re.match(r"^(\d+)x(\d+)", janela.geometry()).groups())
        tela_largura, tela_altura = janela.winfo_screenwidth(), janela.winfo_screenheight()
        escala = janela._get_window_scaling() if hasattr(janela, "_get_window_scaling") else 1
        largura = min(largura, max(1, int(tela_largura / escala) - 40))
        altura = min(altura, max(1, int(tela_altura / escala) - 80))
        minimo_largura, minimo_altura = janela.minsize()
        janela.minsize(min(minimo_largura, largura), min(minimo_altura, altura))
        x = max(0, int((tela_largura - largura * escala) / 2))
        y = max(0, int((tela_altura - altura * escala) / 2))
        janela.geometry(f"{largura}x{altura}+{x}+{y}")
    janela.after_idle(posicionar)


def exibir_preview_pdf_jpg(parent, caminho_jpg, titulo="Pré-visualização"):
    """Abre uma janela CTkToplevel exibindo a imagem JPG gerada a partir de
    um PDF (usada após gerar Pedido de Cotação / Mapa Comparativo / Ordem).
    Centraliza aqui a lógica repetida que antes existia em 3 lugares."""
    try:
        janela = ctk.CTkToplevel(parent)
        janela.title(titulo)
        janela.configure(fg_color=COR_FUNDO_APP)
        from PIL import Image
        img = Image.open(caminho_jpg)
        largura_max, altura_max = 1100, 780
        escala = min(largura_max / img.width, altura_max / img.height, 1)
        tam = (int(img.width * escala), int(img.height * escala))
        ctk_img = ctk.CTkImage(light_image=img, dark_image=img, size=tam)
        janela.geometry(f"{tam[0] + 40}x{tam[1] + 80}")
        centralizar_janela(janela)
        ctk.CTkLabel(janela, image=ctk_img, text="").pack(padx=16, pady=16)
        ctk.CTkButton(janela, text="Fechar", font=FONTE_BOTAO, height=34,
                      command=janela.destroy).pack(pady=(0, 12))
    except Exception:
        pass  # a pré-visualização é apenas um bônus; os arquivos já foram salvos


# ----------------------------------------------------------------------
# APLICAÇÃO (classe BASE — sem os mixins ainda)
# ----------------------------------------------------------------------
# IMPORTANTE sobre a ordem dos imports: Cotacao.py, Mapa.py,
# AdjudicacaoOrdem.py e Emails.py fazem "from GeradorCotacao import ...*"
# para reaproveitar as constantes/funções acima. Se este arquivo importasse
# esses módulos aqui, no topo (nível de módulo), e algum deles fosse
# executado/importado ANTES de GeradorCotacao.py (por exemplo, ao rodar
# "python Cotacao.py" diretamente, ou por causa da ordem de import de uma
# IDE), teríamos um import circular: Cotacao.py -> GeradorCotacao.py ->
# Cotacao.py (ainda incompleto) -> ImportError.
#
# Por isso a classe abaixo (_AppCotacaoBase) NÃO herda os mixins diretamente.
# A composição final (AppCotacao = _AppCotacaoBase + mixins) só acontece
# dentro de criar_app(), chamada no final do arquivo — nesse ponto
# GeradorCotacao.py já terminou de carregar 100%, então não há mais risco
# de ciclo, não importa qual arquivo tenha sido executado primeiro.
#
# Alguns métodos abaixo (_abrir_pdf_pedido_cotacao_ativa e
# _carregar_empresas_mapa_se_existir, chamados por acao_consultar_cotacao)
# usam self._caminho_arquivo_mapa_empresas(...), que só existe em tempo de
# execução — vem do MapaMixin (Mapa.py), misturado pela criar_app(). Isso é
# seguro porque esses métodos só são de fato chamados depois que a classe
# final AppCotacao (já com todos os mixins) foi instanciada; o mesmo padrão
# já é usado em _montar_grupos_botoes, que referencia botões de outros
# mixins (self.acao_criar_pasta_cotacao etc.) do mesmo jeito.
class _AppCotacaoBase(ctk.CTk):
    """Tela principal (base). Os métodos de ação (botões dos 4 grupos)
    vêm dos mixins definidos em Cotacao.py, Mapa.py, AdjudicacaoOrdem.py e
    Emails.py, misturados pela função criar_app() no final deste arquivo."""

    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("dark-blue")

        self.config_dados = carregar_config()

        self.title("Controle de Cotações - Tela de Entrada")
        self.geometry("1600x900")
        centralizar_janela(self)
        self.configure(fg_color=COR_FUNDO_APP)

        self.entradas_cotacao = []          # entries da matriz de cotação
        self.linhas_resultado = []          # cache do último resultado
        self.lista_empresas_mapa = []       # empresas resolvidas (para INCLUIR EMPRESAS MAPA)
        self.mapa_fornecedores = {}         # dados digitados por fornecedor (cadastro de propostas)
        self.indice_fornecedor_atual = 0
        self.janela_wizard = None
        self._blink_jobs = {}               # after() ids do destaque piscante (validação de saldo)
        self.saldos_bd_atual = {}
        self._pasta_propostas_atual = None
        self._pasta_cotacao_atual = None
        self._codigo_cotacao_atual = None
        self._ano_atual_wizard = None
        self._dados_form_atual = None
        self.caminho_arquivo_atual = None   # caminho do .xlsx localizado na pasta MATRIX
        self._chave_arquivo_atual = None    # (pasta_geof, ano, tipo, nome_planilha) usados para localizá-lo

        self._montar_layout()
        self._preencher_defaults()

    # ------------------------------------------------------------------
    # LAYOUT GERAL
    # ------------------------------------------------------------------
    def _montar_layout(self):
        titulo = ctk.CTkLabel(self, text="CONTROLE DE COTAÇÕES",
                               font=FONTE_TITULO, text_color=COR_TEXTO)
        titulo.pack(pady=(8, 6))

        corpo = ctk.CTkFrame(self, fg_color=COR_FUNDO_APP)
        corpo.pack(fill="both", expand=True, padx=20)

        corpo.grid_columnconfigure(0, weight=1)
        corpo.grid_columnconfigure(1, weight=1)

        self._montar_coluna_esquerda(corpo)
        self._montar_coluna_direita(corpo)

        # Tabela de resultados (largura total)
        self._montar_tabela(self)

        # Grupos de botões (largura total) — cada grupo chama métodos dos mixins
        self._montar_grupos_botoes(self)

    # ------------------------------------------------------------------
    # COLUNA ESQUERDA: COTAÇÃO / NOME PLANILHA / ANO / BIMESTRE / OPERADOR
    # ------------------------------------------------------------------
    def _montar_coluna_esquerda(self, parent):
        col = ctk.CTkFrame(parent, fg_color=COR_CARD, corner_radius=12)
        col.grid(row=0, column=0, sticky="nsew", padx=(0, 8), pady=6)

        pad = {"padx": 16, "pady": (6, 2)}

        # --- COTAÇÃO (matriz até 10) ---
        ctk.CTkLabel(col, text="COTAÇÃO (até 10 - código 999)",
                     font=FONTE_SECAO, text_color=COR_TEXTO).pack(anchor="w", **pad)
        ctk.CTkLabel(col, text="Clique/digite em uma caixa para torná-la a cotação ativa",
                     font=("Segoe UI", 11), text_color=COR_TEXTO_SEC).pack(anchor="w", padx=16)

        grid_cot = ctk.CTkFrame(col, fg_color="transparent")
        grid_cot.pack(anchor="w", padx=16, pady=2)
        self.indice_cotacao_ativa = None
        for i in range(QTD_MAX_COTACOES):
            e = ctk.CTkEntry(grid_cot, width=70, height=32, font=FONTE_ENTRY,
                              placeholder_text="999", justify="center",
                              border_width=2, border_color=COR_COTACAO_INATIVA)
            e.grid(row=i // 5, column=i % 5, padx=3, pady=3)
            e.bind("<KeyRelease>", lambda ev, idx=i: self._on_cotacao_digitada(idx))
            e.bind("<FocusIn>", lambda ev, idx=i: self._definir_cotacao_ativa(idx))
            e.bind("<Button-1>", lambda ev, idx=i: self._definir_cotacao_ativa(idx))
            self.entradas_cotacao.append(e)

        ctk.CTkButton(grid_cot, text="⚙ Contatos\nda Ordem", font=("Segoe UI", 11, "bold"),
                      width=76, height=67, fg_color="#3a3a3a", hover_color="#4a4a4a",
                      command=self._abrir_configuracao_contatos
                      ).grid(row=0, column=5, rowspan=2, padx=(10, 0), pady=3, sticky="ns")

        ctk.CTkButton(grid_cot, text="🔄 Reset /\nAtualizar", font=("Segoe UI", 11, "bold"),
                      width=76, height=67, fg_color="#3a3a3a", hover_color="#4a4a4a",
                      command=self.acao_resetar_cotacao_ativa
                      ).grid(row=0, column=6, rowspan=2, padx=(6, 0), pady=3, sticky="ns")

        ctk.CTkButton(col, text="CONSULTAR COTAÇÃO", font=FONTE_BOTAO,
                      fg_color=COR_DESTAQUE, height=32,
                      command=self.acao_consultar_cotacao).pack(fill="x", padx=16, pady=(4, 8))

        # --- NOME PLANILHA (dropdown já com as opções completas e coloridas) ---
        ctk.CTkLabel(col, text="NOME PLANILHA", font=FONTE_SECAO,
                     text_color=COR_TEXTO).pack(anchor="w", **pad)

        self.var_nome_planilha = tk.StringVar(value="")
        self.combo_planilha = ctk.CTkOptionMenu(
            col, values=["Matrix_ANO_APS"], variable=self.var_nome_planilha,
            font=FONTE_ENTRY, dropdown_font=FONTE_ENTRY, height=32,
            command=self._on_selecionar_nome_planilha)
        self.combo_planilha.pack(fill="x", padx=16, pady=2)

        ctk.CTkButton(col, text="CONSULTAR MATRIX", font=FONTE_BOTAO,
                      fg_color=COR_DESTAQUE, height=32,
                      command=self.acao_consultar_matrix).pack(fill="x", padx=16, pady=(4, 8))

        # --- ANO / BIMESTRE-SEMESTRE / OPERADOR (mesma linha) ---
        linha_ano = ctk.CTkFrame(col, fg_color="transparent")
        linha_ano.pack(fill="x", padx=16, pady=(6, 8))
        linha_ano.grid_columnconfigure((0, 1, 2), weight=1)

        sub_ano = ctk.CTkFrame(linha_ano, fg_color="transparent")
        sub_ano.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ctk.CTkLabel(sub_ano, text="ANO", font=FONTE_SECAO,
                     text_color=COR_TEXTO).pack(anchor="w")
        self.entry_ano = ctk.CTkEntry(sub_ano, font=FONTE_ENTRY, height=32, placeholder_text="AAAA")
        self.entry_ano.pack(fill="x", pady=(2, 0))
        self.entry_ano.bind("<KeyRelease>", self._on_ano_digitado)

        sub_bim = ctk.CTkFrame(linha_ano, fg_color="transparent")
        sub_bim.grid(row=0, column=1, sticky="ew", padx=6)
        ctk.CTkLabel(sub_bim, text="BIMESTRE/SEMESTRE", font=FONTE_SECAO,
                     text_color=COR_TEXTO).pack(anchor="w")
        self.var_bimestre = tk.StringVar(value="1")
        ctk.CTkOptionMenu(sub_bim, values=["1", "2", "3", "4", "5", "6"],
                          variable=self.var_bimestre, font=FONTE_ENTRY,
                          dropdown_font=FONTE_ENTRY, height=32).pack(fill="x", pady=(2, 0))

        sub_op = ctk.CTkFrame(linha_ano, fg_color="transparent")
        sub_op.grid(row=0, column=2, sticky="ew", padx=(6, 0))
        ctk.CTkLabel(sub_op, text="OPERADOR", font=FONTE_SECAO,
                     text_color=COR_TEXTO).pack(anchor="w")
        linha_operador = ctk.CTkFrame(sub_op, fg_color="transparent")
        linha_operador.pack(fill="x", pady=(2, 0))
        linha_operador.grid_columnconfigure(0, weight=1)

        self.var_operador = tk.StringVar(value=self.config_dados["ultimo_operador"])
        self.combo_operador = ctk.CTkOptionMenu(
            linha_operador,
            values=self.config_dados["operadores"] + [OPCAO_NOVO_OPERADOR],
            variable=self.var_operador, font=FONTE_ENTRY,
            dropdown_font=FONTE_ENTRY, height=32,
            command=self._on_selecionar_operador)
        self.combo_operador.grid(row=0, column=0, sticky="ew")

        ctk.CTkButton(linha_operador, text="🗑", width=36, height=32,
                      font=FONTE_BOTAO, fg_color="#7a2626",
                      command=self._excluir_operador).grid(row=0, column=1, padx=(4, 0))

    # ------------------------------------------------------------------
    # COLUNA DIREITA: PASTA GEOF / DA / SUPER / SIS / SEI
    # ------------------------------------------------------------------
    def _montar_coluna_direita(self, parent):
        col = ctk.CTkFrame(parent, fg_color=COR_CARD, corner_radius=12)
        col.grid(row=0, column=1, sticky="nsew", padx=(8, 0), pady=6)

        pad = {"padx": 16, "pady": (6, 2)}

        # --- PASTA DE TRABALHO (até \GEOF) ---
        ctk.CTkLabel(col, text="PASTA DE TRABALHO (caminho até \\GEOF)",
                     font=FONTE_SECAO, text_color=COR_TEXTO).pack(anchor="w", **pad)
        linha_pasta = ctk.CTkFrame(col, fg_color="transparent")
        linha_pasta.pack(fill="x", padx=16, pady=2)
        linha_pasta.grid_columnconfigure(0, weight=1)
        self.entry_pasta_geof = ctk.CTkEntry(
            linha_pasta, font=FONTE_ENTRY, height=32,
            placeholder_text=r"Cole aqui o caminho até ...\GEOF")
        self.entry_pasta_geof.grid(row=0, column=0, sticky="ew")
        ctk.CTkButton(linha_pasta, text="Procurar...", width=110, height=32,
                      font=FONTE_BOTAO, command=self._procurar_pasta_geof).grid(row=0, column=1, padx=(6, 0))

        # Campo interno (não exibido na tela): guarda o caminho do arquivo
        # .xlsx localizado dentro de GEOF\TIPO\PDPAS ANO\MATRIX.
        self.entry_caminho_completo = ctk.CTkEntry(col)

        # --- DA NOME / DA MATRICULA (lado a lado) ---
        linha_da = ctk.CTkFrame(col, fg_color="transparent")
        linha_da.pack(fill="x", padx=16, pady=2)
        linha_da.grid_columnconfigure((0, 1), weight=1)

        frame_da_nome, self.entry_da_nome = self._criar_campo(linha_da, "DA - NOME (Diretor Administrativo)")
        frame_da_nome.grid(row=0, column=0, sticky="ew", padx=(0, 6))

        frame_da_mat, self.entry_da_matricula = self._criar_campo(
            linha_da, "DA - MATRÍCULA", placeholder="9.999.999-9")
        frame_da_mat.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        self.entry_da_matricula.bind("<KeyRelease>", lambda e: self._aplicar_mascara_matricula(self.entry_da_matricula))

        # --- SUPER NOME / SUPER MATRICULA (lado a lado) ---
        linha_super = ctk.CTkFrame(col, fg_color="transparent")
        linha_super.pack(fill="x", padx=16, pady=2)
        linha_super.grid_columnconfigure((0, 1), weight=1)

        frame_super_nome, self.entry_super_nome = self._criar_campo(linha_super, "SUPER - NOME (Superintendente)")
        frame_super_nome.grid(row=0, column=0, sticky="ew", padx=(0, 6))

        frame_super_mat, self.entry_super_matricula = self._criar_campo(
            linha_super, "SUPER - MATRÍCULA", placeholder="9.999.999-9")
        frame_super_mat.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        self.entry_super_matricula.bind("<KeyRelease>", lambda e: self._aplicar_mascara_matricula(self.entry_super_matricula))

        # --- SIS (calculado, copiável) ---
        ctk.CTkLabel(col, text="SIS", font=FONTE_SECAO,
                     text_color=COR_TEXTO).pack(anchor="w", **pad)
        linha_sis = ctk.CTkFrame(col, fg_color="transparent")
        linha_sis.pack(fill="x", padx=16, pady=2)
        linha_sis.grid_columnconfigure(0, weight=1)
        self.entry_sis = ctk.CTkEntry(linha_sis, font=FONTE_ENTRY, height=32, state="readonly")
        self.entry_sis.grid(row=0, column=0, sticky="ew")
        ctk.CTkButton(linha_sis, text="Copiar", width=80, height=32, font=FONTE_BOTAO,
                      command=lambda: self._copiar_para_area_transferencia(self.entry_sis.get())
                      ).grid(row=0, column=1, padx=(6, 0))

        # --- SEI (vindo da planilha, exibido 1x, copiável) ---
        ctk.CTkLabel(col, text="SEI", font=FONTE_SECAO,
                     text_color=COR_TEXTO).pack(anchor="w", **pad)
        linha_sei = ctk.CTkFrame(col, fg_color="transparent")
        linha_sei.pack(fill="x", padx=16, pady=(2, 8))
        linha_sei.grid_columnconfigure(0, weight=1)
        self.entry_sei = ctk.CTkEntry(linha_sei, font=FONTE_ENTRY, height=32, state="readonly")
        self.entry_sei.grid(row=0, column=0, sticky="ew")
        ctk.CTkButton(linha_sei, text="Copiar", width=80, height=32, font=FONTE_BOTAO,
                      command=lambda: self._copiar_para_area_transferencia(self.entry_sei.get())
                      ).grid(row=0, column=1, padx=(6, 0))

    def _criar_campo(self, parent, texto_label, placeholder=None):
        """Cria um mini-frame com label + entry, usado para campos lado a lado."""
        frame = ctk.CTkFrame(parent, fg_color="transparent")
        ctk.CTkLabel(frame, text=texto_label, font=("Segoe UI", 14, "bold"),
                     text_color=COR_TEXTO).pack(anchor="w")
        entry = ctk.CTkEntry(frame, font=FONTE_ENTRY, height=32,
                              placeholder_text=placeholder or "")
        entry.pack(fill="x", pady=(2, 0))
        return frame, entry

    # ------------------------------------------------------------------
    # TABELA DE RESULTADOS DA COTAÇÃO
    # ------------------------------------------------------------------
    def _montar_tabela(self, parent):
        frame = ctk.CTkFrame(parent, fg_color=COR_CARD, corner_radius=12)
        frame.pack(fill="both", expand=True, padx=20, pady=(0, 6))

        cabecalho = ctk.CTkFrame(frame, fg_color="transparent")
        cabecalho.pack(fill="x", padx=16, pady=(8, 2))
        ctk.CTkLabel(cabecalho, text="ITENS DA COTAÇÃO (aba Controle)",
                     font=FONTE_SECAO, text_color=COR_TEXTO).pack(side="left")

        # --- CÓDIGO selecionado (copiável) - fica acima da tabela para fácil acesso ---
        linha_codigo = ctk.CTkFrame(cabecalho, fg_color="transparent")
        linha_codigo.pack(side="right")
        ctk.CTkLabel(linha_codigo, text="CÓDIGO selecionado:",
                     font=("Segoe UI", 12), text_color=COR_TEXTO_SEC).pack(side="left", padx=(0, 6))
        self.entry_codigo_selecionado = ctk.CTkEntry(linha_codigo, font=FONTE_ENTRY, height=28,
                                                      width=160, state="readonly")
        self.entry_codigo_selecionado.pack(side="left")
        ctk.CTkButton(linha_codigo, text="Copiar", width=70, height=28, font=FONTE_BOTAO,
                      command=lambda: self._copiar_para_area_transferencia(
                          self.entry_codigo_selecionado.get())).pack(side="left", padx=(6, 0))

        style = ttk.Style()
        style.theme_use("default")
        style.configure("Custom.Treeview",
                        background="#1f1f1f", fieldbackground="#1f1f1f",
                        foreground=COR_TEXTO, rowheight=22, font=FONTE_TABELA)
        style.configure("Custom.Treeview.Heading",
                        background="#2f2f2f", foreground=COR_TEXTO,
                        font=("Segoe UI", 12, "bold"))
        style.map("Custom.Treeview", background=[("selected", COR_DESTAQUE)])

        colunas = ("TR", "ORIGEM", "ITEM", "CODIGO", "DESCRICAO", "UN", "QTD")
        self.tabela = ttk.Treeview(frame, columns=colunas, show="headings",
                                    style="Custom.Treeview", height=6)
        larguras = {"TR": 90, "ORIGEM": 100, "ITEM": 60, "CODIGO": 100,
                    "DESCRICAO": 420, "UN": 60, "QTD": 80}
        for c in colunas:
            self.tabela.heading(c, text=c)
            self.tabela.column(c, width=larguras.get(c, 100), anchor="w")
        self.tabela.pack(fill="both", expand=True, padx=16, pady=(0, 8))
        self.tabela.bind("<<TreeviewSelect>>", self._on_selecionar_linha_tabela)
        self.tabela.bind("<Double-1>", self._on_duplo_clique_tabela)

    def _on_selecionar_linha_tabela(self, event=None):
        selecao = self.tabela.selection()
        if not selecao:
            return
        valores = self.tabela.item(selecao[0], "values")
        codigo = valores[3] if len(valores) > 3 else ""  # índice 3 = coluna CODIGO
        self.entry_codigo_selecionado.configure(state="normal")
        self.entry_codigo_selecionado.delete(0, "end")
        self.entry_codigo_selecionado.insert(0, codigo)
        self.entry_codigo_selecionado.configure(state="readonly")

    def _on_duplo_clique_tabela(self, event=None):
        self._on_selecionar_linha_tabela(event)
        codigo = self.entry_codigo_selecionado.get()
        if codigo:
            self._copiar_para_area_transferencia(codigo)

    # ------------------------------------------------------------------
    # GRUPOS DE BOTÕES (parte inferior) — chamam métodos vindos dos mixins
    # ------------------------------------------------------------------
    def _montar_grupos_botoes(self, parent):
        area = ctk.CTkFrame(parent, fg_color=COR_FUNDO_APP)
        area.pack(fill="x", padx=20, pady=(0, 8))
        area.grid_columnconfigure((0, 1, 2, 3), weight=1)

        self._grupo_botoes(area, col=0, cor=COR_GRUPO_1,
                            titulo="1º Passo – Nova cotação",
                            botoes=[("CRIAR PASTA COTAÇÃO", self.acao_criar_pasta_cotacao),
                                    ("PEDIDO", self.acao_salvar_cotacao),
                                    ("AJUSTE DO PEDIDO", self.acao_ajustar_pedido)])

        self._grupo_botoes(area, col=1, cor=COR_GRUPO_2,
                            titulo="2º Passo – Preparar Mapa Comparativo",
                            botoes=[("ACESSAR PROPOSTAS", self.acao_acessar_propostas),
                                    ("LISTAR EMPRESAS", self.acao_listar_empresas),
                                    ("INCLUIR EMPRESAS MAPA", self.acao_incluir_empresas_mapa),
                                    ("CARREGAR EMPRESAS MAPA", self.acao_carregar_empresas_mapa)])

        self._grupo_botoes(area, col=2, cor=COR_GRUPO_3,
                            titulo="3º Passo – Adjudicação e Ordem",
                            botoes=[("ADJUDICAÇÃO", self._acionar_adjudicacao),
                                    ("ORDEM", self.acao_ordem),
                                    ("SALVAR", self.acao_salvar_passo3)])

        self._grupo_botoes(area, col=3, cor=COR_GRUPO_4,
                            titulo="E-mails",
                            botoes=[("E-mail Cotação", self.acao_email_cotacao),
                                    ("E-mail Ordem", self.acao_email_ordem),
                                    ("FERIADOS", self.acao_gerenciar_feriados)])

    def _acionar_adjudicacao(self):
        """Entrada do botão ADJUDICAÇÃO da tela CONTROLE DE COTAÇÕES.

        O módulo AdjudicacaoOrdem carrega o cadastro da cotação e conduz
        cancelamentos, pesquisa/revisão e vencedor de cada item,
        dados do mapa e geração de PDF/JPG/Excel, salvando as decisões da Ordem.
        """
        return self.acao_adjudicacao()

    def _grupo_botoes(self, parent, col, cor, titulo, botoes):
        frame = ctk.CTkFrame(parent, fg_color=cor, corner_radius=12)
        frame.grid(row=0, column=col, sticky="nsew", padx=4, pady=2)

        ctk.CTkLabel(frame, text=titulo, font=FONTE_SECAO,
                     text_color="#1a1a1a", wraplength=220,
                     justify="center").pack(padx=8, pady=(6, 4))

        for texto, comando in botoes:
            ctk.CTkButton(frame, text=texto, font=FONTE_BOTAO,
                          fg_color="#1a1a1a", hover_color="#333333",
                          text_color="#ffffff", height=26,
                          command=comando).pack(fill="x", padx=10, pady=3)

        ctk.CTkLabel(frame, text="", height=1).pack(pady=1)  # respiro inferior

    # ------------------------------------------------------------------
    # PREENCHIMENTO DE DEFAULTS
    # ------------------------------------------------------------------
    def _preencher_defaults(self):
        c = self.config_dados
        self.entry_pasta_geof.insert(0, c.get("pasta_geof", c.get("pasta_trabalho", "")))
        self.entry_ano.insert(0, c["ano"])
        self.var_bimestre.set(c["bimestre"])
        self.entry_da_nome.insert(0, c["da_nome"])
        self.entry_da_matricula.insert(0, c["da_matricula"])
        self.entry_super_nome.insert(0, c["super_nome"])
        self.entry_super_matricula.insert(0, c["super_matricula"])
        self._sufixo_planilha_atual = c.get("ultima_planilha_tipo", "APS")
        self._atualizar_opcoes_nome_planilha()

    def _salvar_defaults_atuais(self):
        self.config_dados.update({
            "pasta_geof": self.entry_pasta_geof.get().strip(),
            "ano": self.entry_ano.get().strip(),
            "bimestre": self.var_bimestre.get(),
            "da_nome": self.entry_da_nome.get().strip(),
            "da_matricula": self.entry_da_matricula.get().strip(),
            "super_nome": self.entry_super_nome.get().strip(),
            "super_matricula": self.entry_super_matricula.get().strip(),
            "ultimo_operador": self.var_operador.get(),
            "ultima_planilha_tipo": self._sufixo_planilha_atual,
        })
        salvar_config(self.config_dados)

    # ------------------------------------------------------------------
    # LOCALIZAÇÃO DO ARQUIVO MATRIX A PARTIR DA PASTA GEOF
    # ------------------------------------------------------------------
    def _montar_caminho_pasta_matrix(self):
        """GEOF + TIPO (conforme NOME PLANILHA) + 'PDPAS ANO' + 'MATRIX'."""
        pasta_geof = self.entry_pasta_geof.get().strip()
        if not pasta_geof:
            return None
        tipo = getattr(self, "_sufixo_planilha_atual", "APS")
        ano = self.entry_ano.get().strip() or "ANO"
        return os.path.join(pasta_geof, tipo, f"PDPAS {ano}", "MATRIX")

    @staticmethod
    def _localizar_arquivo_matrix(pasta_matrix, nome_planilha):
        if not pasta_matrix or not os.path.isdir(pasta_matrix):
            return None
        alvo = nome_planilha.strip().lower()
        variantes = {alvo, alvo.replace(" ", "_"), alvo.replace(" ", "")}
        try:
            arquivos = os.listdir(pasta_matrix)
        except OSError:
            return None
        for arquivo in arquivos:
            nome, ext = os.path.splitext(arquivo)
            if ext.lower() not in (".xlsx", ".xlsm"):
                continue
            nome_norm = nome.strip().lower()
            if (nome_norm in variantes
                    or nome_norm.replace("_", " ") in variantes
                    or nome_norm.replace(" ", "_") in variantes):
                return os.path.join(pasta_matrix, arquivo)
        return None

    def _obter_caminho_arquivo(self):
        """Retorna o caminho do .xlsx já localizado, ou tenta localizar agora.
        O cache só é reaproveitado se PASTA DE TRABALHO, ANO, TIPO (sufixo
        da planilha) e NOME PLANILHA continuarem exatamente os mesmos de
        quando o arquivo foi localizado pela última vez — caso contrário,
        procura de novo do zero. Sem essa checagem, trocar o NOME PLANILHA
        (ex.: de "HRG" para "HRG INV") continuava usando o arquivo antigo
        em cache — já que ele ainda existe em disco — e a consulta parecia
        não encontrar nada, mesmo com dados reais na planilha nova."""
        chave_atual = (
            self.entry_pasta_geof.get().strip(),
            self.entry_ano.get().strip(),
            getattr(self, "_sufixo_planilha_atual", ""),
            self.var_nome_planilha.get(),
        )
        if (self.caminho_arquivo_atual and self._chave_arquivo_atual == chave_atual
                and os.path.exists(self.caminho_arquivo_atual)):
            return self.caminho_arquivo_atual

        pasta_matrix = self._montar_caminho_pasta_matrix()
        caminho = self._localizar_arquivo_matrix(pasta_matrix, self.var_nome_planilha.get())
        if caminho:
            self.caminho_arquivo_atual = caminho
            self._chave_arquivo_atual = chave_atual
        else:
            self.caminho_arquivo_atual = None
            self._chave_arquivo_atual = None
        return caminho

    # ------------------------------------------------------------------
    # ÁREA DE TRANSFERÊNCIA
    # ------------------------------------------------------------------
    def _copiar_para_area_transferencia(self, texto):
        if not texto:
            return
        self.clipboard_clear()
        self.clipboard_append(texto)
        self.update()

    # ------------------------------------------------------------------
    # CONFIGURAÇÃO DOS CONTATOS USADOS NA ORDEM DE COMPRA (por tipo de planilha)
    # ------------------------------------------------------------------
    def _abrir_configuracao_contatos(self):
        contatos_atuais = self.config_dados.get("contatos_ordem", CONTATOS_ORDEM_PADRAO).copy()

        janela = ctk.CTkToplevel(self)
        janela.title("Contatos da Ordem de Compra (por tipo de planilha)")
        janela.configure(fg_color=COR_FUNDO_APP)
        janela.geometry("760x680")
        centralizar_janela(janela)
        janela.grab_set()

        ctk.CTkLabel(
            janela,
            text="Esses são os contatos citados na ORDEM DE COMPRA. Cada tipo de planilha usa um "
                 "conjunto diferente — altere e clique em Salvar; os novos valores ficam salvos "
                 "como padrão até você mudar de novo.",
            font=FONTE_TABELA, text_color=COR_TEXTO_SEC, wraplength=710, justify="left"
        ).pack(anchor="w", padx=16, pady=(16, 10))

        scroll = ctk.CTkScrollableFrame(janela, fg_color="#161616", height=520)
        scroll.pack(fill="both", expand=True, padx=16, pady=(0, 10))

        # (chave, rótulo mostrado, campo mostrado na ordem)
        secoes = [
            ("APS", [
                ("GERENTE", "GERENTE (também usada em APS INV)"),
                ("NFL", "Farmácia - NFL"),
                ("NUAL_APS", "Almoxarifado - NUAL"),
                ("NECFM", "Engenharia Clínica - NECFM"),
            ]),
            ("APS INV", [
                ("GERENTE", "GERENTE (mesma variável da APS)"),
                ("NPDOC_APSINV", "Chefe do NPDOC"),
            ]),
            ("HRG", [
                ("GERENTE_HRG", "GERENTE (também usada em HRG INV)"),
                ("NFH", "Farmácia - NFH"),
                ("NUAL_HRG", "Almoxarifado - NUAL"),
                ("NECFM_HRG", "Engenharia Clínica - NECFM (HRG)"),
            ]),
            ("HRG INV", [
                ("GERENTE_HRG", "GERENTE (mesma variável do HRG)"),
                ("NPDOC_HRGINV", "Chefe do NPDOC"),
            ]),
        ]

        entradas = {}  # chave -> CTkEntry (uma única entry por chave, mesmo repetida em seções)
        for titulo_secao, campos in secoes:
            ctk.CTkLabel(scroll, text=f"SUFIXOS_PLANILHA = {titulo_secao}",
                         font=("Segoe UI", 14, "bold"), text_color=COR_DESTAQUE
                         ).pack(anchor="w", padx=6, pady=(12, 4))
            for chave, rotulo in campos:
                ctk.CTkLabel(scroll, text=rotulo, font=FONTE_TABELA,
                             text_color=COR_TEXTO).pack(anchor="w", padx=16, pady=(4, 2))
                if chave in entradas:
                    # mesma variável (GERENTE) já tem entry — só mostra o valor atual, sem duplicar campo
                    ctk.CTkLabel(scroll, text=f"(mesmo valor de cima: {entradas[chave].get()})",
                                 font=("Segoe UI", 11), text_color=COR_TEXTO_SEC
                                 ).pack(anchor="w", padx=16, pady=(0, 4))
                    continue
                entry = ctk.CTkEntry(scroll, font=FONTE_ENTRY, height=32)
                entry.pack(fill="x", padx=16, pady=(0, 4))
                entry.insert(0, contatos_atuais.get(chave, ""))
                entradas[chave] = entry

        def salvar():
            for chave, entry in entradas.items():
                contatos_atuais[chave] = entry.get().strip()
            self.config_dados["contatos_ordem"] = contatos_atuais
            salvar_config(self.config_dados)
            janela.destroy()

        ctk.CTkButton(janela, text="Salvar", font=FONTE_BOTAO, height=36,
                      fg_color=COR_DESTAQUE, command=salvar).pack(pady=(0, 16))

    # ------------------------------------------------------------------
    # EVENTOS / MÁSCARAS
    # ------------------------------------------------------------------
    def _procurar_pasta_geof(self):
        caminho = filedialog.askdirectory(title="Selecione a pasta GEOF")
        if caminho:
            self.entry_pasta_geof.delete(0, "end")
            self.entry_pasta_geof.insert(0, caminho)

    def _aplicar_mascara_matricula(self, entry):
        texto = entry.get()
        formatado = formatar_matricula(texto)
        if formatado != texto:
            entry.delete(0, "end")
            entry.insert(0, formatado)
            entry.icursor("end")

    def _on_cotacao_digitada(self, indice):
        entry = self.entradas_cotacao[indice]
        texto = formatar_apenas_digitos(entry.get(), 3)
        if texto != entry.get():
            entry.delete(0, "end")
            entry.insert(0, texto)
            entry.icursor("end")
        self._definir_cotacao_ativa(indice)
        if len(texto) == 3:
            # busca automática da cotação ativa assim que os 3 dígitos são digitados
            self._consultar_cotacao_ativa()

    def _on_ano_digitado(self, event=None):
        texto = formatar_apenas_digitos(self.entry_ano.get(), 4)
        if texto != self.entry_ano.get():
            self.entry_ano.delete(0, "end")
            self.entry_ano.insert(0, texto)
            self.entry_ano.icursor("end")
        self._atualizar_opcoes_nome_planilha()

    # --- NOME PLANILHA: opções completas (com ANO) e cores dinâmicas ---
    def _atualizar_opcoes_nome_planilha(self):
        ano = self.entry_ano.get().strip() or "ANO"
        opcoes = [f"Matrix_{ano}_{suf}" for suf in SUFIXOS_PLANILHA_EXIBICAO]
        self.combo_planilha.configure(values=opcoes)
        valor_atual = f"Matrix_{ano}_{getattr(self, '_sufixo_planilha_atual', 'APS')}"
        self.var_nome_planilha.set(valor_atual)
        self._aplicar_cor_nome_planilha(valor_atual)

    def _on_selecionar_nome_planilha(self, valor_completo):
        for suf in SUFIXOS_PLANILHA:
            if valor_completo.endswith(suf):
                self._sufixo_planilha_atual = suf
                break
        self._aplicar_cor_nome_planilha(valor_completo)

    def _aplicar_cor_nome_planilha(self, valor_completo):
        sufixo = "APS"
        for suf in SUFIXOS_PLANILHA:
            if valor_completo.endswith(suf):
                sufixo = suf
                break
        cores = CORES_NOME_PLANILHA[sufixo]
        fonte = ("Segoe UI", 16, "bold") if cores["bold"] else FONTE_ENTRY
        self.combo_planilha.configure(fg_color=cores["fg"], text_color=cores["text"],
                                      button_color=cores["fg"], font=fonte)

    # --- COTAÇÃO ATIVA ---
    def _definir_cotacao_ativa(self, indice):
        if self.indice_cotacao_ativa == indice:
            return
        if self.indice_cotacao_ativa is not None:
            self.entradas_cotacao[self.indice_cotacao_ativa].configure(
                border_color=COR_COTACAO_INATIVA, fg_color=("gray20", "gray20"))
        self.indice_cotacao_ativa = indice
        self.entradas_cotacao[indice].configure(
            border_color=COR_COTACAO_ATIVA, fg_color="#4a4110")

    def _consultar_cotacao_ativa(self):
        """Busca somente a cotação ativa e preenche SEI/SIS/tabela imediatamente."""
        if self.indice_cotacao_ativa is None:
            return
        codigo = self.entradas_cotacao[self.indice_cotacao_ativa].get().strip()
        ano = self.entry_ano.get().strip()
        caminho = self._obter_caminho_arquivo()
        if not codigo or len(codigo) != 3 or not caminho:
            return
        try:
            linhas = buscar_cotacoes_na_planilha(caminho, [codigo], ano)
        except Exception as e:
            messagebox.showerror("Erro ao consultar", str(e))
            return
        if linhas:
            self.linhas_resultado = linhas
            self._preencher_tabela(linhas)
            self._atualizar_sei_e_sis(linhas, ano or linhas[0]["COTACAO_ANO"])

    def _on_selecionar_operador(self, valor):
        if valor == OPCAO_NOVO_OPERADOR:
            self._novo_operador()

    def _novo_operador(self):
        dialogo = ctk.CTkInputDialog(text="Nome do novo operador:", title="Novo operador")
        nome = dialogo.get_input()
        if nome:
            nome = nome.strip().upper()
            if nome and nome not in self.config_dados["operadores"]:
                self.config_dados["operadores"].append(nome)
                salvar_config(self.config_dados)
            self.combo_operador.configure(values=self.config_dados["operadores"] + [OPCAO_NOVO_OPERADOR])
            self.var_operador.set(nome if nome else self.config_dados["operadores"][0])
        else:
            self.var_operador.set(self.config_dados["operadores"][0])

    def _excluir_operador(self):
        atual = self.var_operador.get()
        if atual == OPCAO_NOVO_OPERADOR:
            return
        if len(self.config_dados["operadores"]) <= 1:
            messagebox.showwarning("Atenção", "É necessário manter ao menos um operador.")
            return
        if messagebox.askyesno("Confirmar", f"Excluir o operador '{atual}'?"):
            self.config_dados["operadores"].remove(atual)
            salvar_config(self.config_dados)
            self.combo_operador.configure(values=self.config_dados["operadores"] + [OPCAO_NOVO_OPERADOR])
            self.var_operador.set(self.config_dados["operadores"][0])

    # ------------------------------------------------------------------
    # AÇÃO: CONSULTAR COTAÇÃO (funcional)
    # ------------------------------------------------------------------
    def _atualizar_itens_cotacao_ativa(self):
        """Núcleo comum de CONSULTAR COTAÇÃO e do botão de Reset/Atualizar:
        consulta a aba Controle DO ZERO para a COTAÇÃO ATIVA e atualiza
        ITENS DA COTAÇÃO, SIS e SEI. Nunca reaproveita dados de uma
        consulta anterior — mesmo que PASTA DE TRABALHO, ANO, NOME
        PLANILHA ou a própria cotação ativa tenham mudado desde então, o
        resultado aqui reflete exatamente o que está preenchido na tela
        agora (self._obter_caminho_arquivo() já garante isso ao invalidar
        o cache do arquivo Matrix quando qualquer um desses campos muda).
        Devolve o ANO usado em caso de sucesso (para quem quiser
        encadear mais passos, como abrir o PDF ou recarregar o mapa de
        empresas) ou None se algo impediu a consulta — já mostrando o
        aviso/erro correspondente nesse caso."""
        ano = self.entry_ano.get().strip()
        if not ano or len(ano) != 4:
            messagebox.showwarning("Atenção", "Informe o ANO com 4 dígitos.")
            return None

        caminho = self._obter_caminho_arquivo()
        if not caminho:
            messagebox.showwarning(
                "Atenção",
                "Não foi possível localizar o arquivo Excel.\n"
                "Confira a PASTA (até \\GEOF) e clique em CONSULTAR MATRIX.")
            return None

        if self.indice_cotacao_ativa is None:
            messagebox.showwarning(
                "Atenção", "Clique numa das caixas de COTAÇÃO para defini-la como ativa.")
            return None
        codigo_ativo = self.entradas_cotacao[self.indice_cotacao_ativa].get().strip()
        if len(codigo_ativo) != 3:
            messagebox.showwarning(
                "Atenção", "A COTAÇÃO ativa precisa ter o código completo (3 dígitos).")
            return None

        try:
            linhas = buscar_cotacoes_na_planilha(caminho, [codigo_ativo], ano)
        except Exception as e:
            messagebox.showerror("Erro ao consultar", str(e))
            return None

        if not linhas:
            messagebox.showinfo("Sem resultados",
                                 "Nenhuma linha encontrada na aba 'Controle' para a cotação ativa.")
        self.linhas_resultado = linhas
        self._preencher_tabela(linhas)
        self._atualizar_sei_e_sis(linhas, ano)
        self._salvar_defaults_atuais()
        return ano

    def acao_consultar_cotacao(self):
        """Botão CONSULTAR COTAÇÃO: atualiza ITENS DA COTAÇÃO/SIS/SEI da
        cotação ativa (ver _atualizar_itens_cotacao_ativa) e, além disso,
        abre o PDF do Pedido já gerado (se existir) e recarrega o cadastro
        de empresas do mapa (se existir) — tudo em segundo plano, sem
        travar a consulta caso algum desses dados ainda não exista."""
        ano = self._atualizar_itens_cotacao_ativa()
        if ano is None:
            return
        self._abrir_pdf_pedido_cotacao_ativa(ano)
        self._carregar_empresas_mapa_se_existir(ano)

    def acao_resetar_cotacao_ativa(self):
        """Botão de Reset/Atualizar (ao lado de Contatos da Ordem): mesma
        rotina do CONSULTAR COTAÇÃO, mas sem abrir o PDF do Pedido nem
        recarregar o cadastro de empresas do mapa — atualiza só ITENS DA
        COTAÇÃO, SIS e SEI da cotação ativa. Útil para forçar um refresh
        rápido depois de trocar qualquer informação na tela (COTAÇÃO,
        NOME PLANILHA, ANO, PASTA DE TRABALHO etc.)."""
        self._atualizar_itens_cotacao_ativa()

    def _preencher_tabela(self, linhas):
        for item in self.tabela.get_children():
            self.tabela.delete(item)
        for l in linhas:
            self.tabela.insert("", "end", values=(
                l["TR"], l["ORIGEM"], formatar_numero_inteiro(l["ITEM"]), l["CODIGO"],
                l["DESCRICAO"], l["UN"], formatar_numero_inteiro(l["QTD"], separador_milhar=True)))

    def _atualizar_sei_e_sis(self, linhas, ano):
        self.entry_sei.configure(state="normal")
        self.entry_sei.delete(0, "end")
        self.entry_sis.configure(state="normal")
        self.entry_sis.delete(0, "end")

        if linhas:
            primeira = linhas[0]
            sei = primeira["SEI"] or ""
            tr = primeira["TR"] or ""
            origem = primeira["ORIGEM"] or ""
            cotacao = primeira["COTACAO"] or ""

            self.entry_sei.insert(0, str(sei))

            nome_planilha_completo = self.var_nome_planilha.get()
            sufixo_planilha = nome_planilha_completo[QTD_CHAR_REMOVER_NOME_PLANILHA:]

            sis_texto = (
                f"{cotacao}/{ano} \u2013 COTAÇÃO REF. TERMO DE REFERÊNCIA "
                f"Nº {tr}/{ano} - {origem}/{sufixo_planilha}/SRSSU."
            )
            self.entry_sis.insert(0, sis_texto)

        self.entry_sei.configure(state="readonly")
        self.entry_sis.configure(state="readonly")

    # ------------------------------------------------------------------
    # CONSULTAR COTAÇÃO — conveniências ligadas à cotação ativa
    # ------------------------------------------------------------------
    def _abrir_pdf_pedido_cotacao_ativa(self, ano):
        """Abre automaticamente o PDF do PEDIDO DE COTAÇÃO já gerado (botão
        PEDIDO, em Cotacao.py) para a cotação ativa, se existir. É apenas
        uma conveniência: se a pasta de trabalho não estiver configurada,
        ou o PEDIDO ainda não tiver sido gerado, a consulta continua
        normalmente, sem exibir erro nenhum."""
        if self.indice_cotacao_ativa is None:
            return
        pasta_geof = self.entry_pasta_geof.get().strip()
        tipo = getattr(self, "_sufixo_planilha_atual", "")
        codigo_ativo = self.entradas_cotacao[self.indice_cotacao_ativa].get().strip()
        if not pasta_geof or not tipo or len(codigo_ativo) != 3:
            return

        pasta_cotacao = os.path.join(pasta_geof, tipo, f"PDPAS {ano}", "COTACAO", codigo_ativo)
        # O nome exato do arquivo inclui o TIPO e a ORIGEM (ex.: "PEDIDO DE
        # COTAÇÃO 155 HRG(NFH).pdf" — ver Cotacao.py/acao_salvar_cotacao),
        # então localizamos por padrão (glob) em vez de remontar o nome
        # exato aqui, para não depender de reconstruir a ORIGEM.
        padrao_arquivo = os.path.join(pasta_cotacao, f"PEDIDO DE COTAÇÃO {codigo_ativo} *.pdf")
        candidatos = glob.glob(padrao_arquivo)
        if not candidatos:
            return
        try:
            os.startfile(candidatos[0])
        except Exception:
            pass  # abrir o PDF é só uma conveniência; a consulta já foi concluída

    def _carregar_empresas_mapa_se_existir(self, ano):
        """Recarrega em segundo plano (sem abrir o assistente de cadastro)
        o cadastro de propostas por fornecedor (MAPA_EMPRESAS_{código}.json)
        da cotação ativa, caso ele já exista — deixando self.mapa_fornecedores
        e self.lista_empresas_mapa atualizados junto com a consulta, prontos
        para ADJUDICAÇÃO/ORDEM sem precisar clicar em CARREGAR EMPRESAS MAPA
        manualmente.

        Sempre limpa self.mapa_fornecedores/self.lista_empresas_mapa antes de
        tentar carregar — se a nova cotação ativa ainda não tiver cadastro
        salvo, o cadastro da cotação anterior não fica "pendurado" em
        memória (o que faria ADJUDICAÇÃO/ORDEM operar sobre dados errados)."""
        self.mapa_fornecedores = {}
        self.lista_empresas_mapa = []

        if self.indice_cotacao_ativa is None:
            return
        pasta_geof = self.entry_pasta_geof.get().strip()
        tipo = getattr(self, "_sufixo_planilha_atual", "")
        codigo_ativo = self.entradas_cotacao[self.indice_cotacao_ativa].get().strip()
        if not pasta_geof or not tipo or len(codigo_ativo) != 3:
            return

        pasta_cotacao = os.path.join(pasta_geof, tipo, f"PDPAS {ano}", "COTACAO", codigo_ativo)
        # _caminho_arquivo_mapa_empresas vem do MapaMixin (Mapa.py), já
        # misturado em AppCotacao — reaproveitado aqui para não duplicar a
        # regra de nomeação do arquivo (MAPA_EMPRESAS_{código}.json).
        caminho_json = self._caminho_arquivo_mapa_empresas(pasta_cotacao, codigo_ativo)
        if not os.path.exists(caminho_json):
            return

        try:
            with open(caminho_json, "r", encoding="utf-8") as f:
                dados_salvos = json.load(f)
        except Exception:
            return

        fornecedores = dados_salvos.get("fornecedores", {})
        if not fornecedores:
            return

        self.mapa_fornecedores = fornecedores
        self.lista_empresas_mapa = sorted(fornecedores.keys(), key=normalizar_texto)
        self._pasta_cotacao_atual = pasta_cotacao
        self._pasta_propostas_atual = os.path.join(pasta_cotacao, "PROPOSTAS")
        self._codigo_cotacao_atual = codigo_ativo
        self._ano_atual_wizard = ano

    # ------------------------------------------------------------------
    # HELPER PARA AÇÕES AINDA NÃO ESPECIFICADAS (placeholders)
    # ------------------------------------------------------------------
    def _stub(self, nome_funcao):
        messagebox.showinfo(nome_funcao, f"Função '{nome_funcao}' ainda não implementada.\n"
                                          "Aguardando especificação.")

    def acao_consultar_matrix(self):
        """Localiza e abre a Matrix selecionada no aplicativo de planilhas."""
        pasta_geof = self.entry_pasta_geof.get().strip()
        ano = self.entry_ano.get().strip()
        if not pasta_geof:
            messagebox.showwarning("Atenção", "Informe o caminho da PASTA DE TRABALHO (até \\GEOF).")
            return
        if len(ano) != 4 or not ano.isdigit():
            messagebox.showwarning("Atenção", "Informe o ANO com 4 dígitos.")
            return

        pasta_matrix = self._montar_caminho_pasta_matrix()
        nome_planilha = self.var_nome_planilha.get().strip()
        if not nome_planilha:
            messagebox.showwarning("Atenção", "Selecione o NOME PLANILHA que deseja abrir.")
            return
        caminho = self._localizar_arquivo_matrix(pasta_matrix, nome_planilha)
        chave_atual = (
            pasta_geof, ano, getattr(self, "_sufixo_planilha_atual", ""), nome_planilha)

        self.entry_caminho_completo.configure(state="normal")
        self.entry_caminho_completo.delete(0, "end")
        if caminho:
            self.caminho_arquivo_atual = caminho
            self._chave_arquivo_atual = chave_atual
            self.entry_caminho_completo.insert(0, caminho)
            self.entry_caminho_completo.configure(text_color="#8ee08e")
        else:
            self.caminho_arquivo_atual = None
            self._chave_arquivo_atual = None
            self.entry_caminho_completo.insert(0, f"Não encontrado em: {pasta_matrix}")
            self.entry_caminho_completo.configure(text_color="#f08e8e")
            messagebox.showerror(
                "Arquivo não encontrado",
                f"Não encontrei um arquivo '{nome_planilha}' (.xlsx/.xlsm) em:\n{pasta_matrix}")
        self.entry_caminho_completo.configure(state="readonly")
        self._salvar_defaults_atuais()

        if caminho:
            try:
                self._abrir_planilha_matrix(caminho)
            except (OSError, subprocess.SubprocessError) as erro:
                messagebox.showerror(
                    "Erro ao abrir a Matrix",
                    f"Não foi possível abrir a planilha:\n{caminho}\n\n"
                    f"Confira se o Excel ou outro aplicativo de planilhas está instalado "
                    f"e associado a arquivos .xlsx/.xlsm.\n\nDetalhes: {erro}")

    @staticmethod
    def _abrir_planilha_matrix(caminho):
        """Abre o arquivo existente sem alterar seus dados, imagens ou modelos."""
        caminho = os.path.abspath(caminho)
        if not os.path.isfile(caminho):
            raise FileNotFoundError(f"Planilha não encontrada: {caminho}")
        if sys.platform == "win32":
            os.startfile(caminho)
        else:
            comando = "open" if sys.platform == "darwin" else "xdg-open"
            subprocess.run([comando, caminho], check=True, capture_output=True, text=True)


# ----------------------------------------------------------------------
# COMPOSIÇÃO FINAL — mistura _AppCotacaoBase com os mixins de cada módulo
# ----------------------------------------------------------------------
def criar_app():
    """Importa os mixins de ação (Cotacao.py, Mapa.py, AdjudicacaoOrdem.py,
    Emails.py) e monta a classe final AppCotacao. Feito aqui — e não no
    topo do arquivo — para evitar o import circular explicado acima."""
    from Cotacao import CotacaoMixin
    from Mapa import MapaMixin
    from AdjudicacaoOrdem import AdjudicacaoOrdemMixin
    from Emails import EmailMixin

    class AppCotacao(_AppCotacaoBase, CotacaoMixin, MapaMixin, AdjudicacaoOrdemMixin, EmailMixin):
        """Tela principal completa: _AppCotacaoBase + os métodos de ação de
        cada um dos 4 módulos (1º Passo, 2º Passo, 3º Passo, E-mails)."""
        pass

    return AppCotacao


if __name__ == "__main__":
    AppCotacao = criar_app()
    app = AppCotacao()
    app.mainloop()