# -*- coding: utf-8 -*-
"""
AdjudicacaoOrdem.py
====================
MÓDULO DO 3º PASSO — Adjudicação e Ordem de Compra
-------------------------------------------------------
Responsável pelos botões do grupo "3º Passo – Adjudicação e Ordem":
    - ADJUDICAÇÃO  (monta o Mapa Comparativo: 1º/2º/3º colocado por item,
                     escolha do vencedor, itens cancelados, PDF + JPG)
    - ORDEM        (gera uma Ordem de Compra em PDF + JPG para cada
                     fornecedor vencedor, com os itens que ele ganhou)
    - SALVAR       (placeholder — salva os defaults da tela)

Consome o cadastro de propostas por fornecedor gerado no 2º Passo
(Mapa.py), lido do arquivo COTACAO\\{código}\\MAPA_EMPRESAS_{código}.json,
e grava nele também as decisões de vencedor/cancelamento tomadas aqui, para
que o botão ORDEM não precise perguntar de novo.

gerar_objeto_automatico() — usada aqui para montar o título do Mapa
Comparativo — vive em GeradorCotacao.py (função compartilhada por todos os
módulos, inclusive Emails.py, que a usa para montar o OBJETO de cada
cotação no E-mail Cotação); aqui só é importada, não redefinida.
"""

import os
import json
import re
from statistics import mean, median
from math import isfinite
from decimal import Decimal, ROUND_FLOOR
from xml.sax.saxutils import escape
import tkinter as tk
from Mensagens import messagebox

import customtkinter as ctk

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepInFrame
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT

from GeradorCotacao import (
    centralizar_janela, maximizar_janela,
    COR_FUNDO_APP, COR_CARD, COR_TEXTO, COR_DESTAQUE,
    FONTE_BOTAO, FONTE_SECAO, FONTE_TABELA,
    COR_CINZA_ADJ, COR_CIANO_ADJ, COR_VERDE_ADJ, COR_AMARELO_ADJ,
    CONTATOS_ORDEM_PADRAO,
    formatar_moeda, normalizar_texto, extrair_nome_codigo_fornecedor, truncar_em_palavra,
    formatar_numero_inteiro,
    gerar_objeto_automatico,
    converter_pdf_para_jpg, exibir_preview_pdf_jpg, montar_cabecalho_institucional,
)
from Mapa import carregar_fornecedores_completo

TITULOS_PESQUISA_PADRAO = ("Média¹ Mapa GDF", "Mediana² Painel Federal")
STATUS_ITEM = {
    "DESERTA": "Não compareceu nenhum interessado: nenhum fornecedor apresentou proposta ou manifestou interesse.",
    "FRACASSADA": "Houve participantes, mas nenhuma proposta ou licitante foi considerado válido: propostas desclassificadas por preços incompatíveis ou participantes inabilitados por não cumprirem os requisitos do edital.",
    "REVOGADA": "Processo encerrado por razões de conveniência e oportunidade, por interesse público, diante de fato novo ou superveniente que tornou a contratação desnecessária ou inadequada.",
    "ANULADA": "Anulação obrigatória do procedimento em razão de ilegalidade insanável em uma de suas etapas, vício que não pode ser corrigido.",
}


def texto_status_item(item):
    status = item.get("status_item")
    return f"DISPENSA DE LICITAÇÃO {status}" if status in STATUS_ITEM else ""


def observacoes_status_itens(itens):
    return "\n".join(f"ITEM {formatar_numero_inteiro(it['numero_item'])} — {texto_status_item(it)}. {STATUS_ITEM[it['status_item']]}"
                     for it in itens if texto_status_item(it))


def corrigir_textos_relatorio(texto):
    """Atualiza também textos salvos, preservando as demais edições do usuário."""
    texto = re.sub(r'(\bITEM\s+)(\d+)\.0+\b', r'\1\2', texto)
    texto = re.sub(r'(para o\(s\) item\(ns\) )([\d., ]+)(?=, uma vez)',
                   lambda m: m[1] + re.sub(r'\b(\d+)\.0+\b', r'\1', m[2]), texto)
    texto = texto.replace(
        'Atenção: conjunto com menos de três valores. A justificativa e os documentos de suporte devem ser conferidos pelo responsável.',
        'Atenção: conjunto com menos de três valores. A justificativa e os documentos de suporte fazem parte do processo.')
    texto = texto.replace(
        'presente cotação, solicitamos de V.Sª o respaldo administrativo para declarar vencedora, por terem apresentado o menor preço e/ou menor prazo de entrega, a(s) seguinte(s) empresa(s) para o(s) respectivo(s) item(ns):',
        'presente cotação.')
    return texto.replace('A comprovação das fontes, das propostas e dos pareceres mencionados deve acompanhar o processo.',
                         'A comprovação das fontes, das propostas e dos pareceres mencionados acompanham o processo.')


def atualizar_status_relatorio(texto, texto_auto):
    """Atualiza os desfechos automáticos sem substituir as demais edições."""
    texto = corrigir_textos_relatorio(texto)
    texto_auto = corrigir_textos_relatorio(texto_auto)
    situacoes = {}
    vencedores = {}
    item = None
    em_itens = False
    observacoes = []
    for linha in texto_auto.splitlines():
        if linha == "3. ANÁLISE POR ITEM":
            em_itens = True
        elif linha.startswith("4. PESQUISA"):
            em_itens = False
        if em_itens and linha.startswith("ITEM "):
            item = linha
        if em_itens and linha.startswith("Situação:"):
            situacoes[item] = linha
        if em_itens and linha.startswith("Empresa vencedora"):
            vencedores[item] = linha
        if linha.startswith("ITEM ") and " — DISPENSA DE LICITAÇÃO " in linha and not em_itens:
            observacoes.append(linha)
    saida = []
    em_itens = False
    em_observacoes = False
    item = None
    for linha in texto.splitlines():
        if linha == "3. ANÁLISE POR ITEM":
            em_itens = True
        elif linha.startswith("4. PESQUISA"):
            em_itens = False
        elif linha == "5. OBSERVAÇÕES DO MAPA COMPARATIVO":
            em_observacoes = True
        elif linha == "6. CONCLUSÃO":
            if em_observacoes:
                saida.extend(observacoes)
            em_observacoes = False
        if em_itens and linha.startswith("ITEM "):
            item = linha
            saida.append(linha)
            if "DISPENSA DE LICITAÇÃO" in situacoes.get(item, ""):
                saida.append(situacoes[item])
            continue
        if em_itens and linha.startswith("Situação:"):
            if "DISPENSA DE LICITAÇÃO" in situacoes.get(item, ""):
                continue
            if linha.startswith("Situação: DISPENSA DE LICITAÇÃO"):
                linha = situacoes.get(item, "Situação: item em análise na adjudicação.")
        if em_itens and linha.startswith("Empresa vencedora"):
            linha = vencedores.get(item, linha)
        if em_observacoes and linha.startswith("ITEM ") and " — DISPENSA DE LICITAÇÃO " in linha:
            continue
        if em_observacoes and observacoes and linha == "Nenhuma observação registrada.":
            continue
        saida.append(linha)
    return "\n".join(saida)


def adicionar_dica_status(botao, texto):
    estado = {"job": None, "janela": None}

    def ocultar(event=None):
        if estado['job'] is not None:
            botao.after_cancel(estado['job'])
            estado['job'] = None
        if estado['janela'] is not None:
            estado['janela'].destroy()
            estado['janela'] = None

    def mostrar():
        estado['job'] = None
        dica = ctk.CTkToplevel(botao)
        estado['janela'] = dica
        dica.overrideredirect(True)
        dica.geometry(f"+{botao.winfo_rootx()}+{max(0, botao.winfo_rooty()-115)}")
        ctk.CTkLabel(dica, text=texto, wraplength=440, justify="left",
                     font=("Segoe UI", 15)).pack(padx=12, pady=10)

    def agendar(event=None):
        ocultar()
        estado['job'] = botao.after(450, mostrar)

    botao.bind("<Enter>", agendar)
    botao.bind("<Leave>", ocultar)
    botao.bind("<ButtonPress-1>", ocultar)
    botao.bind("<Destroy>", ocultar)
TEXTO_PESQUISA_PADRAO = 'Como pesquisa de preços foram consultados:\n1 - O preço médio do Mapa de Preços do GDF, conforme art. 88, inciso I e art. 90, § 1º, do Decreto 44.330/2023;\n2 - A mediana do Painel de Preços Federal no Compras.gov, conforme o art. 23, § 1º, inciso I, da Lei nº 14.133/2021 e art. 88, inciso II, do Decreto 44.330/2023;\n3 - Pesquisa direta com fornecedores, conforme art. 88, inciso IV e art. 108 do Decreto 44.330/2023.\n4 - Considerando o art. 90, § 3º do Decreto 44.330/2023 "Na impossibilidade de cumprimento do disposto no caput, deverá ser juntada aos autos comprovação e justificativa do agente responsável."\n5 - O preço de referência foi fixado com base no art. 102 Decreto 44.330/2023 "O valor de referência de cada item será o menor preço obtido após o cálculo da média e mediana final dos valores válidos contidos na pesquisa de preços."\nEm função do que se apresenta na planilha acima, das necessidades expostas no Termo de Referência e que o(s) produto(s) ofertado(s) atende(m) ao(s) requisito(s) de melhor proposta, prevista no edital referente à presente cotação, solicitamos de V.Sª o respaldo administrativo para declarar vencedora, por terem apresentado o menor preço e/ou menor prazo de entrega, a(s) seguinte(s) empresa(s) para o(s) respectivo(s) item(ns):'

def montar_texto_pesquisa_justificativa(mapa_itens):
    observacoes = [f"ITEM {item['numero_item']} {item['observacaoitem'].strip()}"
                  for item in mapa_itens if not item.get("cancelado") and item.get("observacaoitem", "").strip()]
    return TEXTO_PESQUISA_PADRAO + ("\n\n" + "\n".join(observacoes) if observacoes else "")


def valores_pesquisa(item):
    """Garante as duas posições mesmo em cadastros antigos ou itens cancelados."""
    valores = list(item.get("pesquisas") or [])
    return (valores + [None, None])[:2]


def montar_texto_relatorio_pesquisa(itens, codigo, ano, dados, sei=""):
    """Relata os dados efetivamente usados, sem presumir consultas ou pareceres."""
    def moeda(valor):
        if valor is None:
            return "não informado"
        return "R$ " + f"{valor:,.4f}".replace(",", "X").replace(".", ",").replace("X", ".")

    linhas = ["RELATÓRIO DE PESQUISA DE PREÇOS", f"Cotação: {codigo}/{ano}",
              f"Processo/publicação SEI: {sei or 'não informado'}", "", "1. OBJETO DA PESQUISA",
              dados.get("objeto") or "Objeto descrito nos itens abaixo.", "", "2. METODOLOGIA E CRITÉRIOS",
              "Este relatório apresenta as informações registradas na adjudicação. Os valores são unitários. "
              "Cada pesquisa externa preenchida representa um valor no conjunto, assim como cada proposta de fornecedor mantida. "
              "Campo sem valor informado não participa do cálculo.",
              "A média é a soma dos valores dividida pela quantidade de valores.",
              "A mediana é o valor central da lista ordenada; se a quantidade for par, é a média dos dois valores centrais.",
              "O sistema adota como referência o menor resultado entre média e mediana.",
              "Os limites exibidos correspondem à referência multiplicada por 1,5 (Preço exorbitante) e por 0,5 (Preço inexequível).",
              "Após uma retirada, as medidas são recalculadas. Os limites finais podem diferir daqueles observados no momento da retirada.",
              "Os fundamentos e as justificativas informados pelo responsável estão transcritos ao final.",
              "", "3. ANÁLISE POR ITEM"]
    for item in itens:
        linhas += ["", f"ITEM {item['numero_item']} — {item['codigo']}",
                   f"Descrição: {item['descricao']}", f"Unidade: {item.get('un') or 'não informada'}"]
        if item.get("cancelado"):
            linhas.append("Situação: CANCELADO. Não participa do fechamento da adjudicação; não há conclusão de referência ou vencedor para este item neste relatório.")
        elif texto_status_item(item):
            linhas.append(f"Situação: {texto_status_item(item)}. {STATUS_ITEM[item['status_item']]}")
        else:
            linhas.append("Situação: item em análise na adjudicação.")
        linhas.append("Fontes externas registradas:")
        titulos = item.get("titulos_pesquisa") or dados.get("titulos_pesquisa") or TITULOS_PESQUISA_PADRAO
        pesquisas = valores_pesquisa(item)
        for i, titulo in enumerate(titulos):
            valor = pesquisas[i] if i < len(pesquisas) else None
            linhas.append(f"- {titulo}: {moeda(valor)}" + (" — não integra o cálculo." if valor is None else "."))
        linhas.append("Propostas mantidas após a revisão:")
        propostas = item["propostas"]
        for proposta in propostas:
            linhas.append(f"- {proposta['fornecedor']}: {moeda(proposta['unitario'])}; marca: {proposta.get('marca') or 'não informada'}.")
        if not propostas:
            linhas.append("Nenhuma proposta mantida.")
        historico = item.get("precos_retirados") or []
        for grupo, titulo in [("superior", "FORNECEDORES RETIRADOS - 50% ACIMA"),
                              ("inferior", "FORNECEDORES RETIRADOS - 50% ABAIXO"),
                              ("analise", "FORNECEDORES RETIRADOS - ANÁLISE DA PROPOSTA")]:
            linhas.append(titulo)
            registros = [p for p in historico if p.get("grupo") == grupo]
            for p in registros:
                motivo = p.get("justificativa") or ("Preço acima de 50% da referência" if grupo == "superior" else
                          "Preço abaixo de 50% da referência" if grupo == "inferior" else "Justificativa não registrada")
                linhas.append(f"- {p['fornecedor']}: {moeda(p['unitario'])}. Justificativa: {motivo}. Não participa da média e da mediana finais.")
            if not registros:
                linhas.append("Nenhum registro neste grupo.")
        registrados = {p.get("fornecedor_chave") for p in historico}
        sem_historico = [chave for chave in item.get("fornecedores_retirados", []) if chave not in registrados]
        if sem_historico:
            linhas.append("Retiradas antigas sem detalhamento completo: " + ", ".join(sem_historico) + ". Consultar documentação de suporte para valores e motivos.")
        linhas.append("OBSERVAÇÃOITEM: " + (item.get("observacaoitem") or "Nenhuma observação registrada."))
        if item.get("cancelado"):
            continue
        valores = sorted([p["unitario"] for p in propostas] + [v for v in pesquisas if v is not None])
        linhas.append(f"Conjunto final: {len(propostas)} proposta(s) mantida(s) e {sum(v is not None for v in pesquisas)} valor(es) de pesquisas externas, totalizando {len(valores)} valor(es).")
        if not valores:
            linhas.append("Não há valores para calcular média, mediana ou referência. Item sem vencedor.")
            continue
        media, mediana = mean(valores), median(valores)
        referencia = min(media, mediana)
        linhas += ["Valores em ordem crescente: " + "; ".join(moeda(v) for v in valores) + ".",
                   f"Soma: {moeda(sum(valores))}. Média: {moeda(sum(valores))} ÷ {len(valores)} = {moeda(media)}."]
        meio = len(valores) // 2
        linhas.append(f"Mediana: termo central, posição {meio+1}: {moeda(mediana)}." if len(valores) % 2 else
                      f"Mediana: média dos termos nas posições {meio} e {meio+1}: ({moeda(valores[meio-1])} + {moeda(valores[meio])}) ÷ 2 = {moeda(mediana)}.")
        linhas += [f"REFERÊNCIA FINAL: menor entre média ({moeda(media)}) e mediana ({moeda(mediana)}) = {moeda(referencia)}.",
                   f"Limites finais de revisão: inferior = {moeda(referencia)} × 0,5 = {moeda(referencia*0.5)}; superior = {moeda(referencia)} × 1,5 = {moeda(referencia*1.5)}."]
        if len(valores) < 3:
            linhas.append("Atenção: conjunto com menos de três valores. A justificativa e os documentos de suporte fazem parte do processo.")
        idx = item.get("vencedor_idx")
        if not texto_status_item(item) and isinstance(idx, int) and 0 <= idx < len(propostas):
            p = propostas[idx]
            linhas.append(f"Empresa vencedora selecionada: {p['fornecedor']} — {moeda(p['unitario'])}. A referência estatística e o preço da vencedora são informações distintas.")
        else:
            linhas.append("Empresa vencedora não selecionada.")
    # O mapa inclui as observações automaticamente na justificativa; no relatório
    # elas já estão na análise individual. Remove somente esses blocos conhecidos.
    justificativa = dados.get("texto_pesquisa") or ""
    for item in itens:
        observacao = (item.get("observacaoitem") or "").strip()
        if observacao:
            bloco = f"ITEM {item['numero_item']} {observacao}"
            justificativa = justificativa.replace(bloco, "")
    linhas += ["", "4. PESQUISA DE PREÇOS E JUSTIFICATIVA", justificativa.strip() or "Não informada.",
               "", "5. OBSERVAÇÕES DO MAPA COMPARATIVO", dados.get("observacoes") or "Nenhuma observação registrada.",
               "", "6. CONCLUSÃO", "As referências e as decisões por item estão discriminadas acima. As propostas retiradas não integram os cálculos finais. "
               "A comprovação das fontes, das propostas e dos pareceres mencionados deve acompanhar o processo."]
    return corrigir_textos_relatorio("\n".join(linhas))


def gerar_pdf_relatorio_pesquisa(caminho, texto):
    """Exporta exatamente o texto revisado, com paginação e quebra de parágrafos."""
    doc = SimpleDocTemplate(caminho, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=1.7*cm, bottomMargin=1.7*cm,
                            title="Relatório de Pesquisa de Preços", author="GEOF")
    styles = getSampleStyleSheet()
    corpo = ParagraphStyle("RelatorioCorpo", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=6)
    titulo = ParagraphStyle("RelatorioTitulo", parent=styles["Heading2"], fontSize=12, leading=16, spaceBefore=10)
    story = [montar_cabecalho_institucional(A4[0]-4*cm, ["GOVERNO DO DISTRITO FEDERAL",
        "SECRETARIA DE ESTADO DE SAÚDE", "GEOF — PESQUISA DE PREÇOS"]), Spacer(1, 12)]
    for linha in texto.splitlines():
        if not linha.strip():
            story.append(Spacer(1, 6))
            continue
        destaque = linha.startswith(("ITEM ", "RELATÓRIO DE", "1. ", "2. ", "3. ", "4. ", "5. ", "6. "))
        story.append(Paragraph(escape(linha), titulo if destaque else corpo))
    def pagina(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(A4[0]-2*cm, cm, f"Página {documento.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=pagina, onLaterPages=pagina)


def numero_pesquisa(texto):
    texto = str(texto).strip().replace(" ", "")
    if not texto:
        return None
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    valor = float(texto)
    if not isfinite(valor) or valor <= 0:
        raise ValueError("Informe um valor positivo ou deixe o campo vazio.")
    return valor


def formatar_pesquisa(valor):
    return "" if valor is None else formatar_moeda(valor).replace("R$ ", "")


def recalcular_referencias(item):
    valores = [p["unitario"] for p in item["propostas"]]
    valores += [v for v in valores_pesquisa(item) if v is not None]
    media = mean(valores) if valores else None
    mediana = median(valores) if valores else None
    referencia = min(media, mediana) if valores else None
    item["referencias"] = dict(media=media, mediana=mediana, minimo=referencia,
                               superior=referencia * 1.5 if valores else None,
                               inferior=referencia * 0.5 if valores else None)


def simular_referencia(valores, indice_negociacao):
    """Calcula um cenário independente; None representa um valor excluído.

    x <= média exige x <= média dos outros valores. Para x <= mediana,
    o teto é o elemento central superior dos outros valores ordenados.
    O menor dos dois tetos satisfaz a regra da referência, sem iterações.
    """
    numeros = [Decimal(str(v)) if v is not None else None for v in valores]
    if any(v is not None and (not v.is_finite() or v <= 0) for v in numeros):
        raise ValueError("Use apenas valores positivos e finitos.")
    incluidos = [v for v in numeros if v is not None]
    media = sum(incluidos) / len(incluidos) if incluidos else None
    mediana = median(incluidos) if incluidos else None
    referencia = min(media, mediana) if incluidos else None
    alvo = numeros[indice_negociacao] if 0 <= indice_negociacao < len(numeros) else None
    outros = sorted(v for i, v in enumerate(numeros) if v is not None and i != indice_negociacao)
    limite = None
    if alvo is not None and outros:
        teto = min(sum(outros) / len(outros), outros[len(outros) // 2])
        limite = teto.quantize(Decimal('0.01'), rounding=ROUND_FLOOR)
    return dict(media=media, mediana=mediana, referencia=referencia, limite=limite,
                alvo=alvo, aceita=(alvo <= referencia if alvo is not None else None),
                quantidade=len(incluidos))


def aplicar_retiradas(item, fornecedores=(), pesquisas=()):
    removidos = set(item.get("fornecedores_retirados", [])) | set(fornecedores)
    item["fornecedores_retirados"] = sorted(removidos)
    item["propostas"] = [p for p in item["propostas"] if p.get("fornecedor_chave", p["fornecedor"]) not in removidos]
    item["pesquisas"] = valores_pesquisa(item)
    for indice in pesquisas:
        item["pesquisas"][indice] = None
    item["vencedor_idx"] = 0 if item["propostas"] and not texto_status_item(item) else None
    recalcular_referencias(item)


def retornar_fornecedor(item, chave):
    """Reintegra somente a proposta solicitada e desfaz seus registros de retirada."""
    def identificador(proposta):
        return proposta.get("fornecedor_chave", proposta["fornecedor"])

    original = next((p for p in item.get("propostas_originais", []) if identificador(p) == chave), None)
    if original is None:
        return False
    propostas = list(item["propostas"])
    idx = item.get("vencedor_idx")
    escolhido = propostas[idx] if idx is not None and 0 <= idx < len(propostas) else None
    if not any(identificador(p) == chave for p in propostas):
        propostas.append(original)
    propostas.sort(key=lambda p: p["unitario"])
    item["propostas"] = propostas
    item["fornecedores_retirados"] = [c for c in item.get("fornecedores_retirados", []) if c != chave]
    item["precos_retirados"] = [p for p in item.get("precos_retirados", []) if identificador(p) != chave]
    item["retiradas_pendentes"] = [r for r in item.get("retiradas_pendentes", []) if r[1] != chave]
    item.setdefault("justificativas_retirada", {}).pop(chave, None)
    item["vencedor_idx"] = (next((i for i, p in enumerate(propostas) if p is escolhido), 0)
                             if not texto_status_item(item) else None)
    recalcular_referencias(item)
    return True


def gerar_excel_mapa_comparativo(caminho, codigo, ano, itens, dados):
    """Exporta o modelo A:M; limites em N ficam fora da área de impressão."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    if len(itens) > 12:
        # Cada folha mantém a área exigida, sem cortar itens de cotações maiores.
        from copy import copy
        from tempfile import TemporaryDirectory
        from openpyxl import load_workbook
        wb = Workbook()
        wb.remove(wb.active)
        with TemporaryDirectory() as pasta:
            for pagina, inicio in enumerate(range(0, len(itens), 12), 1):
                arquivo = os.path.join(pasta, f"mapa{pagina}.xlsx")
                gerar_excel_mapa_comparativo(arquivo, codigo, ano, itens[inicio:inicio+12], dados)
                origem = load_workbook(arquivo)
                fonte_nome = f"Valores classificados {pagina}"
                for original in origem:
                    destino = wb.create_sheet(f"Adjudicação {pagina}" if original.title == "Adjudicação" else fonte_nome)
                    for linha in original:
                        for celula in linha:
                            valor = celula.value
                            if isinstance(valor, str) and celula.data_type == 'f':
                                valor = valor.replace("'Valores classificados'!", f"'{fonte_nome}'!")
                            nova = destino.cell(celula.row, celula.column, valor)
                            nova.font = copy(celula.font)
                            nova.fill = copy(celula.fill)
                            nova.border = copy(celula.border)
                            nova.alignment = copy(celula.alignment)
                            nova.protection = copy(celula.protection)
                            nova.number_format = celula.number_format
                    for area in original.merged_cells.ranges:
                        destino.merge_cells(str(area))
                    for chave, dim in original.row_dimensions.items():
                        destino.row_dimensions[chave] = copy(dim)
                    for chave, dim in original.column_dimensions.items():
                        destino.column_dimensions[chave] = copy(dim)
                    destino.page_setup = copy(original.page_setup)
                    destino.sheet_properties = copy(original.sheet_properties)
                    destino.print_area = original.print_area
                origem.close()
        wb.save(caminho)
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudicação"
    for linha, texto in [(1, "GOVERNO DO DISTRITO FEDERAL"),
                         (2, "SECRETARIA DE ESTADO DE SAÚDE"),
                         (3, "PROGRAMA DE DESCENTRALIZAÇÃO PROGRESSIVA DE AÇÕES DE SAÚDE — PDPAS"),
                         (4, "SOLICITAÇÃO DE ADJUDICAÇÃO PARA O PEDIDO DE COTAÇÃO PDPAS"),
                         (5, f"{codigo}/{ano}"), (6, dados.get('objeto', '')),
                         (7, f"MAPA COMPARATIVO {codigo}/{ano}")]:
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=13)
        ws.cell(linha, 1, texto)
    titulos = dados.get('titulos_pesquisa', TITULOS_PESQUISA_PADRAO)
    ws['A8'] = 'Item / Código'
    for c in (2, 5, 8):
        ws.merge_cells(start_row=8, start_column=c, end_row=8, end_column=c+1)
        ws.cell(8, c, f'{(c-2)//3+1}º Colocado / Unitário / Total')
        ws.cell(8, c+2, 'Qtd./Und./Marca')
    ws['K8'], ws['L8'], ws['M8'] = *titulos, 'MÉDIA'
    fontes = wb.create_sheet("Valores classificados")
    fontes.append(["Item", "Fornecedor", "Unitário"])
    linha_fonte = 2
    for indice, item in enumerate(itens):
        r = 9 + indice * 4
        propostas = [] if item.get('cancelado') or texto_status_item(item) else item['propostas']
        ws.cell(r, 1, item['numero_item'])
        ws.cell(r+2, 1, item['codigo'])
        for pos in range(3):
            c = 2 + pos * 3
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c+1)
            if pos < len(propostas):
                proposta = propostas[pos]
                ws.cell(r, c, proposta['fornecedor'])
                ws.cell(r, c+2, proposta['quantidade'])
                ws.cell(r+1, c, proposta['unitario'])
                ws.cell(r+1, c+1, proposta['total'])
                ws.cell(r+1, c+2, item['un'])
                ws.cell(r+2, c, 'Prazo (dias)')
                ws.cell(r+2, c+1, proposta.get('prazo'))
                ws.cell(r+2, c+2, proposta.get('marca') or '')
            else:
                ws.cell(r, c, 'CANCELADO' if item.get('cancelado') else texto_status_item(item) or ('DESERTO' if not propostas else 'PROPOSTA NÃO APRESENTADA'))
        ws.merge_cells(start_row=r+3, start_column=1, end_row=r+3, end_column=3)
        ws.cell(r+3, 1, 'Saldo remanescente')
        ws.cell(r+3, 4, item['saldo'])
        ws.merge_cells(start_row=r+3, start_column=5, end_row=r+3, end_column=10)
        ws.cell(r+3, 5, item['descricao'])
        for j in range(2):
            ws.merge_cells(start_row=r, start_column=11+j, end_row=r+2, end_column=11+j)
            ws.cell(r, 11+j, valores_pesquisa(item)[j])
        inicio = linha_fonte
        # Inclui todos os fornecedores válidos, não apenas os três exibidos.
        for proposta in item['propostas']:
            fontes.append([str(item['numero_item']), proposta['fornecedor'], proposta['unitario']])
            linha_fonte += 1
        argumentos = [f'K{r}', f'L{r}']
        if linha_fonte > inicio:
            argumentos.insert(0, f"'Valores classificados'!C{inicio}:C{linha_fonte-1}")
        args = ','.join(argumentos)
        ws.cell(r, 13, f'=IF(COUNT({args})=0,"",AVERAGE({args}))')
        ws.cell(r+1, 13, 'MEDIANA')
        ws.cell(r+2, 13, f'=IF(COUNT({args})=0,"",MEDIAN({args}))')
        ws.cell(r+3, 13, f'=IF(COUNT(M{r},M{r+2})=0,"",MIN(M{r},M{r+2}))')
        ws.cell(r+2, 14, f'=IF(M{r+3}="","",M{r+3}*1.5)')
        ws.cell(r+3, 14, f'=IF(M{r+3}="","",M{r+3}*0.5)')
        ws.cell(r+2, 15, '50% ACIMA')
        ws.cell(r+3, 15, '50% ABAIXO')
        ws.merge_cells(start_row=r+3, start_column=11, end_row=r+3, end_column=12)
        ws.cell(r+3, 11, 'Preço de referência')
    ws.merge_cells('A57:M57')
    ws['A57'] = dados.get('texto_pesquisa', TEXTO_PESQUISA_PADRAO)
    ws.row_dimensions[57].height = 170
    for r, (numeros, vencedor) in enumerate(montar_declaracao_vencedores(itens), 59):
        ws.cell(r, 1, ', '.join(map(str, numeros)))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
        ws.cell(r, 2, vencedor)
    rodape = [(73, dados.get('observacoes', '')),
              (76, 'Aceita essa solicitação, procederemos à etapa do envio de Ordem de Compra do material cotado.'),
              (79, 'Atenciosamente, COMISSÃO DO PDPAS.'),
              (82, 'ADJUDICAÇÃO E HOMOLOGAÇÃO'),
              (83, 'De acordo, em função do exposto acima.'),
              (85, 'Adjudico e Homologo o presente. Proceda-se à desaplicação dos recursos necessários para fazer face ao pagamento da presente aquisição, assim que os materiais forem entregues e a NOTA FISCAL for atestada.'),
              (88, 'Em tempo, solicitamos envio ao NPDOC para acompanhamento do prazo de entrega.'),
              (90, 'Atendendo ao art. 16. Inciso IV da Portaria SES 473, de 04/12/2023 Declaro que há disponibilidade financeira para cobrir a presente contratação.'),
              (92, f"{dados.get('diretor_nome', '')} — {dados.get('diretor_matricula', '')}")]
    for r, texto in rodape:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
        ws.cell(r, 1, texto)
    ws.row_dimensions[73].height = 55 if dados.get('observacoes') else 8
    ws.row_dimensions[85].height = 30
    ws.row_dimensions[90].height = 25
    lado = Side(style='thin', color='000000')
    borda = Border(left=lado, right=lado, top=lado, bottom=lado)
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            cell.font = Font(size=8)
            if cell.column >= 11 or (cell.row >= 9 and (cell.row-9) % 4 == 1 and cell.column in (2, 3, 5, 6, 8, 9)) or (cell.row >= 9 and (cell.row-9) % 4 == 3 and cell.column == 4):
                cell.number_format = '#,##0.00'
            if 8 <= cell.row <= 8 + 4*len(itens) and cell.column <= 13:
                cell.border = borda
                if cell.row == 8:
                    cell.fill = PatternFill('solid', fgColor='A5A5A5')
                elif (cell.row - 9) % 4 == 3 and cell.column >= 11:
                    cell.fill = PatternFill('solid', fgColor='00FFFF')
    for r in (4, 7, 82):
        ws.cell(r, 1).font = Font(size=12, bold=True)
        ws.cell(r, 1).fill = PatternFill('solid', fgColor='00FFFF')
    ws['A57'].alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
    for c in range(1, 16):
        ws.column_dimensions[get_column_letter(c)].width = 12 if c in (2, 5, 8, 11, 12, 13) else 9
    ws.row_dimensions[8].height = 42
    ws.print_area = 'A1:M92'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    wb.save(caminho)


TEXTO_OBSERVACAO_PADRAO = (
    "A cotação foi finalizada com menos de três propostas para o(s) item(ns) {itens}, uma vez "
    "que, realizada a pesquisa de mercado, não foram apresentadas outras propostas, e conforme "
    "pesquisas realizadas no Painel de Preços Federal, Mapa de Preços do DF o preço apresentado "
    "encontra-se em consonância com a realidade de mercado atual e dentro do limite estabelecido "
    "na tabela CMED.\n\n"
    "Foram apresentadas notas fiscais emitidas para outros contratantes, comprovando previamente "
    "que os preços estão em conformidade com os praticados semelhantes de objetos de mesma "
    "natureza, conforme a Lei 14.133/2021 Art. 23 §4."
)


# ----------------------------------------------------------------------
# MONTAGEM DOS DADOS DO MAPA COMPARATIVO
# ----------------------------------------------------------------------
def montar_mapa_itens_completo(linhas_resultado, mapa_fornecedores, fornecedores_texto_completo,
                                saldos_bd):
    """Monta a lista COMPLETA de itens da cotação (na ordem original do
    Controle), cada um com suas propostas ordenadas por menor preço — usada
    para o Mapa Comparativo, que precisa mostrar até itens sem proposta
    ('DESERTO'), diferente do ranking do Passo 26 (Mapa.py)."""
    resultado = []
    for item in linhas_resultado:
        codigo = str(item.get("CODIGO") or "")
        propostas = []
        for nome_forn, dados in mapa_fornecedores.items():
            texto_completo = fornecedores_texto_completo.get(normalizar_texto(nome_forn), nome_forn)
            for it in dados.get("itens", []):
                if it.get("codigo") != codigo:
                    continue
                unitario = it.get("unitario")
                if unitario is None or unitario <= 0:
                    continue
                propostas.append({
                    "fornecedor": texto_completo,
                    "fornecedor_chave": normalizar_texto(nome_forn),
                    "quantidade": it.get("quantidade"),
                    "unitario": unitario,
                    "total": it.get("total"),
                    "prazo": it.get("prazo"),
                    "marca": it.get("marca"),
                })
        propostas.sort(key=lambda p: p["unitario"])
        resultado.append({
            "numero_item": item.get("ITEM"),
            "codigo": codigo,
            "descricao": item.get("DESCRICAO") or "",
            "un": item.get("UN") or "",
            "propostas": propostas,
            "saldo": saldos_bd.get(normalizar_texto(codigo)),
            "vencedor_idx": 0 if propostas else None,  # padrão = menor preço; pode ser alterado
        })
    return resultado


def montar_declaracao_vencedores(mapa_itens):
    """Agrupa TODOS os itens do mesmo vencedor numa única linha (não só os
    consecutivos!) — por padrão o 1º colocado, mas respeita a escolha manual
    em item['vencedor_idx']. Os números dos itens saem em ordem crescente
    ('1, 3, 6' -> ATS), e cada empresa aparece uma única vez na tabela,
    ordenada pela primeira aparição, como no modelo de Adjudicação."""
    numeros_por_vencedor = {}
    ordem_primeira_aparicao = []
    for item in mapa_itens:
        if item.get("cancelado") or texto_status_item(item):
            continue
        propostas = item["propostas"]
        if not propostas:
            continue
        idx = min(item.get("vencedor_idx", 0) or 0, len(propostas) - 1)
        vencedor = propostas[idx]["fornecedor"]
        numero = item["numero_item"]
        if vencedor not in numeros_por_vencedor:
            numeros_por_vencedor[vencedor] = []
            ordem_primeira_aparicao.append(vencedor)
        numeros_por_vencedor[vencedor].append(numero)

    grupos = []
    for vencedor in ordem_primeira_aparicao:
        numeros_ordenados = sorted(numeros_por_vencedor[vencedor],
                                   key=lambda n: (n is None, n))
        grupos.append((numeros_ordenados, vencedor))
    return grupos


# ----------------------------------------------------------------------
# GERAÇÃO DO PDF — MAPA COMPARATIVO / SOLICITAÇÃO DE ADJUDICAÇÃO
# ----------------------------------------------------------------------
def gerar_pdf_mapa_comparativo(caminho_pdf, codigo_cotacao, ano, mapa_itens, objeto_titulo="",
                                observacoes="", diretor_nome="", diretor_matricula="",
                                sei_publicacao="", tipo_planilha="HRG",
                                publicacao_rotulo="1ª PUBLICAÇÃO", texto_pesquisa=TEXTO_PESQUISA_PADRAO,
                                titulos_pesquisa=TITULOS_PESQUISA_PADRAO):
    """Gera o Mapa Comparativo / Solicitação de Adjudicação replicando o
    layout usado pela SES/PDPAS (cabeçalho, barra de título, tabela de
    1º/2º/3º colocado por item com saldo remanescente e descrição, tabela de
    empresas vencedoras e o texto de adjudicação/homologação). Formato
    retrato, igual ao modelo original (uma única página).

    As cores das barras de título seguem o tipo de planilha (extraídas dos
    modelos reais): APS = verde claro; APS INV = verde claro + amarelo;
    HRG = azul claro; HRG INV = azul claro + amarelo. Cada item pode ter um
    'vencedor_idx' (0=1º colocado, 1=2º, 2=3º) definido manualmente — o
    colocado escolhido recebe um destaque cinza na tabela, e é ele (não
    necessariamente o de menor preço) que aparece na tabela de empresas
    vencedoras."""
    is_hrg = "HRG" in tipo_planilha.upper()
    is_inv = "INV" in tipo_planilha.upper()
    cor_base = COR_CIANO_ADJ if is_hrg else COR_VERDE_ADJ
    sufixo_letterhead = "/HRG" if is_hrg else ""
    sigla_comissao = "HRG" if is_hrg else "SRSSU"

    doc = SimpleDocTemplate(caminho_pdf, pagesize=A4,
                             topMargin=1 * cm, bottomMargin=1 * cm,
                             leftMargin=1.2 * cm, rightMargin=1.2 * cm)
    styles = getSampleStyleSheet()
    largura_util = A4[0] - 2.4 * cm

    estilo_barra = ParagraphStyle("Barra", parent=styles["Normal"], fontSize=12,
                                  alignment=TA_CENTER, fontName="Helvetica-Bold")
    estilo_objeto = ParagraphStyle("Objeto", parent=styles["Normal"], fontSize=10.5,
                                   alignment=TA_CENTER, fontName="Helvetica-Bold", leading=13)
    estilo_mapa_titulo = ParagraphStyle("MapaTitulo", parent=styles["Normal"], fontSize=14,
                                        alignment=TA_CENTER, fontName="Helvetica-Bold")
    estilo_cel_item = ParagraphStyle("CelItem", parent=styles["Normal"], fontSize=7,
                                     alignment=TA_CENTER, fontName="Helvetica-Bold", leading=10)
    estilo_cel_normal = ParagraphStyle("CelNormal", parent=styles["Normal"], fontSize=6,
                                       alignment=TA_CENTER, leading=8)
    estilo_cel_bold = ParagraphStyle("CelBold", parent=styles["Normal"], fontSize=6,
                                     alignment=TA_CENTER, fontName="Helvetica-Bold", leading=8)
    estilo_desc = ParagraphStyle("Desc", parent=styles["Normal"], fontSize=7.5, leading=9,
                                 fontName="Helvetica-Bold", alignment=TA_LEFT)
    estilo_corpo = ParagraphStyle("Corpo", parent=styles["Normal"], fontSize=9, leading=12,
                                  alignment=TA_JUSTIFY, firstLineIndent=26)
    estilo_corpo_centro = ParagraphStyle("CorpoCentro", parent=estilo_corpo, alignment=TA_CENTER,
                                         firstLineIndent=0)
    estilo_sei = ParagraphStyle("Sei", parent=styles["Normal"], fontSize=9, leading=12,
                                textColor=colors.red, fontName="Helvetica-Bold", alignment=TA_CENTER)
    estilo_rodape = ParagraphStyle("Rodape", parent=styles["Normal"], fontSize=7.5, leading=9.5,
                                   alignment=TA_CENTER, fontName="Helvetica-Bold")

    story = [
        montar_cabecalho_institucional(largura_util, [
            "GOVERNO DO DISTRITO FEDERAL",
            "SECRETARIA DE ESTADO DE SAÚDE",
            "SUPERINTENDÊNCIA DA REGIÃO DE SAÚDE SUL",
            "DIRETORIA ADMINISTRATIVA",
            f"PROGRAMA DE DESCENTRALIZAÇÃO PROGRESSIVA DE AÇÕES DE SAÚDE - PDPAS{sufixo_letterhead}",
        ]),
        Spacer(1, 8),
    ]

    barra1 = Table([[Paragraph("SOLICITAÇÃO DE ADJUDICAÇÃO PARA O PEDIDO DE COTAÇÃO PDPAS",
                               estilo_barra)]], colWidths=[largura_util])
    barra1.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), cor_base),
                                ("BOX", (0, 0), (-1, -1), 0.5, colors.black),
                                ("TOPPADDING", (0, 0), (-1, -1), 3),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    story.append(barra1)

    # A segunda linha (cotação/ano - publicação) só ganha fundo amarelo nas
    # variantes "...INV"; nas variantes normais (APS/HRG) fica sem destaque,
    # igual aos modelos reais.
    texto_barra2 = f"{codigo_cotacao}/{ano} – {publicacao_rotulo}"
    if is_inv:
        barra2 = Table([[Paragraph(texto_barra2, estilo_barra)]], colWidths=[largura_util])
        barra2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), COR_AMARELO_ADJ),
                                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        story.append(barra2)
    else:
        story.append(Paragraph(texto_barra2, estilo_barra))
    story.append(Spacer(1, 8))

    if objeto_titulo:
        story.append(Paragraph(objeto_titulo, estilo_objeto))
        story.append(Spacer(1, 6))

    story.append(Paragraph(f"MAPA COMPARATIVO {codigo_cotacao}/{ano}", estilo_mapa_titulo))
    story.append(Spacer(1, 8))

    frac = [0.05, 0.085, 0.065, 0.05, 0.085, 0.065, 0.05, 0.085, 0.065, 0.05, 0.11, 0.11, 0.13]
    col_widths = [largura_util * f for f in frac]

    linhas = [[
        Paragraph("Item /<br/>Código", estilo_cel_bold),
        Paragraph("1º Colocado /<br/>Unitário / Total", estilo_cel_bold), "",
        Paragraph("Qtd./Und./<br/>Marca", estilo_cel_bold),
        Paragraph("2º Colocado /<br/>Unitário / Total", estilo_cel_bold), "",
        Paragraph("Qtd./Und./<br/>Marca", estilo_cel_bold),
        Paragraph("3º Colocado /<br/>Unitário / Total", estilo_cel_bold), "",
        Paragraph("Qtd./Und./<br/>Marca", estilo_cel_bold),
    ]]
    linhas[0] += [Paragraph(escape(titulo), estilo_cel_bold) for titulo in titulos_pesquisa]
    linhas[0].append(Paragraph("MÉDIA", estilo_cel_bold))
    estilos_extra = [
        ("SPAN", (1, 0), (2, 0)), ("SPAN", (4, 0), (5, 0)), ("SPAN", (7, 0), (8, 0)),
        ("BACKGROUND", (0, 0), (-1, 0), COR_CINZA_ADJ),
    ]

    COR_VENCEDOR_ADJ = colors.HexColor("#D9D9D9")   # cinza claro — destaque do vencedor
    CORES_ALTERNADAS = [colors.HexColor("#D9D9D9"), colors.HexColor("#A5A5A5")]  # cinza claro/escuro

    linha_idx = 1
    for indice_item, item in enumerate(mapa_itens):
        cancelado = bool(item.get("cancelado"))
        propostas = [] if cancelado or texto_status_item(item) else item["propostas"]
        top3 = propostas[:3]
        cor_bloco = CORES_ALTERNADAS[indice_item % 2]
        linha_topo_bloco = linha_idx

        # linha grossa separando este bloco do item anterior (exceto no 1º)
        if indice_item > 0:
            estilos_extra.append(("LINEABOVE", (0, linha_topo_bloco), (-1, linha_topo_bloco),
                                  1.6, colors.black))

        # ---- linha 1: nome do colocado / quantidade ----
        linha1 = [Paragraph(str(item["numero_item"] or ""), estilo_cel_item)]
        for pos in range(3):
            if pos < len(top3):
                p = top3[pos]
                linha1 += [Paragraph(p["fornecedor"], estilo_cel_bold), "",
                           Paragraph(str(p["quantidade"] or ""), estilo_cel_normal)]
            else:
                if cancelado:
                    texto_vazio = "CANCELADO"
                else:
                    texto_vazio = texto_status_item(item) or ("DESERTO" if not item["propostas"] else "PROPOSTA NÃO APRESENTADA")
                linha1 += [Paragraph(texto_vazio, estilo_cel_bold), "", ""]
        pesquisas = valores_pesquisa(item)
        ref = item.get("referencias", {})
        linha1 += [Paragraph(formatar_pesquisa(v), estilo_cel_normal) for v in pesquisas]
        linha1.append(Paragraph(formatar_pesquisa(ref.get("media")), estilo_cel_normal))
        estilos_extra += [("SPAN", (10, linha_idx), (10, linha_idx+2)),
                          ("SPAN", (11, linha_idx), (11, linha_idx+2))]
        linhas.append(linha1)
        estilos_extra += [
            ("SPAN", (1, linha_idx), (2, linha_idx)), ("SPAN", (4, linha_idx), (5, linha_idx)),
            ("SPAN", (7, linha_idx), (8, linha_idx)),
            ("BACKGROUND", (0, linha_idx), (0, linha_idx), cor_bloco),
        ]
        linha_idx += 1

        # ---- linha 2: unitário / total / UN ----
        linha2 = [""]
        for pos in range(3):
            if pos < len(top3):
                p = top3[pos]
                linha2 += [Paragraph(formatar_moeda(p["unitario"]), estilo_cel_normal),
                           Paragraph(formatar_moeda(p["total"]), estilo_cel_normal),
                           Paragraph(item["un"], estilo_cel_normal)]
            else:
                linha2 += ["", "", ""]
        linha2 += ["", "", Paragraph("MEDIANA", estilo_cel_bold)]
        linhas.append(linha2)
        estilos_extra.append(("BACKGROUND", (0, linha_idx), (0, linha_idx), cor_bloco))
        linha_idx += 1

        # ---- linha 3: código / prazo / marca ----
        linha3 = [Paragraph(item["codigo"], estilo_cel_bold)]
        for pos in range(3):
            linha3.append(Paragraph("Prazo (dias)", estilo_cel_bold))
            if pos < len(top3):
                p = top3[pos]
                linha3 += [Paragraph(str(p["prazo"] or ""), estilo_cel_normal),
                           Paragraph(p["marca"] or "", estilo_cel_normal)]
            else:
                linha3 += ["", ""]
        linha3 += ["", "", Paragraph(formatar_pesquisa(ref.get("mediana")), estilo_cel_normal)]
        linhas.append(linha3)
        estilos_extra += [
            ("BACKGROUND", (0, linha_idx), (0, linha_idx), cor_bloco),
            ("BACKGROUND", (1, linha_idx), (1, linha_idx), cor_bloco),
            ("BACKGROUND", (4, linha_idx), (4, linha_idx), cor_bloco),
            ("BACKGROUND", (7, linha_idx), (7, linha_idx), cor_bloco),
        ]
        linha_idx += 1

        # ---- linha 4: saldo remanescente / descrição (1 linha só) ----
        saldo_texto = formatar_moeda(item["saldo"]) if item["saldo"] is not None else "-"
        descricao_1_linha = truncar_em_palavra((item["descricao"] or "").strip(), 60)
        if descricao_1_linha and len(descricao_1_linha) < len((item["descricao"] or "").strip()):
            descricao_1_linha += "..."
        linha4 = [Paragraph("Saldo remanescente", estilo_cel_bold), "", "",
                  Paragraph(saldo_texto, estilo_cel_bold),
                  Paragraph(descricao_1_linha, estilo_desc), "", "", "", "", ""]
        linha4 += [Paragraph("Preço de referência", estilo_cel_bold), "",
                   Paragraph(formatar_pesquisa(ref.get("minimo")), estilo_cel_bold)]
        estilos_extra += [("SPAN", (10, linha_idx), (11, linha_idx)),
                          ("BACKGROUND", (10, linha_idx), (12, linha_idx), cor_base)]
        linhas.append(linha4)
        estilos_extra += [
            ("SPAN", (0, linha_idx), (2, linha_idx)), ("SPAN", (4, linha_idx), (9, linha_idx)),
            ("BACKGROUND", (0, linha_idx), (2, linha_idx), cor_bloco),
        ]

        # ---- destaque do vencedor (não é sempre o 1º colocado — pode ter
        # sido alterado manualmente, ex.: desistência do menor preço) ----
        vencedor_pos = item.get("vencedor_idx")
        if vencedor_pos is not None and vencedor_pos < len(top3):
            col_inicio = 1 + 3 * vencedor_pos
            estilos_extra.append(
                ("BACKGROUND", (col_inicio, linha_topo_bloco), (col_inicio + 2, linha_idx - 1),
                 COR_VENCEDOR_ADJ))

        linha_idx += 1

    tabela = Table(linhas, colWidths=col_widths)
    tabela.setStyle(TableStyle(estilos_extra + [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(tabela)
    story.append(Spacer(1, 10))

    story.append(Paragraph(escape(texto_pesquisa).replace("\n", "<br/>"), estilo_corpo))
    story.append(Spacer(1, 6))

    grupos_vencedores = montar_declaracao_vencedores(mapa_itens)
    if grupos_vencedores:
        dados_decl = [["Item(ns)", "Empresa"]]
        for numeros, vencedor in grupos_vencedores:
            dados_decl.append([", ".join(str(n) for n in numeros), vencedor])
        tabela_decl = Table(dados_decl, colWidths=[largura_util * 0.15, largura_util * 0.85])
        tabela_decl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), COR_CINZA_ADJ),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(tabela_decl)
        story.append(Spacer(1, 10))

    if observacoes:
        for paragrafo in observacoes.split("\n\n"):
            paragrafo = paragrafo.strip()
            if paragrafo:
                story.append(Paragraph(paragrafo, estilo_corpo))
                story.append(Spacer(1, 4))
        story.append(Spacer(1, 2))

    story.append(Paragraph(
        "Aceita essa solicitação, procederemos à etapa do envio de Ordem de Compra do material "
        "cotado.", estilo_corpo))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Atenciosamente,", estilo_corpo_centro))
    story.append(Paragraph(f"COMISSÃO DO PDPAS/{sigla_comissao}.", estilo_corpo_centro))
    story.append(Spacer(1, 12))

    story.append(Paragraph("ADJUDICAÇÃO E HOMOLOGAÇÃO", estilo_mapa_titulo))
    story.append(Spacer(1, 4))
    story.append(Paragraph("De acordo, em função do exposto acima.", estilo_corpo_centro))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        "<b>Adjudico e Homologo</b> o presente. Proceda-se à desaplicação dos recursos necessários "
        "para fazer face ao pagamento da presente aquisição, assim que os materiais forem entregues "
        "e a NOTA FISCAL for atestada.", estilo_corpo))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        "Em tempo, solicitamos envio ao <b>NPDOC</b> para acompanhamento do prazo de entrega.",
        estilo_corpo))
    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "Atendendo ao art. 16. Inciso IV da Portaria SES 473, de 04/12/2023 Declaro que há "
        "<b>disponibilidade financeira</b> para cobrir a presente contratação.", estilo_corpo_centro))
    if diretor_nome:
        story.append(Paragraph(
            f"Diretor(a) Administrativo(a)/SRSSU: {diretor_nome}, Matrícula: {diretor_matricula}",
            estilo_corpo_centro))
    story.append(Spacer(1, 10))
    if sei_publicacao:
        story.append(Paragraph(
            f"A Publicação no Diário Oficial do DF será realizada pelo SEI {sei_publicacao}",
            estilo_sei))

    story.append(Spacer(1, 14))
    if is_hrg:
        linha_rodape2 = "PDPAS/HRG - GEOF / Diretoria Administrativa / SRSSU / SES"
        linha_rodape3 = "Área Especial Nº 01, Lote Único - Setor Central GAMA/DF. CEP: 72.405-901"
        linha_rodape4 = "CNPJ: 00.394.700/0006-12 - Fone: (61) 3449-7105"
    else:
        linha_rodape2 = "PDPAS / APS-SUL - GEOF / Diretoria Administrativa / SRSSU / SES"
        linha_rodape3 = ""
        linha_rodape4 = "CNPJ: 00.394.700/0051-77 - Fone: (61) 3449-7105"
    linhas_rodape = ["Brasília – Patrimônio Cultural da Humanidade", linha_rodape2]
    if linha_rodape3:
        linhas_rodape.append(linha_rodape3)
    linhas_rodape.append(linha_rodape4)
    story.append(Paragraph("<br/>".join(linhas_rodape), estilo_rodape))

    if len(mapa_itens) <= 12:
        doc.build([KeepInFrame(doc.width, doc.height - 12, story, mode="shrink")])
    else:
        doc.build(story)


# ----------------------------------------------------------------------
# GERAÇÃO DO PDF — ORDEM DE COMPRA
# ----------------------------------------------------------------------
def gerar_pdf_ordem_compra(caminho_pdf, codigo_cotacao, ano, tr_numero, origem, fornecedor_info,
                            itens_vencidos, tipo_planilha, sei_texto, diretor_nome, diretor_matricula,
                            contatos=None):
    """Gera a ORDEM DE COMPRA de um fornecedor vencedor, com os itens que
    ele ganhou. Usa o mesmo esquema de cores da Adjudicação (APS=verde,
    HRG=azul/ciano, "...INV"=+amarelo), mas no Ordem a cor amarela fica na
    PRIMEIRA barra (ao contrário da Adjudicação), conforme os modelos reais.

    'contatos' é o dict configurável (ver CONTATOS_ORDEM_PADRAO) com os
    nomes/telefones do GERENTE, NFL, NUAL, NECFM, NPDOC e NFH — o bloco de
    "Informações técnicas/Aceite NF" muda de acordo com o tipo_planilha
    exato (APS / APS INV / HRG / HRG INV), cada um com seu próprio
    conjunto de contatos."""
    contatos = {**CONTATOS_ORDEM_PADRAO, **(contatos or {})}
    tipo_norm = tipo_planilha.strip().upper()
    is_hrg = "HRG" in tipo_norm
    is_inv = "INV" in tipo_norm
    cor_base = COR_CIANO_ADJ if is_hrg else COR_VERDE_ADJ
    cor_barra1 = COR_AMARELO_ADJ if is_inv else cor_base
    cor_barra2 = cor_base

    doc = SimpleDocTemplate(caminho_pdf, pagesize=A4,
                             topMargin=1 * cm, bottomMargin=1 * cm,
                             leftMargin=1.2 * cm, rightMargin=1.2 * cm)
    styles = getSampleStyleSheet()
    largura_util = A4[0] - 2.4 * cm

    estilo_barra = ParagraphStyle("BarraOrd", parent=styles["Normal"], fontSize=12,
                                  alignment=TA_CENTER, fontName="Helvetica-Bold")
    estilo_valor = ParagraphStyle("ValorOrd", parent=styles["Normal"], fontSize=9, leading=12)
    estilo_cel = ParagraphStyle("CelOrd", parent=styles["Normal"], fontSize=8.5, leading=10.5)
    estilo_cel_center = ParagraphStyle("CelCenterOrd", parent=estilo_cel, alignment=TA_CENTER)
    estilo_corpo = ParagraphStyle("CorpoOrd", parent=styles["Normal"], fontSize=8.5, leading=11)
    estilo_corpo_centro = ParagraphStyle("CorpoCentroOrd", parent=estilo_corpo, alignment=TA_CENTER)
    estilo_corpo_centro_bold = ParagraphStyle("CorpoCentroBoldOrd", parent=estilo_corpo_centro,
                                              fontName="Helvetica-Bold")
    estilo_rodape = ParagraphStyle("RodapeOrd", parent=styles["Normal"], fontSize=7.5, leading=9.5,
                                   alignment=TA_CENTER, fontName="Helvetica-Bold")

    story = [
        montar_cabecalho_institucional(largura_util, [
            "GOVERNO DO DISTRITO FEDERAL",
            "SECRETARIA DE ESTADO DE SAÚDE",
            "SUPERINTENDÊNCIA DA REGIÃO DE SAÚDE SUL",
            "DIRETORIA ADMINISTRATIVA",
            "GERÊNCIA DE ORÇAMENTO E FINANÇAS/SRSSU - PDPAS",
        ]),
        Spacer(1, 8),
    ]

    barra1 = Table([[Paragraph(f"ORDEM DE COMPRA - COTAÇÃO {codigo_cotacao}/{ano}", estilo_barra)]],
                    colWidths=[largura_util])
    barra1.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), cor_barra1),
                                ("BOX", (0, 0), (-1, -1), 0.5, colors.black),
                                ("TOPPADDING", (0, 0), (-1, -1), 3),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    story.append(barra1)

    # O TR já vem como "033/2026" (com o ano embutido) direto da aba Controle,
    # então só concatena "/{ano}" se ele ainda não tiver uma barra.
    if tr_numero and "/" in str(tr_numero):
        texto_barra2 = f"TR Nº {tr_numero} - {origem}"
    elif tr_numero:
        texto_barra2 = f"TR Nº {tr_numero}/{ano} - {origem}"
    else:
        texto_barra2 = origem or ""
    if is_inv:
        barra2 = Table([[Paragraph(texto_barra2, estilo_barra)]], colWidths=[largura_util])
        barra2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), cor_barra2),
                                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        story.append(barra2)
    else:
        story.append(Paragraph(texto_barra2, estilo_barra))
    story.append(Spacer(1, 10))

    # ---- bloco com dados do fornecedor ----
    linhas_forn = [
        [Paragraph(f"<b>FORNECEDOR:</b> {fornecedor_info.get('razao_social') or fornecedor_info.get('nome', '')}",
                   estilo_valor),
         Paragraph(f"<b>CÓDIGO SES:</b> {fornecedor_info.get('codigo_ses', '')}", estilo_valor)],
        [Paragraph(f"<b>CONTATO:</b> {fornecedor_info.get('contato', '')}", estilo_valor),
         Paragraph(f"<b>SEI:</b> {sei_texto or ''}", estilo_valor)],
        [Paragraph(f"<b>FONE:</b> {fornecedor_info.get('telefone', '')}", estilo_valor), ""],
        [Paragraph(f"<b>CNPJ:</b> {fornecedor_info.get('cnpj', '')}", estilo_valor), ""],
        [Paragraph(f"<b>BANCO:</b> {fornecedor_info.get('banco', '')}", estilo_valor),
         Paragraph(f"<b>AGÊNCIA:</b> {fornecedor_info.get('agencia', '')} &nbsp;&nbsp; "
                   f"<b>CONTA:</b> {fornecedor_info.get('conta', '')}", estilo_valor)],
    ]
    tabela_forn = Table(linhas_forn, colWidths=[largura_util * 0.62, largura_util * 0.38])
    tabela_forn.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(tabela_forn)
    story.append(Spacer(1, 8))

    story.append(Paragraph(
        "Solicitamos os seguintes materiais a serem entregues conforme cotação. A compra só será "
        "reconhecida com o atesto da Nota Fiscal, que deverá acompanhar os materiais abaixo "
        "solicitados:", estilo_corpo))
    story.append(Spacer(1, 6))

    # ---- tabela de itens vencidos por este fornecedor ----
    dados_tabela = [["Item", "Cod.", "Descrição", "Qtd.", "Und.", "Valor Unitário",
                     "Valor Total", "Prazo de Entrega"]]
    total_geral = 0
    for item, proposta in itens_vencidos:
        dados_tabela.append([
            str(item["numero_item"] or ""), item["codigo"],
            Paragraph(item["descricao"] or "", estilo_cel),
            str(proposta["quantidade"] or ""), item["un"],
            formatar_moeda(proposta["unitario"]), formatar_moeda(proposta["total"]),
            f'{proposta["prazo"]} dias úteis.' if proposta["prazo"] else "",
        ])
        total_geral += proposta["total"] or 0

    tabela_itens = Table(
        dados_tabela,
        colWidths=[1.2 * cm, 1.8 * cm, largura_util - 12.4 * cm, 1.3 * cm, 1.3 * cm,
                   2.4 * cm, 2.4 * cm, 2.4 * cm],
        repeatRows=1)
    tabela_itens.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (1, -1), "CENTER"), ("ALIGN", (3, 0), (-1, -1), "CENTER"),
    ]))
    story.append(tabela_itens)

    tabela_total = Table([["TOTAL", "", formatar_moeda(total_geral)]],
                         colWidths=[3 * cm, largura_util - 3 * cm - 3.5 * cm, 3.5 * cm])
    tabela_total.setStyle(TableStyle([
        ("SPAN", (0, 0), (1, 0)),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#D9D9D9")),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (0, 0), "CENTER"), ("ALIGN", (2, 0), (2, 0), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
    ]))
    story.append(tabela_total)
    story.append(Spacer(1, 10))

    # ---- dados para nota fiscal / entrega ----
    barra_nf = Table([[Paragraph("Dados para emissão de Nota Fiscal e local para entrega:",
                                 estilo_barra)]], colWidths=[largura_util])
    barra_nf.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), cor_barra1),
                                  ("TOPPADDING", (0, 0), (-1, -1), 3),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                                  ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
    story.append(barra_nf)
    story.append(Spacer(1, 4))

    if is_hrg:
        nome_unidade = "HOSPITAL REGIONAL DO GAMA - DF"
        cnpj_unidade = "00.394.700/0006-12"
        insc_municipal = "0732950500936"
        endereco_unidade = "ÁREA ESPECIAL Nº 01 - Setor Central GAMA/DF. CEP: 72.405-901"
    else:
        nome_unidade = "SUPERINTENDENCIA DA REGIAO DE SAUDE SUL - APS"
        cnpj_unidade = "00.394.700/0051-77"
        insc_municipal = "0732950502475"
        endereco_unidade = ("GAOAPS-SUL no Setor Oeste, SOE Q. 12, Praça 1, Área Especial - "
                            "GAMA OESTE, CEP: 72.425-120.")

    story.append(Paragraph(f"<b>{nome_unidade}</b>", estilo_corpo))
    story.append(Paragraph(f"<b>CNPJ:</b> {cnpj_unidade} &nbsp;&nbsp; "
                           f"<b>Inscrição Municipal:</b> {insc_municipal}", estilo_corpo))
    story.append(Paragraph(f"<b>Endereço:</b> {endereco_unidade}", estilo_corpo))
    story.append(Spacer(1, 6))

    if not is_hrg:
        story.append(Paragraph("***LOCAL DE ENTREGA***", estilo_corpo_centro_bold))
        story.append(Paragraph(endereco_unidade, estilo_corpo_centro))
        story.append(Paragraph("Ao lado da Igreja Nossa Senhora Aparecida. Coordenadas no Maps "
                               "-16.007452594187583, -48.075027885099104", estilo_corpo_centro))
        story.append(Paragraph("Plus Code do Google Maps XWVF+2X9 Gama, Brasília - DF",
                               estilo_corpo_centro))
        story.append(Spacer(1, 6))

    story.append(Paragraph(
        "Os pagamentos da Região Sul são efetuados por TED, entre os dias 20 e 30 de cada mês. "
        "NÃO EMITIR BOLETO!", estilo_corpo_centro_bold))
    story.append(Spacer(1, 4))

    if tipo_norm == "APS":
        story.append(Paragraph(
            f"<b>Informações técnicas/Aceite NF:</b> {contatos.get('GERENTE', '')} - GERENTE;",
            estilo_corpo))
        story.append(Paragraph(f"Farmácia-NFL - {contatos.get('NFL', '')}", estilo_corpo))
        story.append(Paragraph(f"Almoxarifado-NUAL - {contatos.get('NUAL_APS', '')}", estilo_corpo))
        story.append(Paragraph(
            f"ENGENHARIA CLINICA/NECFM - {contatos.get('NECFM', '')}", estilo_corpo))
    elif tipo_norm == "APS INV":
        story.append(Paragraph(
            f"<b>Informações técnicas/Aceite NF:</b> {contatos.get('GERENTE', '')} - GERENTE;",
            estilo_corpo))
        story.append(Paragraph("Entrega no Núcleo de Patrimônio (NPDOC/APS) de segunda a QUINTA de "
                               "08:00 as 11:00 e 14:00 as 17:00", estilo_corpo))
        story.append(Paragraph(f"Chefe do NPDOC - {contatos.get('NPDOC_APSINV', '')}", estilo_corpo))
    elif tipo_norm == "HRG":
        linhas_aceite = [
            [Paragraph("<b>Informações técnicas/Aceite NF:</b>", estilo_corpo)],
            [Paragraph(f"<b>Farmácia NFH</b> - {escape(contatos['NFH'])}", estilo_corpo)],
            [Paragraph(f"<b>Almoxarifado NUAL</b> - {escape(contatos['NUAL_HRG'])}", estilo_corpo)],
            [Paragraph(f"<b>Engenharia Clínica NECFM</b> - {escape(contatos['NECFM_HRG'])}", estilo_corpo)],
        ]
        tabela_aceite = Table(linhas_aceite, colWidths=[largura_util])
        tabela_aceite.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(tabela_aceite)
    else:  # HRG INV — atenção: aqui NÃO entra NECFM, só o Chefe do NPDOC
        story.append(Paragraph("<b>Informações técnicas/Aceite NF:</b>", estilo_corpo))
        story.append(Paragraph("Entrega no Núcleo de Patrimônio (NPDOC/HRG) de segunda a QUINTA de "
                               "08:00 as 11:00 e 14:00 as 17:00", estilo_corpo))
        story.append(Paragraph(f"Chefe do NPDOC - {contatos.get('NPDOC_HRGINV', '')}", estilo_corpo))
    story.append(Spacer(1, 8))

    story.append(Paragraph(
        "*Favor discriminar na Nota o valor UNITÁRIO e a QUANTIDADE de cada item, conforme esta "
        "Ordem de Compra.", estilo_corpo))
    story.append(Paragraph(
        "*Nota Fiscal empresa Não enquadrada Simples Nacional terá Imposto de Renda retido no "
        "percentual: 1,2% p/material; e 4,8% p/serviço, IN RFB nº 2145 26/06/2023", estilo_corpo))
    story.append(Paragraph(
        "*PAGAMENTO de empresas do DF, acima de R$ 5.000,00 somente em conta do BRB. Decreto "
        "Distrital Nº 32.767, de 17/02/2011 artigo 6º caput.", estilo_corpo))
    story.append(Paragraph(
        "*Caso seja necessário a instalação e/ou montagem deve ser agendada com o Chefe do "
        "Núcleo Citado acima", estilo_corpo))
    story.append(Spacer(1, 14))

    if diretor_nome:
        story.append(Paragraph(diretor_nome, estilo_corpo_centro_bold))
        story.append(Paragraph("Diretor(a) Administrativo(a)/SRSSU", estilo_corpo_centro))
        story.append(Paragraph(f"Matrícula: {diretor_matricula}", estilo_corpo_centro))

    story.append(Spacer(1, 12))
    story.append(Paragraph(
        "Brasília-Patrimônio Cultural da Humanidade<br/>"
        "GEOF / Diretoria Administrativa / SRSSU / SES<br/>"
        "AREA ESPECIAL Nº 01 - Setor Central GAMA/DF. CEP: 72.405-901<br/>"
        "Contato PREFERENCIALMENTE por e-mails, FoneVOIP: (61)3449-7105", estilo_rodape))

    doc.build(story)


# ----------------------------------------------------------------------
# MIXIN — métodos de ação ligados aos botões do 3º Passo
# ----------------------------------------------------------------------
class AdjudicacaoOrdemMixin:
    """Métodos herdados por AppCotacao. Usa o cadastro de propostas
    (self.mapa_fornecedores / arquivo MAPA_EMPRESAS_*.json) preparado no
    2º Passo (Mapa.py) para montar o Mapa Comparativo e as Ordens de Compra."""

    def acao_adjudicacao(self):
        if not self.linhas_resultado:
            messagebox.showwarning(
                "Atenção",
                "Nenhum item de cotação carregado. Clique em CONSULTAR COTAÇÃO antes.")
            return

        contexto = self._preparar_contexto_wizard()
        if contexto is None:
            return
        pasta_propostas, codigo_cotacao, ano = contexto

        caminho_json = self._caminho_arquivo_mapa_empresas(self._pasta_cotacao_atual, codigo_cotacao)
        if not os.path.exists(caminho_json):
            messagebox.showwarning(
                "Nenhum cadastro encontrado",
                f"Não há propostas cadastradas em:\n{caminho_json}\n\n"
                "Use o botão INCLUIR EMPRESAS MAPA primeiro.")
            return
        try:
            with open(caminho_json, "r", encoding="utf-8") as f:
                dados_salvos = json.load(f)
        except Exception as e:
            messagebox.showerror("Erro ao carregar cadastro", str(e))
            return

        self.mapa_fornecedores = dados_salvos.get("fornecedores", {})
        if not self.mapa_fornecedores:
            messagebox.showwarning("Atenção", "O arquivo salvo não contém fornecedores.")
            return

        self._gerar_mapa_comparativo()

    def _montar_dados_ordens_ativas(self):
        """Monta os dados necessários para gerar as Ordens de Compra da
        cotação ativa — em PDF (acao_ordem, aqui) ou por e-mail
        (acao_email_ordem, em Emails.py): consulta o cadastro de propostas
        já salvo, aplica as decisões de vencedor/cancelamento definidas na
        Adjudicação e agrupa os itens vencidos por fornecedor.

        Devolve um dict com 'ordens_por_fornecedor' (nome_normalizado ->
        (texto_completo, [(item, proposta), ...])), 'fornecedores_completo',
        'codigo_cotacao', 'ano', 'tr_numero', 'origem' e 'sei_texto' — ou
        None se algo impedir a montagem (já mostrando o aviso/erro
        correspondente nesse caso)."""
        if not self.linhas_resultado:
            messagebox.showwarning(
                "Atenção", "Nenhum item de cotação carregado. Clique em CONSULTAR COTAÇÃO antes.")
            return None

        contexto = self._preparar_contexto_wizard()
        if contexto is None:
            return None
        pasta_propostas, codigo_cotacao, ano = contexto

        caminho_json = self._caminho_arquivo_mapa_empresas(self._pasta_cotacao_atual, codigo_cotacao)
        if not os.path.exists(caminho_json):
            messagebox.showwarning(
                "Nenhum cadastro encontrado",
                f"Não há propostas cadastradas em:\n{caminho_json}\n\n"
                "Use INCLUIR EMPRESAS MAPA e depois ADJUDICAÇÃO primeiro.")
            return None
        try:
            with open(caminho_json, "r", encoding="utf-8") as f:
                dados_salvos = json.load(f)
        except Exception as e:
            messagebox.showerror("Erro ao carregar cadastro", str(e))
            return None

        mapa_fornecedores = dados_salvos.get("fornecedores", {})
        decisoes = dados_salvos.get("decisoes_itens", {})
        if not mapa_fornecedores:
            messagebox.showwarning("Atenção", "O arquivo salvo não contém fornecedores.")
            return None
        if not decisoes:
            messagebox.showwarning(
                "Atenção",
                "Ainda não há vencedores definidos para esta cotação.\n"
                "Clique em ADJUDICAÇÃO primeiro para escolher os vencedores de cada item.")
            return None

        try:
            fornecedores_completo = carregar_fornecedores_completo(self._obter_caminho_arquivo())
        except Exception as e:
            messagebox.showerror("Erro ao ler aba Fornecedores", str(e))
            return None
        fornecedores_texto = {chave: dados["texto_completo"]
                              for chave, dados in fornecedores_completo.items()}

        mapa_itens = montar_mapa_itens_completo(
            self.linhas_resultado, mapa_fornecedores, fornecedores_texto, self.saldos_bd_atual)

        # aplica as decisões (vencedor escolhido / cancelado) salvas na Adjudicação
        for item in mapa_itens:
            decisao = decisoes.get(str(item["numero_item"]))
            if decisao:
                aplicar_retiradas(item, decisao.get("fornecedores_retirados", []))
                item["vencedor_idx"] = decisao.get("vencedor_idx")
                item["cancelado"] = bool(decisao.get("cancelado"))
                item["status_item"] = decisao.get("status_item", "")

        # agrupa os itens vencidos por fornecedor (só o vencedor de cada item)
        ordens_por_fornecedor = {}   # nome_normalizado -> (texto_completo, [ (item, proposta), ... ])
        for item in mapa_itens:
            if item.get("cancelado") or texto_status_item(item):
                continue
            propostas = item["propostas"]
            if not propostas:
                continue
            idx = min(item.get("vencedor_idx", 0) or 0, len(propostas) - 1)
            proposta_vencedora = propostas[idx]
            texto_vencedor = proposta_vencedora["fornecedor"]
            nome_vencedor, _ = extrair_nome_codigo_fornecedor(texto_vencedor)
            chave = normalizar_texto(nome_vencedor)
            if chave not in ordens_por_fornecedor:
                ordens_por_fornecedor[chave] = (texto_vencedor, [])
            ordens_por_fornecedor[chave][1].append((item, proposta_vencedora))

        if not ordens_por_fornecedor:
            messagebox.showinfo(
                "Nenhum vencedor", "Nenhum item tem vencedor definido (todos cancelados ou sem "
                                   "propostas).")
            return None

        tr_numero = self.linhas_resultado[0].get("TR") if self.linhas_resultado else ""
        origem = self.linhas_resultado[0].get("ORIGEM") if self.linhas_resultado else ""
        sei_texto = self.entry_sei.get().strip()

        return {
            "ordens_por_fornecedor": ordens_por_fornecedor,
            "fornecedores_completo": fornecedores_completo,
            "codigo_cotacao": codigo_cotacao,
            "ano": ano,
            "tr_numero": tr_numero,
            "origem": origem,
            "sei_texto": sei_texto,
        }

    def acao_ordem(self):
        dados = self._montar_dados_ordens_ativas()
        if dados is None:
            return
        ordens_por_fornecedor = dados["ordens_por_fornecedor"]
        fornecedores_completo = dados["fornecedores_completo"]
        codigo_cotacao = dados["codigo_cotacao"]
        ano = dados["ano"]
        tr_numero = dados["tr_numero"]
        origem = dados["origem"]
        sei_texto = dados["sei_texto"]
        diretor_nome = self.entry_da_nome.get().strip()
        diretor_matricula = self.entry_da_matricula.get().strip()

        arquivos_gerados = []
        for chave, (texto_vencedor, itens_vencidos) in ordens_por_fornecedor.items():
            info_fornecedor = fornecedores_completo.get(chave, {"nome": texto_vencedor})
            nome_arquivo_seguro = re_sub_nome_arquivo(info_fornecedor.get("nome", chave))
            base = f"ORDEM {codigo_cotacao} {nome_arquivo_seguro}"
            caminho_pdf = os.path.join(self._pasta_cotacao_atual, base + ".pdf")
            caminho_jpg_base = os.path.join(self._pasta_cotacao_atual, base)
            try:
                gerar_pdf_ordem_compra(
                    caminho_pdf, codigo_cotacao, ano, tr_numero, origem, info_fornecedor,
                    itens_vencidos, self._sufixo_planilha_atual, sei_texto,
                    diretor_nome, diretor_matricula,
                    contatos=self.config_dados.get("contatos_ordem", CONTATOS_ORDEM_PADRAO))
                caminho_jpg = converter_pdf_para_jpg(caminho_pdf, caminho_jpg_base)
                arquivos_gerados.append((info_fornecedor.get("nome", texto_vencedor),
                                         caminho_pdf, caminho_jpg))
            except Exception as e:
                messagebox.showerror("Erro ao gerar Ordem de Compra",
                                     f"Fornecedor {texto_vencedor}:\n{e}")
                return

        resumo = "\n".join(f"• {nome} → {os.path.basename(pdf)}"
                           for nome, pdf, _jpg in arquivos_gerados)
        messagebox.showinfo(
            "Ordens de Compra geradas",
            f"{len(arquivos_gerados)} ordem(ns) de compra gerada(s) em:\n"
            f"{self._pasta_cotacao_atual}\n\n{resumo}")

        # pré-visualiza a primeira ordem gerada
        if arquivos_gerados:
            _nome, _pdf, caminho_jpg = arquivos_gerados[0]
            exibir_preview_pdf_jpg(self, caminho_jpg, "Ordem de Compra")

    def acao_salvar_passo3(self):
        self._salvar_defaults_atuais()
        self._stub("SALVAR")

    # ==================================================================
    # PASSO 27 — Mapa Comparativo (tela + PDF + JPG)
    # ==================================================================
    def _gerar_mapa_comparativo(self):
        if not self.mapa_fornecedores:
            return
        fornecedores_texto = self._obter_fornecedores_texto_completo()
        mapa_itens = montar_mapa_itens_completo(
            self.linhas_resultado, self.mapa_fornecedores, fornecedores_texto, self.saldos_bd_atual)
        if not mapa_itens:
            return

        caminho_json = self._caminho_arquivo_mapa_empresas(
            self._pasta_cotacao_atual, self._codigo_cotacao_atual)
        try:
            with open(caminho_json, "r", encoding="utf-8") as arquivo:
                decisoes = json.load(arquivo).get("decisoes_itens", {})
        except Exception as erro:
            messagebox.showerror("Erro ao carregar decisões", str(erro))
            return
        for item in mapa_itens:
            decisao = decisoes.get(str(item["numero_item"]), {})
            item["propostas_originais"] = list(item["propostas"])
            for campo in ("pesquisas", "precos_retirados", "observacaoitem", "cancelado", "status_item",
                          "titulos_pesquisa", "justificativas_retirada", "retiradas_pendentes"):
                if campo in decisao:
                    item[campo] = decisao[campo]
            aplicar_retiradas(item, decisao.get("fornecedores_retirados", []))
            item["vencedor_idx"] = decisao.get("vencedor_idx", item["vencedor_idx"])

        # Cancelados ficam no documento, mas não participam do fechamento por item.
        mapa_itens = self._perguntar_itens_cancelados(mapa_itens)
        ativos = [item for item in mapa_itens if not item.get("cancelado")]
        pesquisas = self._perguntar_pesquisas_precos(ativos)
        if pesquisas is None:
            return

        # 3) monta os textos automáticos (objeto e observações padrão)
        objeto_auto = gerar_objeto_automatico(mapa_itens)
        itens_poucas_propostas = [formatar_numero_inteiro(it["numero_item"]) for it in mapa_itens
                                   if not it.get("cancelado") and not texto_status_item(it) and 0 < len(it["propostas"]) < 3]
        observacoes_auto = (TEXTO_OBSERVACAO_PADRAO.format(itens=", ".join(itens_poucas_propostas))
                             if itens_poucas_propostas else "")
        observacoes_auto = "\n\n".join(t for t in (observacoes_auto, observacoes_status_itens(mapa_itens)) if t)

        dados_extra = self._perguntar_dados_finalizacao(
            objeto_auto, observacoes_auto, montar_texto_pesquisa_justificativa(mapa_itens))
        dados_extra["titulos_pesquisa"] = pesquisas

        texto_auto = montar_texto_relatorio_pesquisa(mapa_itens, self._codigo_cotacao_atual,
            self._ano_atual_wizard, dados_extra, self.entry_sei.get().strip())
        texto_relatorio = self._editar_relatorio_pesquisa(texto_auto)
        if texto_relatorio is None:
            return
        caminho_relatorio = os.path.join(self._pasta_cotacao_atual,
            "RELATÓRIO DA PESQUISA.PDF")

        base = f"MAPA COMPARATIVO {self._codigo_cotacao_atual}"
        caminho_pdf = os.path.join(self._pasta_cotacao_atual, base + ".pdf")
        caminho_jpg_base = os.path.join(self._pasta_cotacao_atual, base)
        try:
            gerar_pdf_mapa_comparativo(
                caminho_pdf, self._codigo_cotacao_atual, self._ano_atual_wizard, mapa_itens,
                objeto_titulo=dados_extra.get("objeto", ""),
                observacoes=dados_extra.get("observacoes", ""),
                diretor_nome=self.entry_da_nome.get().strip(),
                diretor_matricula=self.entry_da_matricula.get().strip(),
                sei_publicacao=self.entry_sei.get().strip(),  # já vem preenchido na tela principal
                tipo_planilha=self._sufixo_planilha_atual,
                texto_pesquisa=dados_extra["texto_pesquisa"], titulos_pesquisa=pesquisas)
            caminho_excel = os.path.join(self._pasta_cotacao_atual, base + ".xlsx")
            dados_extra["diretor_nome"] = self.entry_da_nome.get().strip()
            dados_extra["diretor_matricula"] = self.entry_da_matricula.get().strip()
            gerar_excel_mapa_comparativo(caminho_excel, self._codigo_cotacao_atual,
                                        self._ano_atual_wizard, mapa_itens, dados_extra)
            caminho_jpg = converter_pdf_para_jpg(caminho_pdf, caminho_jpg_base)
            gerar_pdf_relatorio_pesquisa(caminho_relatorio, texto_relatorio)
        except Exception as e:
            messagebox.showerror("Erro ao gerar Mapa Comparativo", str(e))
            return

        if not self._salvar_decisoes_mapa(mapa_itens, dados_extra):
            return
        # Atualiza também o relatório de preços já existente, sem abrir outra tela.
        from Mapa import montar_ranking_itens, gerar_pdf_ranking
        excluidos = {it["codigo"]: it.get("fornecedores_retirados", []) for it in mapa_itens}
        ranking = montar_ranking_itens(self.mapa_fornecedores, fornecedores_texto, excluidos)
        try:
            gerar_pdf_ranking(os.path.join(self._pasta_cotacao_atual,
                f"MAPA COMPARATIVO PRÉVIO RESUMIDO {self._codigo_cotacao_atual}.pdf"),
                self._codigo_cotacao_atual, self._ano_atual_wizard, ranking)
        except Exception as e:
            messagebox.showwarning("Relação por menor preço", f"Mapa salvo, mas não foi possível atualizar a relação de preços: {e}")
        messagebox.showinfo(
            "Mapa Comparativo gerado",
            f"Arquivos salvos em:\n{caminho_pdf}\n{caminho_jpg}\n{caminho_excel}\n{caminho_relatorio}")

        exibir_preview_pdf_jpg(self, caminho_jpg, "Mapa Comparativo")

    def _editar_relatorio_pesquisa(self, texto_auto):
        caminho = self._caminho_arquivo_mapa_empresas(self._pasta_cotacao_atual, self._codigo_cotacao_atual)
        try:
            with open(caminho, "r", encoding="utf-8") as arquivo:
                salvo = json.load(arquivo).get("relatorio_pesquisa", {})
        except Exception as erro:
            messagebox.showerror("Erro ao carregar relatório", str(erro))
            return None
        janela = ctk.CTkToplevel(self)
        janela.title("Relatório de Pesquisa de Preços — revisar e editar")
        janela.geometry("1280x900")
        maximizar_janela(janela)
        janela.grab_set()
        rodape = ctk.CTkFrame(janela)
        rodape.pack(side="bottom", fill="x", padx=16, pady=12)
        aviso = ("Os dados da pesquisa mudaram desde a última versão. O texto salvo foi preservado; "
                 "use ATUALIZAR COM DADOS ATUAIS para reconstruí-lo e substituir as edições anteriores."
                 if salvo and salvo.get("base") != texto_auto else
                 "Revise o texto antes de gerar o PDF. SALVAR TEXTO guarda as edições e atualiza o PDF do relatório.")
        ctk.CTkLabel(janela, text=aviso, font=("Segoe UI", 18), wraplength=1150,
                     justify="left").pack(fill="x", padx=16, pady=12)
        texto = ctk.CTkTextbox(janela, font=("Segoe UI", 18), wrap="word")
        texto.pack(fill="both", expand=True, padx=16, pady=8)
        texto.insert("1.0", salvo.get("texto", texto_auto))
        resultado = []
        def salvar(gerar_pdf=True):
            conteudo = texto.get("1.0", "end-1c")
            if not conteudo.strip():
                messagebox.showwarning("Relatório vazio", "Preencha o texto do relatório.", parent=janela)
                return False
            try:
                with open(caminho, "r", encoding="utf-8") as arquivo:
                    dados = json.load(arquivo)
                dados["relatorio_pesquisa"] = {"texto": conteudo, "base": base[0]}
                with open(caminho + ".tmp", "w", encoding="utf-8") as arquivo:
                    json.dump(dados, arquivo, ensure_ascii=False, indent=2)
                os.replace(caminho + ".tmp", caminho)
                if gerar_pdf:
                    gerar_pdf_relatorio_pesquisa(
                        os.path.join(self._pasta_cotacao_atual, "RELATÓRIO DA PESQUISA.PDF"), conteudo)
                return True
            except Exception as erro:
                messagebox.showerror("Erro ao salvar relatório", str(erro), parent=janela)
                return False
        base = [salvo.get("base", texto_auto)]
        def atualizar():
            if messagebox.askyesno("Atualizar relatório", "Substituir o texto editado pelo texto gerado com os dados atuais?", parent=janela):
                texto.delete("1.0", "end")
                texto.insert("1.0", texto_auto)
                base[0] = texto_auto
        def gerar():
            if salvar():
                resultado.append(texto.get("1.0", "end-1c"))
                janela.destroy()
        def fechar():
            if salvar(gerar_pdf=False):
                janela.destroy()
        ctk.CTkButton(rodape, text="SALVAR TEXTO", command=salvar, height=44).pack(side="left", padx=8)
        ctk.CTkButton(rodape, text="ATUALIZAR COM DADOS ATUAIS", command=atualizar, height=44).pack(side="left", padx=8)
        ctk.CTkButton(rodape, text="SALVAR E GERAR DOCUMENTOS", command=gerar, height=44).pack(side="right", padx=8)
        janela.protocol("WM_DELETE_WINDOW", fechar)
        self.wait_window(janela)
        return resultado[0] if resultado else None

    def _salvar_campos_telas(self, itens, campos):
        """Atualiza apenas os campos da tela, sem apagar decisões de outros itens."""
        caminho = self._caminho_arquivo_mapa_empresas(
            self._pasta_cotacao_atual, self._codigo_cotacao_atual)
        try:
            with open(caminho, "r", encoding="utf-8") as arquivo:
                dados = json.load(arquivo)
            decisoes = dados.setdefault("decisoes_itens", {})
            for item in itens:
                decisao = decisoes.setdefault(str(item["numero_item"]), {})
                decisao["codigo"] = item["codigo"]
                for campo in campos:
                    decisao[campo] = item.get(campo)
            temporario = caminho + ".tmp"
            with open(temporario, "w", encoding="utf-8") as arquivo:
                json.dump(dados, arquivo, ensure_ascii=False, indent=2)
            os.replace(temporario, caminho)
            return True
        except Exception as erro:
            messagebox.showerror("Erro ao salvar tela", str(erro))
            return False

    def _abrir_simulador_referencia(self, parent, item, pesquisas, titulos):
        """Janela descartável: não altera o item nem grava arquivos da cotação."""
        janela = ctk.CTkToplevel(parent)
        janela.title(f"Simulação da referência — Item {item['numero_item']}")
        janela.geometry("1100x780")
        janela.minsize(900, 600)
        maximizar_janela(janela)
        janela.grab_set()
        fonte = ("Segoe UI", 17)
        ctk.CTkLabel(janela, text=f"SIMULAÇÃO — ITEM {item['numero_item']} — {item['codigo']}",
                     font=("Segoe UI", 21, "bold")).pack(anchor="w", padx=20, pady=(12, 6))
        ctk.CTkLabel(janela, text="Edite qualquer valor, desmarque para retirar ou adicione uma fonte. "
                     "Selecione em NEGOCIAR o preço que deseja ajustar.\n"
                     "Os demais valores incluídos ficam fixos no cálculo do teto. Nada será salvo na cotação.",
                     font=fonte, justify="left", wraplength=1000).pack(anchor="w", padx=20, pady=6)
        rodape = ctk.CTkFrame(janela, fg_color="transparent")
        rodape.pack(side="bottom", fill="x", padx=20, pady=12)
        resultado = ctk.CTkLabel(janela, text="", font=fonte, justify="left", wraplength=1000)
        resultado.pack(side="bottom", fill="x", padx=20, pady=10)
        grade = ctk.CTkScrollableFrame(janela)
        grade.pack(fill="both", expand=True, padx=20, pady=8)
        grade.grid_columnconfigure(1, weight=1)
        for coluna, texto in enumerate(("INCLUIR", "FONTE / FORNECEDOR", "VALOR UNITÁRIO", "NEGOCIAR")):
            ctk.CTkLabel(grade, text=texto, font=FONTE_BOTAO).grid(row=0, column=coluna, padx=12, pady=8)
        linhas = []
        alvo = tk.IntVar(master=janela, value=0)

        def alterado():
            resultado.configure(text="Cenário alterado. Clique em CALCULAR SIMULAÇÃO.", text_color=COR_TEXTO)

        def adicionar(nome="Nova fonte", valor=None, incluir=True):
            i = len(linhas)
            var = tk.BooleanVar(master=janela, value=incluir)
            ctk.CTkCheckBox(grade, text="", variable=var, width=30,
                            command=alterado).grid(row=i+1, column=0, padx=12, pady=6)
            nome_entry = ctk.CTkEntry(grade, font=fonte, height=36)
            nome_entry.grid(row=i+1, column=1, padx=12, pady=6, sticky="ew")
            nome_entry.insert(0, nome)
            entry = ctk.CTkEntry(grade, font=fonte, height=36, width=170)
            entry.grid(row=i+1, column=2, padx=12, pady=6)
            texto_valor = formatar_pesquisa(valor)
            if valor is not None and Decimal(str(valor)) != Decimal(str(valor)).quantize(Decimal('.01')):
                texto_valor = f"{valor:,.4f}".replace(",", "X").replace(".", ",").replace("X", ".")
            entry.insert(0, texto_valor)
            ctk.CTkRadioButton(grade, text="", variable=alvo, value=i, width=30, command=alterado).grid(
                row=i+1, column=3, padx=12, pady=6)
            linhas.append((var, nome_entry, entry))
            alterado()
            entry.bind("<KeyRelease>", lambda ev: alterado())

        chaves_ativas = {p.get("fornecedor_chave", p["fornecedor"]) for p in item["propostas"]}
        propostas = item.get("propostas_originais") or item["propostas"]
        for p in propostas:
            adicionar(p["fornecedor"], p["unitario"],
                      p.get("fornecedor_chave", p["fornecedor"]) in chaves_ativas)
        for i, valor in enumerate(pesquisas):
            adicionar(f"Pesquisa {i+1} — {titulos[i]}", valor, valor is not None)
        escolhido = item.get("vencedor_idx") or 0
        if item["propostas"]:
            chave = item["propostas"][min(escolhido, len(item["propostas"])-1)].get("fornecedor_chave",
                        item["propostas"][min(escolhido, len(item["propostas"])-1)]["fornecedor"])
            alvo.set(next((i for i, p in enumerate(propostas)
                           if p.get("fornecedor_chave", p["fornecedor"]) == chave), 0))

        def calcular(usar_teto=False):
            try:
                valores = [numero_pesquisa(campo.get()) if var.get() else None
                           for var, _, campo in linhas]
                dados = simular_referencia(valores, alvo.get())
            except ValueError:
                resultado.configure(text="Valor inválido. Use números positivos, como 2.950,00, ou deixe vazio.",
                                    text_color="#FF7777")
                return
            if usar_teto and dados['limite'] is not None and dados['limite'] > 0:
                campo = linhas[alvo.get()][2]
                campo.delete(0, "end")
                campo.insert(0, formatar_pesquisa(dados['limite']))
                return calcular()
            if not dados['quantidade']:
                resultado.configure(text="Inclua pelo menos um valor para calcular a referência.", text_color=COR_TEXTO)
                return
            texto = (f"Valores incluídos: {dados['quantidade']}    |    Média: {formatar_pesquisa(dados['media'])}"
                     f"    |    Mediana: {formatar_pesquisa(dados['mediana'])}\n"
                     f"REFERÊNCIA SIMULADA: {formatar_pesquisa(dados['referencia'])}\n")
            if dados['alvo'] is None:
                texto += "Inclua e preencha o valor selecionado em NEGOCIAR para calcular seu teto."
            else:
                nome = linhas[alvo.get()][1].get()
                texto += (f"{nome}: {formatar_pesquisa(dados['alvo'])} — "
                          + ("DENTRO" if dados['aceita'] else "ACIMA") + " da referência simulada.\n")
                if dados['limite'] is None:
                    texto += "Sem outros valores incluídos, a referência é o próprio preço; não há teto independente."
                else:
                    texto += (f"Maior preço em centavos que fica dentro da referência recalculada: "
                              f"{formatar_pesquisa(dados['limite'])}.\n"
                              "Regra: menor entre média e mediana, mantendo os demais valores deste cenário. "
                              "As retiradas são manuais; os limites de 50% não excluem valores automaticamente.")
            resultado.configure(text=texto, text_color=COR_TEXTO)

        def fechar():
            janela.destroy()
            parent.grab_set()

        ctk.CTkButton(rodape, text="ADICIONAR VALOR", command=adicionar).pack(side="left", padx=4)
        ctk.CTkButton(rodape, text="CALCULAR SIMULAÇÃO", command=calcular).pack(side="left", padx=4)
        ctk.CTkButton(rodape, text="TESTAR PREÇO LIMITE", command=lambda: calcular(True)).pack(side="left", padx=4)
        ctk.CTkButton(rodape, text="FECHAR SIMULAÇÃO", command=fechar).pack(side="right", padx=4)
        janela.protocol("WM_DELETE_WINDOW", fechar)
        calcular()

    def _perguntar_pesquisas_precos(self, mapa_itens):
        """Fecha pesquisa, revisão e vencedor de um item antes de avançar."""
        ativos = [item for item in mapa_itens if not item.get("cancelado")]
        if not ativos:
            return list(TITULOS_PESQUISA_PADRAO)
        for item in ativos:
            item.setdefault("propostas_originais", list(item["propostas"]))
            for campo in ("fornecedores_retirados", "precos_retirados", "retiradas_pendentes"):
                item.setdefault(campo, [])
            item.setdefault("justificativas_retirada", {})
        janela = ctk.CTkToplevel(self)
        janela.title("Pesquisa de preços por item — revisão e vencedor")
        janela.geometry("1280x940")
        maximizar_janela(janela)
        janela.minsize(1050, 780)
        janela.grab_set()
        fonte = ("Segoe UI", 17)
        destaque = ("Segoe UI", 19, "bold")
        linha_item = ctk.CTkFrame(janela, fg_color="transparent")
        linha_item.pack(fill="x", padx=20, pady=(10, 6))
        contador_item = ctk.CTkLabel(linha_item, text="", font=destaque, anchor="e", width=120)
        contador_item.pack(side="right", padx=(16, 0))
        titulo_item = ctk.CTkLabel(linha_item, text="", font=destaque, wraplength=0, anchor="w", width=1)
        titulo_item.pack(side="left", fill="x", expand=True)
        cabecalho = ctk.CTkFrame(janela)
        cabecalho.pack(fill="x", padx=20, pady=(0, 8))
        titulos, entradas = [], []
        def formatar_unitario(campo):
            try:
                valor = numero_pesquisa(campo.get())
            except ValueError:
                return  # A validação ao recalcular mantém o valor inválido visível.
            campo.delete(0, "end")
            campo.insert(0, formatar_pesquisa(valor))

        for i, default in enumerate(TITULOS_PESQUISA_PADRAO):
            bloco = ctk.CTkFrame(cabecalho, fg_color="transparent")
            bloco.pack(side="left", fill="x", expand=True, padx=(6, 40) if i == 0 else (6, 6), pady=8)
            ctk.CTkLabel(bloco, text=f"PESQUISA DE PREÇOS {i+1} —", font=("Segoe UI", 14, "bold")).pack(side="left", padx=(0, 4))
            titulo = ctk.CTkEntry(bloco, font=("Segoe UI", 16), height=38, width=210)
            titulo.pack(side="left", padx=4)
            titulo.insert(0, default)
            titulos.append(titulo)
            ctk.CTkLabel(bloco, text=f"UNITÁRIO {i+1}", font=("Segoe UI", 14, "bold")).pack(side="left", padx=4)
            valor = ctk.CTkEntry(bloco, font=("Segoe UI", 20), height=38, width=120)
            valor.pack(side="left", padx=(4, 0))
            valor.bind("<FocusOut>", lambda event, campo=valor: formatar_unitario(campo))
            entradas.append(valor)
        corpo = ctk.CTkScrollableFrame(janela)
        corpo.pack(fill="both", expand=True, padx=20)
        rodape = ctk.CTkFrame(janela, border_width=2, border_color=COR_DESTAQUE)
        rodape.pack(fill="x", padx=20, pady=10)
        botoes = ctk.CTkFrame(janela, fg_color="transparent")
        botoes.pack(fill="x", padx=20, pady=(0, 16))
        estado = {"indice": 0}
        resultado = []
        selecoes = []
        justificativas = {}
        entry_observacao = None
        vencedor = tk.IntVar(master=janela, value=0)
        aplicado = [None]

        def atuais():
            return ativos[estado["indice"]]

        def nomes():
            return [e.get().strip() or default for e, default in zip(titulos, TITULOS_PESQUISA_PADRAO)]

        def ler():
            try:
                return [numero_pesquisa(e.get()) for e in entradas]
            except ValueError:
                messagebox.showerror("Preço inválido", "Use valores positivos, como 1.234,56, ou deixe vazio.", parent=janela)
                return None

        def abrir_simulacao():
            precos = ler()
            if precos is not None:
                item_simulado = dict(atuais(), vencedor_idx=vencedor.get())
                self._abrir_simulador_referencia(janela, item_simulado, precos, nomes())

        def selecionar_status(status):
            if ler() is None:
                return
            item = atuais()
            preservar_selecoes()
            item["status_item"] = "" if item.get("status_item") == status else status
            vencedor.set(-1 if texto_status_item(item) else (0 if item["propostas"] else -1))
            atualizar_revisao()
            salvar_atual()

        def selecionar_empresa():
            preservar_selecoes()
            atuais()["status_item"] = ""
            atualizar_revisao()
            salvar_atual()

        def indice_vencedor():
            item = atuais()
            return vencedor.get() if item["propostas"] and not texto_status_item(item) else None

        def preservar_selecoes():
            atuais()["retiradas_pendentes"] = [[grupo, c] for grupo, c, var in selecoes if var.get()]
            atuais()["justificativas_retirada"] = {c: campo.get() for c, campo in justificativas.items()}

        def retornar(chave):
            precos = ler()
            if precos is None:
                return
            item = atuais()
            item["pesquisas"] = precos
            item["vencedor_idx"] = indice_vencedor()
            item["retiradas_pendentes"] = [[grupo, c] for grupo, c, var in selecoes if var.get()]
            item["justificativas_retirada"] = {c: campo.get() for c, campo in justificativas.items()}
            if not retornar_fornecedor(item, chave):
                messagebox.showwarning("Retornar fornecedor", "A proposta original não foi localizada no cadastro deste item.", parent=janela)
                return
            vencedor.set(item["vencedor_idx"] if item["vencedor_idx"] is not None else -1)
            aplicado[0] = list(precos)
            atualizar_revisao()
            salvar_atual()

        def mostrar_retirado(parent, proposta, largura):
            linha = ctk.CTkFrame(parent, fg_color="transparent")
            linha.pack(fill="x", padx=10, pady=5)
            chave = proposta.get("fornecedor_chave", proposta["fornecedor"])
            ctk.CTkButton(linha, text="RETORNAR", width=110, font=FONTE_BOTAO,
                          command=lambda c=chave: retornar(c)).pack(side="left", padx=(0, 10))
            ctk.CTkLabel(linha, text=f"{proposta['fornecedor']} — {formatar_pesquisa(proposta['unitario'])} — "
                         f"{proposta.get('justificativa', 'Retirada por limite de 50%')}",
                         font=fonte, text_color="#FF5555", wraplength=largura,
                         justify="left").pack(side="left", fill="x", expand=True)

        def atualizar_revisao():
            item = atuais()
            status_texto = texto_status_item(item)
            if status_texto:
                vencedor.set(-1)
            for status, botao in botoes_status.items():
                selecionado = item.get("status_item") == status
                botao.configure(fg_color="#805000" if selecionado else COR_DESTAQUE,
                                text_color="#FFFF80" if selecionado else "#FFFFFF")
            for widget in revisao.winfo_children():
                widget.destroy()
            selecoes.clear()
            justificativas.clear()
            ref = item["referencias"]
            texto = "   |   ".join(f"{nome}: {formatar_pesquisa(ref[chave]) or '—'}" for nome, chave in
                [("Média", "media"), ("Mediana", "mediana")])
            linha_resumo = ctk.CTkFrame(revisao, fg_color="transparent")
            linha_resumo.pack(fill="x", pady=6)
            resumo_frame = ctk.CTkFrame(linha_resumo, fg_color="transparent")
            resumo_frame.pack(anchor="center")
            ctk.CTkLabel(resumo_frame, text=texto, font=destaque).pack(side="left", padx=12)
            referencia_card = ctk.CTkFrame(resumo_frame, fg_color="#174D3B", border_width=2, border_color="#65E6AC")
            referencia_card.pack(side="left", padx=16)
            ctk.CTkLabel(referencia_card, text=f"REFERÊNCIA: {formatar_pesquisa(ref['minimo']) or '—'}",
                         font=("Segoe UI", 23, "bold"), text_color="#FFFFFF").pack(padx=12, pady=6)
            ctk.CTkButton(linha_resumo, text="SIMULAR REFERÊNCIA", font=FONTE_BOTAO,
                          command=abrir_simulacao).pack(pady=6)
            valores = [(p["fornecedor"], p["unitario"], "fornecedor", p.get("fornecedor_chave", p["fornecedor"]))
                       for p in item["propostas"]]
            for titulo, chave, acima in [("PREÇOS ACIMA DE 50%", "superior", True),
                                        ("PREÇOS ABAIXO DE 50%", "inferior", False)]:
                limite = ref[chave]
                grupo = ctk.CTkFrame(revisao, fg_color="transparent")
                grupo.pack(fill="x", pady=(10, 4))
                validos = ctk.CTkFrame(grupo, fg_color="transparent")
                validos.pack(side="left", fill="both", expand=True, padx=6)
                removidos = ctk.CTkFrame(grupo, fg_color="transparent")
                removidos.pack(side="left", fill="both", expand=True, padx=6)
                ctk.CTkLabel(validos, text=f"{titulo} — limite {formatar_pesquisa(limite) or '—'}", font=destaque).pack(anchor="w")
                ctk.CTkLabel(removidos, text=f"FORNECEDORES RETIRADOS - 50% {'ACIMA' if acima else 'ABAIXO'}", font=destaque, text_color="#FF5555").pack(anchor="w")
                encontrados = [] if limite is None else [v for v in valores if (v[1] > limite if acima else v[1] < limite)]
                for nome, valor, _, identificador in encontrados:
                    var = tk.BooleanVar(master=janela, value=False)
                    ctk.CTkCheckBox(validos, text=f"RETIRAR VALOR — {nome}: {formatar_pesquisa(valor)}",
                        variable=var, font=fonte, height=32).pack(anchor="w", padx=10, pady=5)
                    var.set([chave, identificador] in item.get("retiradas_pendentes", []))
                    selecoes.append((chave, identificador, var))
                if not encontrados:
                    ctk.CTkLabel(validos, text="Nenhum fornecedor fora deste limite.", font=fonte).pack(anchor="w", padx=10)
                historico = [p for p in item.get("precos_retirados", []) if p.get("grupo") == chave]
                for proposta in historico:
                    mostrar_retirado(removidos, proposta, 400)
                if not historico:
                    ctk.CTkLabel(removidos, text="Nenhum fornecedor retirado.", font=fonte, text_color="#FF5555").pack(anchor="w", padx=10)
            for widget in rodape.winfo_children():
                widget.destroy()
            ctk.CTkLabel(rodape, text="CLASSIFICAÇÃO ATUAL - selecione a empresa vencedora deste item ou marque RETIRAR VALOR caso a proposta enviada não atenda ao descritivo do Termo de Referência, ou conforme Parecer Técnico", font=destaque, wraplength=1400, justify="left").pack(anchor="w", padx=12, pady=8)
            classificacao = ctk.CTkScrollableFrame(rodape, height=180)
            classificacao.pack(fill="x", padx=10, pady=(0, 8))
            for pos, proposta in enumerate(item["propostas"]):
                nome, _ = extrair_nome_codigo_fornecedor(proposta["fornecedor"])
                nome = truncar_em_palavra(nome, 60)
                linha = ctk.CTkFrame(classificacao, fg_color="transparent")
                linha.pack(fill="x", padx=8, pady=4)
                ctk.CTkRadioButton(linha, text=(f"{status_texto} — " if status_texto else "") +
                    f"{pos+1}º colocado — {nome} — Unitário: {formatar_pesquisa(proposta['unitario'])}",
                    variable=vencedor, value=pos, command=selecionar_empresa,
                    font=fonte, height=30).pack(side="left", padx=(8, 12))
                chave = proposta.get("fornecedor_chave", proposta["fornecedor"])
                grupo = "analise"
                var = tk.BooleanVar(master=janela, value=[grupo, chave] in item.get("retiradas_pendentes", []))
                ctk.CTkCheckBox(linha, text="RETIRAR VALOR", variable=var,
                    font=fonte, height=32).pack(side="left", padx=12)
                motivo = ctk.CTkEntry(linha, placeholder_text="Justificativa da retirada (obrigatória)",
                                      font=fonte, width=600)
                motivo.pack(side="left", fill="x", expand=True, padx=(12, 8))
                motivo.insert(0, item.get("justificativas_retirada", {}).get(chave,
                    "Proposta enviada não atende ao descritivo do Termo de Referência conforme"))
                selecoes.append((grupo, chave, var))
                justificativas[chave] = motivo
            historico_analise = [p for p in item.get("precos_retirados", []) if p.get("grupo") == "analise"]
            if historico_analise:
                ctk.CTkLabel(revisao, text="FORNECEDORES RETIRADOS — ANÁLISE DA PROPOSTA",
                             font=destaque, text_color="#FF5555").pack(anchor="w", padx=10)
                for proposta in historico_analise:
                    mostrar_retirado(revisao, proposta, 900)
            if not item["propostas"]:
                ctk.CTkLabel(rodape, text=status_texto or "Sem propostas classificadas — item sem vencedor.", font=fonte).pack(pady=10)

        def recalcular():
            precos = ler()
            if precos is None:
                return False
            item = atuais()
            idx = vencedor.get()
            escolhido = item["propostas"][idx] if 0 <= idx < len(item["propostas"]) else None
            fornecedores = [chave for _, chave, var in selecoes if var.get()]
            if any(var.get() and grupo == "analise" and not justificativas[chave].get().strip() for grupo, chave, var in selecoes):
                messagebox.showwarning("Justificativa obrigatória",
                    "Informe a justificativa de cada fornecedor marcado para retirada.", parent=janela)
                return False
            historico = item.setdefault("precos_retirados", [])
            registrados = set()
            for grupo, chave, var in selecoes:
                if var.get() and chave not in registrados:
                    registrados.add(chave)
                    proposta = next(p for p in item["propostas"] if p.get("fornecedor_chave", p["fornecedor"]) == chave)
                    historico.append(dict(fornecedor=proposta["fornecedor"], unitario=proposta["unitario"], grupo=grupo,
                        fornecedor_chave=chave, justificativa=(justificativas[chave].get().strip() if grupo == "analise"
                            else "Preço acima de 50% da referência" if grupo == "superior"
                            else "Preço abaixo de 50% da referência")))
            item["pesquisas"] = precos
            item["retiradas_pendentes"] = []
            item["justificativas_retirada"] = {chave: campo.get() for chave, campo in justificativas.items()}
            aplicar_retiradas(item, fornecedores)
            for i, entry in enumerate(entradas):
                entry.delete(0, "end")
                entry.insert(0, formatar_pesquisa(item["pesquisas"][i]))
            vencedor.set(-1 if texto_status_item(item) else next((i for i, p in enumerate(item["propostas"]) if p is escolhido), 0))
            aplicado[0] = list(item["pesquisas"])
            atualizar_revisao()
            return True

        def mostrar_item():
            item = atuais()
            descricao = " ".join(item["descricao"].split())
            titulo_item.configure(text=f"ITEM {item['numero_item']} — {item['codigo']} — {descricao}")
            contador_item.configure(text=f"Item {estado['indice']+1}/{len(ativos)}")
            janela.title(f"Pesquisa de preços — Item {estado['indice']+1} de {len(ativos)}")
            for widget in corpo.winfo_children():
                widget.destroy()
            selecoes.clear()
            item["pesquisas"] = valores_pesquisa(item)
            recalcular_referencias(item)
            vencedor.set(-1 if texto_status_item(item) else (item.get("vencedor_idx") or 0))
            for campo, texto in zip(titulos, item.get("titulos_pesquisa") or nomes()):
                campo.delete(0, "end")
                campo.insert(0, texto)
            for i, entry in enumerate(entradas):
                entry.delete(0, "end")
                entry.insert(0, formatar_pesquisa(item["pesquisas"][i]))
            nonlocal revisao, entry_observacao
            ctk.CTkLabel(corpo, text="OBSERVAÇÃOITEM — será incluída na pesquisa de preços e justificativa", font=fonte).pack(anchor="w", padx=12, pady=(10, 4))
            entry_observacao = ctk.CTkTextbox(corpo, font=("Segoe UI", 18), height=90)
            entry_observacao.pack(fill="x", padx=12, pady=(0, 8))
            entry_observacao.insert("1.0", item.get("observacaoitem", ""))
            revisao = ctk.CTkFrame(corpo, fg_color="transparent")
            revisao.pack(fill="x", pady=8)
            aplicado[0] = list(item["pesquisas"])
            atualizar_revisao()
            voltar_botao.configure(state="normal" if estado["indice"] else "disabled")
            avancar.configure(text="CONFIRMAR EMPRESA / STATUS\nE PRÓXIMO" if estado["indice"] < len(ativos)-1 else "CONFIRMAR EMPRESA / STATUS\nE CONCLUIR")

        campos_pesquisa = ("pesquisas", "titulos_pesquisa", "observacaoitem", "vencedor_idx", "status_item",
                           "fornecedores_retirados", "precos_retirados", "referencias",
                           "justificativas_retirada", "retiradas_pendentes")

        def salvar_atual():
            precos = ler()
            if precos is None:
                return False
            item = atuais()
            item["pesquisas"] = precos
            recalcular_referencias(item)
            item["titulos_pesquisa"] = nomes()
            item["observacaoitem"] = entry_observacao.get("1.0", "end").strip()
            item["vencedor_idx"] = indice_vencedor()
            item["justificativas_retirada"] = {chave: campo.get() for chave, campo in justificativas.items()}
            item["retiradas_pendentes"] = [[grupo, chave] for grupo, chave, var in selecoes if var.get()]
            return self._salvar_campos_telas([item], campos_pesquisa)

        def fechar():
            if salvar_atual():
                janela.destroy()

        def limpar():
            item = atuais()
            item["propostas"] = list(item["propostas_originais"])
            item.update(pesquisas=[None, None], titulos_pesquisa=list(TITULOS_PESQUISA_PADRAO),
                        observacaoitem="", status_item="", vencedor_idx=0 if item["propostas"] else None,
                        fornecedores_retirados=[], precos_retirados=[],
                        justificativas_retirada={}, retiradas_pendentes=[])
            recalcular_referencias(item)
            self._salvar_campos_telas([item], campos_pesquisa)
            mostrar_item()

        def confirmar_item():
            precos = ler()
            if precos is None:
                return
            if precos != aplicado[0] or any(var.get() for _, _, var in selecoes):
                messagebox.showwarning("Recalcular", "Clique em RECALCULAR para atualizar os preços e aplicar as retiradas antes de confirmar.", parent=janela)
                return
            item = atuais()
            item["vencedor_idx"] = indice_vencedor()
            item["observacaoitem"] = entry_observacao.get("1.0", "end").strip()
            if not salvar_atual():
                return
            if estado["indice"] == len(ativos)-1:
                resultado.extend(nomes())
                janela.destroy()
            else:
                estado["indice"] += 1
                mostrar_item()

        def voltar():
            if estado["indice"] == 0 or not recalcular():
                return
            item = atuais()
            item["vencedor_idx"] = indice_vencedor()
            item["observacaoitem"] = entry_observacao.get("1.0", "end").strip()
            if not salvar_atual():
                return
            estado["indice"] -= 1
            mostrar_item()

        revisao = None
        voltar_botao = ctk.CTkButton(botoes, text="VOLTAR", font=("Segoe UI", 14), height=44, width=85, command=voltar)
        voltar_botao.pack(side="left", padx=(0, 12))
        ctk.CTkButton(botoes, text="RECALCULAR", font=("Segoe UI", 16, "bold"), height=44, width=135, command=recalcular).pack(side="left", padx=(0, 12))
        ctk.CTkButton(botoes, text="LIMPAR DADOS", font=("Segoe UI", 14), width=105, command=limpar).pack(side="left", padx=8)
        janela.protocol("WM_DELETE_WINDOW", fechar)
        botoes_status = {}
        for status, descricao in STATUS_ITEM.items():
            botao = ctk.CTkButton(botoes, text=status, font=("Segoe UI", 14, "bold"),
                                width=90, height=44, command=lambda s=status: selecionar_status(s))
            botao.pack(side="left", padx=3)
            adicionar_dica_status(botao, f"DISPENSA DE LICITAÇÃO {status}\n{descricao}")
            botoes_status[status] = botao
        avancar = ctk.CTkButton(botoes, text="", font=("Segoe UI", 13), width=210, height=44, command=confirmar_item)
        avancar.pack(side="left", padx=(8, 0))
        mostrar_item()
        self.wait_window(janela)
        return resultado or None

    def _perguntar_itens_cancelados(self, mapa_itens):
        """Tela simples para marcar quais itens devem aparecer como
        CANCELADO no Mapa Comparativo — mesmo que já tenham valores e
        fornecedores lançados, eles ficam suprimidos (igual ao DESERTO, mas
        com o texto 'CANCELADO' nas 3 colunas de colocado)."""
        janela = ctk.CTkToplevel(self)
        janela.title("Itens cancelados")
        janela.configure(fg_color=COR_FUNDO_APP)
        janela.geometry("1200x800")
        janela.minsize(800, 500)
        maximizar_janela(janela)
        janela.grab_set()

        # Reserva o rodapé antes do conteúdo expansível para manter os botões visíveis.
        rodape = ctk.CTkFrame(janela, fg_color="transparent")
        rodape.pack(side="bottom", fill="x", padx=16, pady=(8, 16))

        ctk.CTkLabel(
            janela,
            text="Marque os itens que devem aparecer como CANCELADO no Mapa Comparativo "
                 "(os valores e fornecedores desse item ficam suprimidos, mesmo que já tenham "
                 "sido lançados). Deixe todos desmarcados se nenhum item foi cancelado.",
            font=("Segoe UI", 20, "bold"), text_color=COR_TEXTO, wraplength=850, justify="left"
        ).pack(anchor="w", padx=16, pady=(16, 8))

        scroll = ctk.CTkScrollableFrame(janela, fg_color="#161616", height=200)
        scroll.pack(fill="both", expand=True, padx=16, pady=(0, 10))

        variaveis = {}
        def salvar():
            for item in mapa_itens:
                item["cancelado"] = variaveis[item["numero_item"]].get()
            return self._salvar_campos_telas(mapa_itens, ("cancelado",))

        def limpar():
            for var in variaveis.values():
                var.set(False)
            salvar()

        for item in mapa_itens:
            var = tk.BooleanVar(value=bool(item.get("cancelado", False)))
            variaveis[item["numero_item"]] = var
            texto = f"Item {item['numero_item']} — {item['codigo']} — {item['descricao'][:75]}"
            ctk.CTkCheckBox(scroll, text=texto, variable=var, font=("Segoe UI", 20), height=36, command=salvar,
                            text_color=COR_TEXTO).pack(anchor="w", padx=10, pady=5)

        def confirmar():
            if salvar():
                janela.destroy()

        ctk.CTkButton(rodape, text="LIMPAR DADOS", font=("Segoe UI", 20),
                      height=44, command=limpar).pack(side="left", padx=(0, 12))
        ctk.CTkButton(rodape, text="Confirmar e continuar", font=FONTE_BOTAO, height=44,
                      fg_color=COR_DESTAQUE, command=confirmar).pack(side="right")
        janela.protocol("WM_DELETE_WINDOW", confirmar)
        self.wait_window(janela)
        return mapa_itens

    def _perguntar_dados_finalizacao(self, objeto_auto="", observacoes_auto="", pesquisa_auto=TEXTO_PESQUISA_PADRAO):
        """Pergunta/confirma o objeto (título do mapa) e as observações —
        ambos já vêm pré-preenchidos automaticamente, mas continuam
        editáveis. O SEI de publicação é lido direto do campo SEI da tela
        principal, sem precisar perguntar de novo."""
        janela = ctk.CTkToplevel(self)
        janela.title("Dados do Mapa Comparativo")
        janela.configure(fg_color=COR_FUNDO_APP)
        janela.geometry("1200x950")
        maximizar_janela(janela)
        janela.minsize(950, 750)
        janela.grab_set()

        conteudo = ctk.CTkScrollableFrame(janela)
        conteudo.pack(fill="both", expand=True, padx=16, pady=10)
        ctk.CTkLabel(conteudo, text="OBJETO da cotação (título do mapa):", font=("Segoe UI", 20, "bold"),
                     text_color=COR_TEXTO).pack(anchor="w", padx=16, pady=(16, 4))
        entry_objeto = ctk.CTkTextbox(conteudo, height=100, font=("Segoe UI", 18))
        entry_objeto.pack(fill="x", padx=16)
        if objeto_auto:
            entry_objeto.insert("1.0", objeto_auto)

        ctk.CTkLabel(conteudo, text="Pesquisa de preços e justificativa (editável):",
                     font=("Segoe UI", 20, "bold")).pack(anchor="w", padx=16, pady=(12, 4))
        entry_pesquisa = ctk.CTkTextbox(conteudo, height=350, font=("Segoe UI", 18))
        entry_pesquisa.pack(fill="x", padx=16)
        entry_pesquisa.insert("1.0", pesquisa_auto)

        ctk.CTkLabel(conteudo, text="Observações (editável — já preenchido quando aplicável):",
                     font=("Segoe UI", 20, "bold"), text_color=COR_TEXTO).pack(anchor="w", padx=16, pady=(12, 4))
        entry_obs = ctk.CTkTextbox(conteudo, height=180, font=("Segoe UI", 18))
        entry_obs.pack(fill="x", padx=16)
        if observacoes_auto:
            entry_obs.insert("1.0", observacoes_auto)

        resultado = {}

        def confirmar():
            resultado["texto_pesquisa"] = entry_pesquisa.get("1.0", "end").strip()
            resultado["objeto"] = entry_objeto.get("1.0", "end").strip()
            resultado["observacoes"] = entry_obs.get("1.0", "end").strip()
            janela.destroy()

        ctk.CTkButton(janela, text="Gerar Mapa Comparativo", font=("Segoe UI", 20, "bold"), height=46,
                      fg_color=COR_DESTAQUE, command=confirmar).pack(pady=18)
        janela.protocol("WM_DELETE_WINDOW", confirmar)
        self.wait_window(janela)
        return resultado

    def _salvar_decisoes_mapa(self, mapa_itens, dados_extra=None):
        """Grava no mesmo MAPA_EMPRESAS_{cotação}.json quem foi o vencedor
        escolhido (ou item cancelado) para cada item, para o botão ORDEM
        poder reaproveitar sem perguntar de novo."""
        caminho_json = self._caminho_arquivo_mapa_empresas(
            self._pasta_cotacao_atual, self._codigo_cotacao_atual)
        conteudo = {}
        if os.path.exists(caminho_json):
            try:
                with open(caminho_json, "r", encoding="utf-8") as f:
                    conteudo = json.load(f)
            except Exception as e:
                messagebox.showerror("Erro ao salvar", f"Não foi possível ler o cadastro existente: {e}")
                return False
        conteudo["cotacao"] = self._codigo_cotacao_atual
        conteudo["ano"] = self._ano_atual_wizard
        conteudo["fornecedores"] = self.mapa_fornecedores
        if dados_extra is not None:
            conteudo["dados_mapa"] = dados_extra
        conteudo["decisoes_itens"] = {
            str(item["numero_item"]): {
                "codigo": item["codigo"],
                "vencedor_idx": item.get("vencedor_idx"),
                "cancelado": bool(item.get("cancelado")),
                "status_item": item.get("status_item", ""),
                "fornecedores_retirados": item.get("fornecedores_retirados", []),
                "precos_retirados": item.get("precos_retirados", []),
                "observacaoitem": item.get("observacaoitem", ""),
                "pesquisas": valores_pesquisa(item),
                "titulos_pesquisa": item.get("titulos_pesquisa", list(TITULOS_PESQUISA_PADRAO)),
                "justificativas_retirada": item.get("justificativas_retirada", {}),
                "retiradas_pendentes": item.get("retiradas_pendentes", []),
                "referencias": item.get("referencias", {}),
            }
            for item in mapa_itens
        }
        try:
            with open(caminho_json, "w", encoding="utf-8") as f:
                json.dump(conteudo, f, ensure_ascii=False, indent=2)
        except Exception as e:
            messagebox.showerror("Erro ao salvar", f"Não foi possível salvar as decisões:\n{e}")
            return False
        return True


def re_sub_nome_arquivo(nome):
    """Remove caracteres inválidos para nome de arquivo do Windows."""
    import re
    return re.sub(r"[^\w\- ]", "", nome or "").strip()
