import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
import fitz
import AdjudicacaoOrdem as adj


def item(precos=(5650, 7500, 12990), pesquisas=(8000, 9213.5), numero=1):
    return dict(numero_item=numero, codigo=f'VME{numero:02}', descricao='Conserto ventilador pulmonar',
                un='Serviço', saldo=65492.11, vencedor_idx=0, pesquisas=list(pesquisas),
                propostas=[dict(fornecedor=f'Fornecedor {i}', fornecedor_chave=f'F{i}',
                                unitario=v, total=v, quantidade=1, prazo=10, marca='Serviço')
                           for i, v in enumerate(precos)])


class PesquisasTest(unittest.TestCase):
    def test_exemplo_imagem(self):
        it = item()
        adj.recalcular_referencias(it)
        self.assertEqual(it['referencias'], dict(media=8670.7, mediana=8000,
                         minimo=8000, superior=12000, inferior=4000))

    def test_milhar_vazio_invalido(self):
        self.assertEqual(adj.numero_pesquisa('1.234,56'), 1234.56)
        self.assertEqual(adj.formatar_pesquisa(1234.56), '1.234,56')
        self.assertIsNone(adj.numero_pesquisa(''))
        for valor in ['-1', '0', 'nan', 'inf', 'abc']:
            with self.assertRaises(ValueError):
                adj.numero_pesquisa(valor)

    def test_retirada_por_item_e_pesquisa(self):
        primeiro, segundo = item(), item(numero=2)
        adj.aplicar_retiradas(primeiro, ['F2'], [1])
        self.assertEqual([p['unitario'] for p in primeiro['propostas']], [5650, 7500])
        self.assertEqual(primeiro['referencias']['media'], 7050)
        self.assertEqual(primeiro['referencias']['mediana'], 7500)
        self.assertEqual(primeiro['referencias']['minimo'], 7050)
        self.assertEqual(len(segundo['propostas']), 3)
        adj.aplicar_retiradas(primeiro, ['F0', 'F1'], [0])
        self.assertTrue(all(v is None for v in primeiro['referencias'].values()))
        self.assertEqual(adj.montar_declaracao_vencedores([primeiro]), [])

    def test_todos_fornecedores_e_excel(self):
        it = item(precos=(10, 20, 30, 100), pesquisas=(40, 50))
        adj.recalcular_referencias(it)
        self.assertAlmostEqual(it['referencias']['media'], 250/6)
        with tempfile.TemporaryDirectory() as d:
            arquivo = Path(d)/'mapa.xlsx'
            adj.gerar_excel_mapa_comparativo(arquivo, '178', 2026, [it, item(numero=2)], {})
            wb = load_workbook(arquivo)
            ws = wb['Adjudicação']
            self.assertEqual(ws['K8'].value, 'Média¹ Mapa GDF')
            self.assertEqual(ws['L8'].value, 'Mediana² Painel Federal')
            self.assertEqual(ws['K9'].value, 40)
            self.assertIn('C2:C5', ws['M9'].value)
            self.assertIn('MEDIAN(', ws['M11'].value)
            self.assertIn('MIN(M9,M11)', ws['M12'].value)
            self.assertIn('M12*1.5', ws['N11'].value)
            self.assertIn('M12*0.5', ws['N12'].value)
            self.assertIn('MEDIAN(', ws['M15'].value)
            self.assertEqual(ws['A57'].value, adj.TEXTO_PESQUISA_PADRAO)
            self.assertIn('A57:M57', str(ws.merged_cells))
            self.assertIn('$A$1:$M$92', ws.print_area)
            self.assertEqual(ws['K9'].number_format, '#,##0.00')
            wb.close()

    def test_pdf_sem_excluido(self):
        it = item()
        adj.aplicar_retiradas(it, ['F2'])
        with tempfile.TemporaryDirectory() as d:
            arquivo = Path(d)/'mapa.pdf'
            adj.gerar_pdf_mapa_comparativo(str(arquivo), '178', '2026', [it],
                                          texto_pesquisa='Texto vigente personalizado',
                                          titulos_pesquisa=('Pesquisa atual 1', 'Pesquisa atual 2'))
            with fitz.open(arquivo) as pdf:
                texto = ''.join(p.get_text() for p in pdf)
            self.assertIn('Fornecedor 0', texto)
            self.assertNotIn('Fornecedor 2', texto)
            self.assertIn('Texto vigente personalizado', texto)
            self.assertIn('Pesquisa atual 1', ' '.join(texto.split()))
            self.assertIn('MEDIANA', texto)

    def test_paginacao_excel_sem_perda(self):
        itens = [item(numero=i) for i in range(1, 14)]
        with tempfile.TemporaryDirectory() as d:
            caminho = Path(d)/'mapa.xlsx'
            adj.gerar_excel_mapa_comparativo(caminho, '178', 2026, itens, {})
            wb = load_workbook(caminho)
            self.assertEqual(wb['Adjudicação 2']['A9'].value, 13)
            self.assertIn("'Valores classificados 2'!", wb['Adjudicação 2']['M9'].value)
            self.assertEqual(wb['Adjudicação 2']['A57'].value, adj.TEXTO_PESQUISA_PADRAO)
            self.assertIn('$A$1:$M$92', wb['Adjudicação 2'].print_area)
            wb.close()

    def test_ranking_respeita_retirada_por_item(self):
        from Mapa import montar_ranking_itens
        fornecedores = {'A': {'itens': [dict(codigo='1', unitario=10), dict(codigo='2', unitario=10)]},
                        'B': {'itens': [dict(codigo='1', unitario=20)]}}
        ranking = dict(montar_ranking_itens(fornecedores, {}, {'1': ['A']}))
        self.assertEqual([p['fornecedor'] for p in ranking['1']['propostas']], ['B'])
        self.assertEqual([p['fornecedor'] for p in ranking['2']['propostas']], ['A'])

    def test_pdf_cinco_itens_uma_pagina(self):
        itens = [item(numero=i) for i in range(1, 6)]
        for it in itens:
            adj.recalcular_referencias(it)
        with tempfile.TemporaryDirectory() as d:
            arquivo = Path(d)/'mapa.pdf'
            adj.gerar_pdf_mapa_comparativo(str(arquivo), '178', '2026', itens)
            with fitz.open(arquivo) as pdf:
                self.assertEqual(len(pdf), 1)
                self.assertIn('ADJUDICAÇÃO E HOMOLOGAÇÃO', pdf[0].get_text())

    def test_ordem_recarrega_exclusoes(self):
        class Entry:
            def get(self):
                return ''
        class App(adj.AdjudicacaoOrdemMixin):
            linhas_resultado = [dict(ITEM=1, CODIGO='VME01', DESCRICAO='Teste', UN='un')]
            saldos_bd_atual = {}
            entry_sei = Entry()
            def _preparar_contexto_wizard(self):
                return ('propostas', '178', '2026')
            def _obter_caminho_arquivo(self):
                return 'matrix.xlsx'
            def _caminho_arquivo_mapa_empresas(self, *args):
                return self.arquivo
        with tempfile.TemporaryDirectory() as d:
            app = App()
            app._pasta_cotacao_atual = d
            app.arquivo = str(Path(d)/'mapa.json')
            fornecedores = {f'F{i}': dict(itens=[dict(codigo='VME01', unitario=v, quantidade=1,
                                      total=v, prazo=10, marca='')]) for i,v in enumerate([1, 100])}
            Path(app.arquivo).write_text(json.dumps(dict(fornecedores=fornecedores,
                decisoes_itens={'1': dict(vencedor_idx=0, fornecedores_retirados=['F0'])})))
            with patch.object(adj, 'carregar_fornecedores_completo', return_value={}):
                dados = app._montar_dados_ordens_ativas()
            vencedor, itens = next(iter(dados['ordens_por_fornecedor'].values()))
            self.assertEqual(vencedor, 'F1')
            self.assertEqual(itens[0][1]['unitario'], 100)


class RelatorioPesquisaTest(unittest.TestCase):
    def test_metodologia_e_observacoes_sem_repeticao_na_justificativa(self):
        primeiro, segundo = item(), item(numero=2)
        primeiro['observacaoitem'] = 'Nota fiscal complementar.\nNão integra o cálculo.'
        segundo['observacaoitem'] = 'Análise técnica registrada.'
        dados = {'texto_pesquisa': adj.montar_texto_pesquisa_justificativa([primeiro, segundo]) +
                 '\nComplemento manual da justificativa.'}
        texto = adj.montar_texto_relatorio_pesquisa([primeiro, segundo], '178', '2026', dados)
        topico4 = texto.split('4. PESQUISA DE PREÇOS E JUSTIFICATIVA')[1].split('5. OBSERVAÇÕES DO MAPA COMPARATIVO')[0]
        for observacao in (primeiro['observacaoitem'], segundo['observacaoitem']):
            self.assertEqual(texto.count(observacao), 1)
            self.assertNotIn(observacao, topico4)
        self.assertIn(adj.TEXTO_PESQUISA_PADRAO, topico4)
        self.assertIn('Complemento manual da justificativa.', topico4)
        self.assertIn('1,5 (Preço exorbitante) e por 0,5 (Preço inexequível)', texto)
        self.assertNotIn('não comprova se houve consulta', texto)
        self.assertNotIn('A retirada depende da seleção', texto)
        self.assertNotIn('Este documento não presume', texto)
        # A justificativa do mapa permanece com as observações originais.
        self.assertIn(primeiro['observacaoitem'], dados['texto_pesquisa'])

    def test_conteudo_calculos_retiradas_cancelado_e_pdf(self):
        it = item(precos=(10, 20, 300), pesquisas=(None, 30))
        it['observacaoitem'] = 'NF de outro CNPJ: apenas reforço, fora do cálculo.'
        it['precos_retirados'] = [dict(fornecedor='Fornecedor 2', fornecedor_chave='F2',
            unitario=300, grupo='superior', justificativa='Preço acima de 50%'),
            dict(fornecedor='Retirado técnico', unitario=15, grupo='analise', justificativa='Despacho 123'),
            dict(fornecedor='Retirado baixo', unitario=1, grupo='inferior', justificativa='Preço abaixo de 50%')]
        adj.aplicar_retiradas(it, ['F2'])
        cancelado = item(numero=2)
        cancelado['cancelado'] = True
        dados = dict(objeto='Objeto <teste> & equipamentos', texto_pesquisa='Justificativa editada',
                     observacoes='Observação geral')
        texto = adj.montar_texto_relatorio_pesquisa([it, cancelado], '178', '2026', dados, 'SEI 123')
        for esperado in ('R$ 20,0000', 'totalizando 3 valor(es)', 'Despacho 123',
                         '50% ACIMA', '50% ABAIXO', 'ANÁLISE DA PROPOSTA',
                         'Justificativa editada', 'Observação geral', it['observacaoitem'],
                         'CANCELADO', 'não informado', 'SEI 123'):
            self.assertIn(esperado, texto)
        with tempfile.TemporaryDirectory() as pasta:
            caminho = str(Path(pasta) / 'relatorio.pdf')
            adj.gerar_pdf_relatorio_pesquisa(caminho, texto + '\n' + ('Texto longo de revisão. ' * 800))
            with fitz.open(caminho) as pdf:
                self.assertGreater(len(pdf), 1)
                extraido = ''.join(p.get_text() for p in pdf)
                self.assertIn('Objeto <teste> & equipamentos', extraido)
                self.assertIn('Despacho 123', extraido)
                self.assertIn('Página 2', extraido)


if __name__ == '__main__':
    unittest.main()
