"""Integração do comando do botão principal com o fluxo de adjudicação.

Os diálogos são respondidos sem um display; a geração e a persistência são reais.
"""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook
import GeradorCotacao as principal
import AdjudicacaoOrdem as adjudicacao


class Widget:
    def __init__(self, *args, **kwargs):
        pass
    def pack(self, *args, **kwargs):
        pass
    def grid(self, *args, **kwargs):
        pass
    def grid_columnconfigure(self, *args, **kwargs):
        pass


class Button(Widget):
    botoes = {}
    def __init__(self, *args, text, command, **kwargs):
        self.command = command
        self.botoes[text] = self
    def invoke(self):
        return self.command()


class BotaoAdjudicacaoTest(unittest.TestCase):
    def criar_app_sem_display(self):
        classe = principal.criar_app()
        app = classe.__new__(classe)
        Button.botoes = {}
        with patch.object(principal.ctk, 'CTkFrame', Widget), \
             patch.object(principal.ctk, 'CTkLabel', Widget), \
             patch.object(principal.ctk, 'CTkButton', Button):
            app._montar_grupos_botoes(None)
        return app

    def test_botao_avisa_se_nao_ha_cotacao(self):
        app = self.criar_app_sem_display()
        app.linhas_resultado = []
        with patch.object(adjudicacao.messagebox, 'showwarning') as aviso:
            Button.botoes['ADJUDICAÇÃO'].invoke()
        aviso.assert_called_once()
        self.assertIn('CONSULTAR COTAÇÃO', aviso.call_args.args[1])

    def test_botao_executa_fluxo_anterior_e_novas_etapas(self):
        with tempfile.TemporaryDirectory() as pasta:
            app = self.criar_app_sem_display()
            app.linhas_resultado = [dict(ITEM=1, CODIGO='VME01', DESCRICAO='Conserto', UN='Serviço')]
            app.saldos_bd_atual = {}
            app._pasta_cotacao_atual = pasta
            app._codigo_cotacao_atual = '178'
            app._ano_atual_wizard = '2026'
            app._sufixo_planilha_atual = 'HRG'
            for nome in ('entry_da_nome', 'entry_da_matricula', 'entry_sei'):
                setattr(app, nome, Mock(get=Mock(return_value='')))
            cadastro = Path(pasta)/'MAPA_EMPRESAS_178.json'
            fornecedores = {nome: dict(itens=[dict(codigo='VME01', unitario=valor,
                quantidade=1, total=valor, prazo=10, marca='Serviço')])
                for nome, valor in [('A', 5650), ('B', 7500), ('C', 12990)]}
            cadastro.write_text(json.dumps(dict(fornecedores=fornecedores)), encoding='utf-8')
            app._preparar_contexto_wizard = Mock(return_value=(pasta, '178', '2026'))
            app._caminho_arquivo_mapa_empresas = Mock(return_value=str(cadastro))
            app._obter_fornecedores_texto_completo = Mock(return_value={})
            etapas = []
            def pesquisas(itens):
                etapas.append('pesquisas')
                for item in itens:
                    item['pesquisas'] = [8000, 9213.5]
                    adjudicacao.recalcular_referencias(item)
                return list(adjudicacao.TITULOS_PESQUISA_PADRAO)
            def revisar(itens, titulos):
                etapas.append('revisao')
                adjudicacao.aplicar_retiradas(itens[0], ['C'])
                return True
            def cancelados(itens):
                etapas.append('cancelados')
                itens[0]['cancelado'] = False
                return itens
            def vencedores(itens):
                etapas.append('vencedores')
                itens[0]['vencedor_idx'] = 0
                return itens
            def dados(*args):
                etapas.append('dados_mapa')
                return dict(objeto='Conserto', observacoes='', texto_pesquisa='Justificativa vigente')
            app._perguntar_pesquisas_precos = pesquisas
            app._revisar_precos = revisar
            app._perguntar_itens_cancelados = cancelados
            app._perguntar_vencedores = vencedores
            app._perguntar_dados_finalizacao = dados
            with patch.object(adjudicacao.messagebox, 'showinfo'), \
                 patch.object(adjudicacao.messagebox, 'showwarning') as aviso, \
                 patch.object(adjudicacao.messagebox, 'showerror') as erro, \
                 patch.object(adjudicacao, 'exibir_preview_pdf_jpg') as preview:
                Button.botoes['ADJUDICAÇÃO'].invoke()
            erro.assert_not_called()
            aviso.assert_not_called()
            self.assertEqual(etapas, ['pesquisas', 'revisao', 'cancelados', 'vencedores', 'dados_mapa'])
            preview.assert_called_once()
            for nome in ('MAPA COMPARATIVO 178.pdf', 'MAPA COMPARATIVO 178.xlsx',
                         'RELACAO MENOR PRECO 178.pdf'):
                self.assertTrue((Path(pasta)/nome).is_file(), nome)
            self.assertTrue(list(Path(pasta).glob('MAPA COMPARATIVO 178*.jpg')))
            salvo = json.loads(cadastro.read_text(encoding='utf-8'))
            self.assertEqual(salvo['fornecedores'], fornecedores)
            self.assertEqual(salvo['decisoes_itens']['1']['fornecedores_retirados'], ['C'])
            self.assertEqual(salvo['decisoes_itens']['1']['vencedor_idx'], 0)
            self.assertEqual(salvo['dados_mapa']['texto_pesquisa'], 'Justificativa vigente')
            wb = load_workbook(Path(pasta)/'MAPA COMPARATIVO 178.xlsx')
            self.assertNotIn('C', [wb['Valores classificados'].cell(r, 2).value
                                for r in range(2, wb['Valores classificados'].max_row+1)])
            wb.close()


if __name__ == '__main__':
    unittest.main()
