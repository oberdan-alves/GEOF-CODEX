# -*- coding: utf-8 -*-
"""
Mapa.py
=======
MÓDULO DO 2º PASSO — Preparar Mapa Comparativo
--------------------------------------------------
Responsável pelos botões do grupo "2º Passo – Preparar Mapa Comparativo":
    - ACESSAR PROPOSTAS   (abre no Explorer a pasta PROPOSTAS da cotação
                             ativa, criando-a se ainda não existir)
    - LISTAR EMPRESAS      (casa arquivos da pasta PROPOSTAS com a aba
                             Fornecedores da planilha Matrix)
    - INCLUIR EMPRESAS MAPA (assistente/wizard de cadastro de propostas,
                             item a item, por fornecedor)
    - CARREGAR EMPRESAS MAPA (recarrega um cadastro já salvo em disco)

Também inclui:
    - Leitura da aba "Fornecedores" e da aba "BD" (saldos) da planilha
    - Montagem do ranking de propostas por menor preço (Passo 26) e
      geração do respectivo PDF

O cadastro de propostas por fornecedor é salvo em
    COTACAO\\{código}\\MAPA_EMPRESAS_{código}.json
e é reaproveitado pelo módulo AdjudicacaoOrdem.py (3º Passo) para montar o
Mapa Comparativo e as Ordens de Compra.
"""

import os
import json
from xml.sax.saxutils import escape
from decimal import Decimal, ROUND_HALF_UP
from math import isfinite
import tkinter as tk
from Mensagens import messagebox

import customtkinter as ctk
from openpyxl import load_workbook

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from GeradorCotacao import (
    centralizar_janela,
    COR_FUNDO_APP, COR_CARD, COR_TEXTO, COR_TEXTO_SEC, COR_DESTAQUE,
    FONTE_ENTRY, FONTE_SECAO, FONTE_BOTAO, FONTE_TABELA,
    normalizar_texto, extrair_nome_codigo_fornecedor, corrigir_acentuacao,
    formatar_data, formatar_apenas_digitos, formatar_moeda,
)


def numero_proposta(texto):
    """Aceita decimal com ponto/vírgula e milhar no formato brasileiro."""
    if texto is None or not str(texto).strip():
        return None
    texto = str(texto).strip().replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        numero = float(texto)
        return numero if isfinite(numero) else None
    except ValueError:
        return None


def formatar_quantidade_proposta(valor):
    numero = numero_proposta(valor)
    if numero is None:
        return "" if valor is None else str(valor)
    return format(Decimal(str(numero)).normalize(), "f").replace(".", ",")


def formatar_unitario_proposta(valor):
    numero = numero_proposta(valor)
    if numero is None:
        return "" if valor is None else str(valor)
    return format(Decimal(str(numero)).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP), ".4f").replace(".", ",")


# ----------------------------------------------------------------------
# LEITURA DA ABA "Fornecedores"
# ----------------------------------------------------------------------
def carregar_fornecedores(caminho_arquivo):
    """Lê a coluna A da aba 'Fornecedores' e devolve uma lista de dicts:
    texto_completo (ex.: 'ETICA (703252)'), nome (sem código), nome_normalizado, codigo."""
    wb = load_workbook(caminho_arquivo, data_only=True, read_only=True)
    if "Fornecedores" not in wb.sheetnames:
        raise ValueError("A aba 'Fornecedores' não foi encontrada na planilha.")
    ws = wb["Fornecedores"]
    fornecedores = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        valor = corrigir_acentuacao(row[0])
        if not valor:
            continue
        nome, codigo = extrair_nome_codigo_fornecedor(valor)
        fornecedores.append({
            "texto_completo": str(valor).strip(),
            "nome": nome,
            "nome_normalizado": normalizar_texto(nome),
            "codigo": codigo,
        })
    wb.close()
    return fornecedores


def carregar_fornecedores_completo(caminho_arquivo):
    """Lê a aba 'Fornecedores' com todas as colunas usadas na ORDEM DE COMPRA:
    A=texto completo, B=Razão social, C=Contato, D=Telefone, E=Código SES,
    F=CNPJ, G=Banco, H=Nº do banco, I=Agência, J=Nº Conta.
    Devolve um dict {nome_normalizado (da coluna A, sem código): {dados}}."""
    wb = load_workbook(caminho_arquivo, data_only=True, read_only=True)
    if "Fornecedores" not in wb.sheetnames:
        raise ValueError("A aba 'Fornecedores' não foi encontrada na planilha.")
    ws = wb["Fornecedores"]
    fornecedores = {}
    for row in ws.iter_rows(min_row=2, max_col=10, values_only=True):
        valor_a = corrigir_acentuacao(row[0])
        if not valor_a:
            continue
        nome, codigo = extrair_nome_codigo_fornecedor(valor_a)
        chave = normalizar_texto(nome)
        fornecedores[chave] = {
            "texto_completo": str(valor_a).strip(),
            "nome": nome,
            "codigo": codigo,
            "razao_social": corrigir_acentuacao(row[1]) or "",
            "contato": corrigir_acentuacao(row[2]) or "",
            "telefone": row[3] or "",
            "codigo_ses": row[4] or codigo or "",
            "cnpj": row[5] or "",
            "banco": corrigir_acentuacao(row[6]) or "",
            "banco_numero": row[7] or "",
            "agencia": row[8] or "",
            "conta": row[9] or "",
        }
    wb.close()
    return fornecedores


def buscar_candidatos_fornecedor(nome_extraido, fornecedores):
    """Encontra fornecedores cujo nome (sem código) é igual ao nome extraído do
    arquivo de proposta, ou começa com ele seguido de espaço (ex.: 'ELLO' casa
    com 'ELLO 0001' e 'ELLO 0002'). Ignora acentuação."""
    alvo = normalizar_texto(nome_extraido)
    candidatos = []
    for f in fornecedores:
        if f["nome_normalizado"] == alvo or f["nome_normalizado"].startswith(alvo + " "):
            candidatos.append(f)
    return candidatos


def listar_arquivos_propostas(pasta_propostas, codigo_cotacao):
    """Retorna a lista de nomes de empresa extraídos dos arquivos
    'PROPOSTA {codigo} {NOME}' dentro da pasta informada, ou None se a pasta
    não existir."""
    if not os.path.isdir(pasta_propostas):
        return None
    prefixo = f"PROPOSTA {codigo_cotacao} ".upper()
    encontrados = []
    for arquivo in sorted(os.listdir(pasta_propostas)):
        caminho_completo = os.path.join(pasta_propostas, arquivo)
        if not os.path.isfile(caminho_completo):
            continue
        nome_sem_ext, _ = os.path.splitext(arquivo)
        if nome_sem_ext.upper().startswith(prefixo):
            nome_extraido = nome_sem_ext[len(prefixo):].strip()
            if nome_extraido:
                encontrados.append(nome_extraido)
    return encontrados


# ----------------------------------------------------------------------
# SALDOS DISPONÍVEIS (aba BD) - usados no cadastro de propostas
# ----------------------------------------------------------------------
def carregar_saldos_bd(caminho_arquivo):
    """Lê a aba 'BD': coluna A = CÓDIGO, coluna C = valor disponível (saldo).
    Devolve um dict {codigo_normalizado: saldo (float ou None)}."""
    wb = load_workbook(caminho_arquivo, data_only=True, read_only=True)
    if "BD" not in wb.sheetnames:
        raise ValueError("A aba 'BD' não foi encontrada na planilha.")
    ws = wb["BD"]
    saldos = {}
    for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        codigo = row[0]
        saldo = row[2] if len(row) > 2 else None
        if codigo is None or str(codigo).strip() == "":
            continue
        chave = normalizar_texto(codigo)
        try:
            saldos[chave] = float(saldo) if saldo is not None else None
        except (TypeError, ValueError):
            saldos[chave] = None
    wb.close()
    return saldos


# ----------------------------------------------------------------------
# PASSO 26 — RANKING DE PROPOSTAS POR MENOR PREÇO (tela + PDF)
# ----------------------------------------------------------------------
def montar_ranking_itens(mapa_fornecedores, fornecedores_texto_completo, excluidos_por_codigo=None):
    """A partir do cadastro de propostas por fornecedor, monta, para cada
    código de item, a lista de propostas ordenada por menor preço unitário.
    Retorna uma lista de tuplas (codigo, {"descricao","un","propostas":[...]})
    na ordem de primeira aparição dos itens."""
    itens_index = {}
    ordem_codigos = []
    for nome_forn, dados in mapa_fornecedores.items():
        texto_completo = fornecedores_texto_completo.get(normalizar_texto(nome_forn), nome_forn)
        for item in dados.get("itens", []):
            unitario = item.get("unitario")
            if unitario is None or unitario <= 0:
                continue  # fornecedor não cotou este item
            codigo = item["codigo"]
            if normalizar_texto(nome_forn) in (excluidos_por_codigo or {}).get(codigo, []):
                continue
            if codigo not in itens_index:
                itens_index[codigo] = {"numero_item": item.get("numero_item"), "descricao": item.get("descricao", ""),
                                        "un": item.get("un", ""), "propostas": []}
                ordem_codigos.append(codigo)
            itens_index[codigo]["propostas"].append({
                "fornecedor": texto_completo,
                "quantidade": item.get("quantidade"),
                "unitario": unitario,
                "total": item.get("total"),
                "prazo": item.get("prazo"),
                "marca": item.get("marca"),
            })
    for codigo in ordem_codigos:
        itens_index[codigo]["propostas"].sort(key=lambda p: p["unitario"])
    return [(codigo, itens_index[codigo]) for codigo in ordem_codigos]


def gerar_pdf_ranking(caminho_pdf, codigo_cotacao, ano, ranking):
    """Gera o PDF com a relação de propostas por item, da menor para a maior,
    destacando (negrito) o 1º, 2º e 3º colocados; o 1º com fonte maior."""
    doc = SimpleDocTemplate(caminho_pdf, pagesize=landscape(A4),
                             topMargin=1.2 * cm, bottomMargin=1.2 * cm,
                             leftMargin=1 * cm, rightMargin=1 * cm)
    styles = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle("TituloRel", parent=styles["Title"], fontSize=15)
    story = [
        Paragraph(f"MAPA COMPARATIVO PRÉVIO RESUMIDO — Cotação {codigo_cotacao}/{ano}",
                  titulo_estilo),
        Spacer(1, 10),
    ]

    dados_tabela = [["ITEM", "CÓDIGO", "FORNECEDOR", "QUANTIDADE", "UNITÁRIO", "TOTAL",
                     "PRAZO", "MARCA"]]
    from reportlab.pdfbase.pdfmetrics import stringWidth
    larguras = [1.5 * cm, 2.2 * cm, 6 * cm, 2.6 * cm, 2.6 * cm, 2.6 * cm, 2 * cm, 3.5 * cm]
    estilo_descricao = ParagraphStyle("ItemResumo", parent=styles["Normal"],
        fontName="Helvetica-Bold", fontSize=10, leading=12, textColor=colors.red)
    estilos_extra = []
    linha_idx = 1
    for numero_item, (codigo, info) in enumerate(ranking, start=1):
        numero = formatar_quantidade_proposta(info.get("numero_item") if info.get("numero_item") is not None else numero_item)
        descricao = " ".join(str(info.get("descricao", "")).split())
        limite = sum(larguras[2:]) - 12
        if stringWidth(descricao, "Helvetica-Bold", 10) > limite:
            while descricao and stringWidth(descricao + "...", "Helvetica-Bold", 10) > limite:
                descricao = descricao[:-1]
            descricao = descricao.rstrip() + "..."
        dados_tabela.append([numero, codigo, Paragraph(escape(descricao), estilo_descricao), "", "", "", "", ""])
        estilos_extra += [("SPAN", (2, linha_idx), (7, linha_idx)),
            ("TEXTCOLOR", (0, linha_idx), (-1, linha_idx), colors.red),
            ("FONTNAME", (0, linha_idx), (-1, linha_idx), "Helvetica-Bold"),
            ("BACKGROUND", (0, linha_idx), (-1, linha_idx), colors.HexColor("#fff0f0")),
            ("ALIGN", (2, linha_idx), (7, linha_idx), "LEFT")]
        if info["propostas"]:
            estilos_extra.append(("NOSPLIT", (0, linha_idx), (-1, linha_idx+1)))
        linha_idx += 1
        for posicao, proposta in enumerate(info["propostas"], start=1):
            dados_tabela.append([
                f"{posicao}º COLOCADO", "", proposta["fornecedor"],
                formatar_quantidade_proposta(proposta["quantidade"]), formatar_moeda(proposta["unitario"]),
                formatar_moeda(proposta["total"]),
                f'{proposta["prazo"]} dias' if proposta["prazo"] else "",
                proposta["marca"] or "",
            ])
            estilos_extra.append(("SPAN", (0, linha_idx), (1, linha_idx)))
            if posicao <= 3:
                estilos_extra.append(("FONTNAME", (0, linha_idx), (-1, linha_idx), "Helvetica-Bold"))
                estilos_extra.append(("FONTSIZE", (0, linha_idx), (-1, linha_idx),
                                      12 if posicao == 1 else 10))
                cor_fundo = colors.HexColor("#d7f2d7") if posicao == 1 else colors.HexColor("#f2f2f2")
                estilos_extra.append(("BACKGROUND", (0, linha_idx), (-1, linha_idx), cor_fundo))
            linha_idx += 1

    tabela = Table(dados_tabela, repeatRows=1,
                    colWidths=larguras)
    estilo_base = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f2f2f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ] + estilos_extra
    tabela.setStyle(TableStyle(estilo_base))
    story.append(tabela)
    doc.build(story)


# ----------------------------------------------------------------------
# MIXIN — métodos de ação ligados aos botões do 2º Passo
# ----------------------------------------------------------------------
class MapaMixin:
    """Métodos herdados por AppCotacao. Cuida da coleta das propostas dos
    fornecedores (wizard) e monta o cadastro salvo em
    MAPA_EMPRESAS_{cotação}.json, reaproveitado no 3º Passo."""

    # ------------------------------------------------------------------
    # ACESSAR PROPOSTAS
    # ------------------------------------------------------------------
    def acao_acessar_propostas(self):
        """Abre no Explorer a pasta PROPOSTAS da cotação ativa:
        PASTA DE TRABALHO/SUFIXO_PLANILHA/PDPAS {ANO}/COTACAO/{código}/PROPOSTAS
        (ex.: G:/Meu Drive/GEOF/HRG/PDPAS 2026/COTACAO/155/PROPOSTAS), criando-a
        se ainda não existir. Reaproveita a mesma validação/criação de
        pasta usada em CRIAR PASTA COTAÇÃO e AJUSTE DO PEDIDO (Cotacao.py)."""
        contexto = self._obter_pasta_cotacao_ativa(incluir_propostas=True)
        if contexto is None:
            return
        pasta_propostas, codigo_cotacao, ano = contexto
        try:
            os.startfile(pasta_propostas)
        except Exception as e:
            messagebox.showerror("Erro ao abrir pasta", f"Não foi possível abrir:\n{pasta_propostas}\n\n{e}")

    # ------------------------------------------------------------------
    # LISTAR EMPRESAS
    # ------------------------------------------------------------------
    def acao_listar_empresas(self):
        pasta_geof = self.entry_pasta_geof.get().strip()
        ano = self.entry_ano.get().strip()

        if not pasta_geof:
            messagebox.showwarning("Atenção", "Informe a PASTA DE TRABALHO (caminho até \\GEOF).")
            return
        if not ano or len(ano) != 4:
            messagebox.showwarning("Atenção", "Informe o ANO com 4 dígitos.")
            return
        if self.indice_cotacao_ativa is None:
            messagebox.showwarning("Atenção", "Clique numa das caixas de COTAÇÃO para defini-la como ativa.")
            return

        codigo_cotacao = self.entradas_cotacao[self.indice_cotacao_ativa].get().strip()
        if len(codigo_cotacao) != 3:
            messagebox.showwarning("Atenção", "A COTAÇÃO ativa precisa ter o código completo (3 dígitos).")
            return

        caminho_arquivo = self._obter_caminho_arquivo()
        if not caminho_arquivo:
            messagebox.showwarning(
                "Atenção",
                "Não foi possível localizar o arquivo Matrix (aba Fornecedores).\n"
                "Clique em CONSULTAR MATRIX antes de listar as empresas.")
            return

        pasta_propostas = os.path.join(
            pasta_geof, self._sufixo_planilha_atual, f"PDPAS {ano}", "COTACAO",
            codigo_cotacao, "PROPOSTAS")

        nomes_extraidos = listar_arquivos_propostas(pasta_propostas, codigo_cotacao)
        if nomes_extraidos is None:
            messagebox.showerror("Pasta não encontrada",
                                  f"Pasta de PROPOSTAS não encontrada em:\n{pasta_propostas}")
            return
        if not nomes_extraidos:
            messagebox.showinfo(
                "Nenhum arquivo encontrado",
                f"Nenhum arquivo 'PROPOSTA {codigo_cotacao} ...' foi encontrado em:\n{pasta_propostas}")
            return

        try:
            fornecedores = carregar_fornecedores(caminho_arquivo)
        except Exception as e:
            messagebox.showerror("Erro ao ler aba Fornecedores", str(e))
            return

        resolvidos = []
        erros = []
        for nome_extraido in nomes_extraidos:
            candidatos = buscar_candidatos_fornecedor(nome_extraido, fornecedores)
            if len(candidatos) == 1:
                resolvidos.append(candidatos[0]["nome"])
            elif len(candidatos) > 1:
                escolhido = self._perguntar_escolha_fornecedor(nome_extraido, candidatos)
                if escolhido:
                    resolvidos.append(escolhido["nome"])
                else:
                    erros.append(nome_extraido)
            else:
                erros.append(nome_extraido)

        resolvidos_ordenados = sorted(resolvidos, key=normalizar_texto)
        linhas_erro = [f"{nome} não localizado verifique antes de incluir empresas no mapa!"
                       for nome in erros]

        self.lista_empresas_mapa = resolvidos_ordenados
        texto_final = ", ".join(resolvidos_ordenados + linhas_erro)
        self._mostrar_resultado_empresas(texto_final, codigo_cotacao, pasta_propostas)

    def _perguntar_escolha_fornecedor(self, nome_extraido, candidatos):
        """Abre uma janela modal para o usuário escolher, entre os candidatos
        encontrados na aba Fornecedores, qual é a empresa participante."""
        janela = ctk.CTkToplevel(self)
        janela.title(f"Selecionar empresa - {nome_extraido}")
        janela.configure(fg_color=COR_FUNDO_APP)
        janela.geometry(f"540x{140 + 36 * len(candidatos)}")
        centralizar_janela(janela)
        janela.grab_set()

        ctk.CTkLabel(
            janela,
            text=f'O arquivo indica "{nome_extraido}", mas há mais de uma empresa '
                 "correspondente na aba Fornecedores.\nSelecione a empresa participante:",
            font=FONTE_ENTRY, text_color=COR_TEXTO, justify="left",
            wraplength=490).pack(padx=16, pady=(16, 10), anchor="w")

        var_escolha = tk.IntVar(value=0)
        for i, cand in enumerate(candidatos):
            ctk.CTkRadioButton(janela, text=cand["texto_completo"], variable=var_escolha,
                               value=i, font=FONTE_ENTRY, text_color=COR_TEXTO
                               ).pack(anchor="w", padx=30, pady=4)

        resultado = {"valor": None}

        def confirmar():
            resultado["valor"] = candidatos[var_escolha.get()]
            janela.destroy()

        ctk.CTkButton(janela, text="Confirmar", font=FONTE_BOTAO, height=32,
                      command=confirmar).pack(pady=16)

        janela.protocol("WM_DELETE_WINDOW", confirmar)
        self.wait_window(janela)
        return resultado["valor"]

    def _mostrar_resultado_empresas(self, texto_final, codigo_cotacao, pasta_propostas):
        janela = ctk.CTkToplevel(self)
        janela.title(f"Empresas listadas - Cotação {codigo_cotacao}")
        janela.configure(fg_color=COR_FUNDO_APP)
        janela.geometry("760x380")
        centralizar_janela(janela)
        janela.grab_set()

        ctk.CTkLabel(janela, text=f"Empresas encontradas para a COTAÇÃO {codigo_cotacao}",
                     font=FONTE_SECAO, text_color=COR_TEXTO).pack(anchor="w", padx=16, pady=(16, 2))
        ctk.CTkLabel(janela, text=pasta_propostas, font=("Segoe UI", 11),
                     text_color=COR_TEXTO_SEC).pack(anchor="w", padx=16, pady=(0, 10))

        caixa = ctk.CTkTextbox(janela, font=FONTE_ENTRY, height=200, wrap="word")
        caixa.pack(fill="both", expand=True, padx=16, pady=(0, 10))
        caixa.insert("1.0", texto_final)
        caixa.configure(state="disabled")

        linha_botoes = ctk.CTkFrame(janela, fg_color="transparent")
        linha_botoes.pack(pady=(0, 16))
        ctk.CTkButton(linha_botoes, text="Copiar", font=FONTE_BOTAO, height=32,
                      command=lambda: self._copiar_para_area_transferencia(texto_final)
                      ).pack(side="left", padx=6)
        ctk.CTkButton(linha_botoes, text="Fechar", font=FONTE_BOTAO, height=32,
                      command=janela.destroy).pack(side="left", padx=6)

    # ==================================================================
    # PASSO 25 — INCLUIR EMPRESAS MAPA (cadastro de propostas por fornecedor)
    # ==================================================================
    def acao_incluir_empresas_mapa(self):
        if not self.lista_empresas_mapa:
            messagebox.showwarning(
                "Atenção",
                "Nenhuma empresa foi listada ainda. Clique em LISTAR EMPRESAS antes.")
            return
        if not self.linhas_resultado:
            messagebox.showwarning(
                "Atenção",
                "Nenhum item de cotação carregado. Clique em CONSULTAR COTAÇÃO antes.")
            return

        contexto = self._preparar_contexto_wizard()
        if contexto is None:
            return
        pasta_propostas, codigo_cotacao, ano = contexto

        # Tenta carregar cadastro já salvo anteriormente (para continuar de onde parou)
        self.mapa_fornecedores = {}
        caminho_json = self._caminho_arquivo_mapa_empresas(self._pasta_cotacao_atual, codigo_cotacao)
        if os.path.exists(caminho_json):
            try:
                with open(caminho_json, "r", encoding="utf-8") as f:
                    dados_salvos = json.load(f)
                self.mapa_fornecedores = dados_salvos.get("fornecedores", {})
            except Exception as e:
                messagebox.showwarning("Aviso", f"Não foi possível ler o cadastro salvo:\n{e}")

        self.indice_fornecedor_atual = 0
        self._dados_form_atual = None
        self._abrir_janela_wizard()

    def acao_carregar_empresas_mapa(self):
        contexto = self._preparar_contexto_wizard()
        if contexto is None:
            return
        pasta_propostas, codigo_cotacao, ano = contexto

        caminho_json = self._caminho_arquivo_mapa_empresas(self._pasta_cotacao_atual, codigo_cotacao)
        if not os.path.exists(caminho_json):
            messagebox.showwarning(
                "Nenhum cadastro encontrado",
                f"Não há um cadastro salvo em:\n{caminho_json}\n\n"
                "Use o botão INCLUIR EMPRESAS MAPA para começar um novo cadastro.")
            return
        try:
            with open(caminho_json, "r", encoding="utf-8") as f:
                dados_salvos = json.load(f)
        except Exception as e:
            messagebox.showerror("Erro ao carregar", f"Não foi possível ler o arquivo:\n{e}")
            return

        self.mapa_fornecedores = dados_salvos.get("fornecedores", {})
        if not self.mapa_fornecedores:
            messagebox.showwarning("Atenção", "O arquivo salvo não contém fornecedores.")
            return

        # Usa a lista de fornecedores que estão no arquivo salvo (ordem alfabética),
        # já que pode ter sido gerado numa sessão anterior com outra listagem.
        self.lista_empresas_mapa = sorted(self.mapa_fornecedores.keys(), key=normalizar_texto)
        self.indice_fornecedor_atual = 0
        self._dados_form_atual = None
        self._abrir_janela_wizard()

    def _preparar_contexto_wizard(self):
        """Valida os campos e monta os caminhos das pastas + carrega os
        saldos da aba BD. Retorna (pasta_propostas, codigo_cotacao, ano) ou
        None se algo estiver faltando (já exibe o aviso/erro correspondente).

        pasta_propostas (…\\COTACAO\\{código}\\PROPOSTAS) é usada só para ler
        os arquivos de proposta dos fornecedores. Os documentos gerados pelo
        app (MAPA_EMPRESAS, MAPA_COMPARATIVO, RELAÇÃO, ORDEM) são salvos em
        _pasta_cotacao_atual (…\\COTACAO\\{código}, um nível acima), para não
        misturar com as propostas."""
        pasta_geof = self.entry_pasta_geof.get().strip()
        ano = self.entry_ano.get().strip()
        codigo_cotacao = (self.entradas_cotacao[self.indice_cotacao_ativa].get().strip()
                           if self.indice_cotacao_ativa is not None else "")
        if not pasta_geof or not ano or not codigo_cotacao:
            messagebox.showwarning("Atenção", "Confira PASTA DE TRABALHO, ANO e a COTAÇÃO ativa.")
            return None

        pasta_cotacao = os.path.join(
            pasta_geof, self._sufixo_planilha_atual, f"PDPAS {ano}", "COTACAO", codigo_cotacao)
        pasta_propostas = os.path.join(pasta_cotacao, "PROPOSTAS")
        if not os.path.isdir(pasta_propostas):
            messagebox.showerror("Pasta não encontrada",
                                  f"Pasta de PROPOSTAS não encontrada em:\n{pasta_propostas}")
            return None

        self._pasta_propostas_atual = pasta_propostas
        self._pasta_cotacao_atual = pasta_cotacao
        self._codigo_cotacao_atual = codigo_cotacao
        self._ano_atual_wizard = ano

        try:
            self.saldos_bd_atual = carregar_saldos_bd(self._obter_caminho_arquivo())
        except Exception as e:
            messagebox.showerror("Erro ao ler aba BD", str(e))
            return None

        return pasta_propostas, codigo_cotacao, ano

    @staticmethod
    def _caminho_arquivo_mapa_empresas(pasta_cotacao, codigo_cotacao):
        return os.path.join(pasta_cotacao, f"MAPA_EMPRESAS_{codigo_cotacao}.json")

    def _abrir_janela_wizard(self):
        if self.janela_wizard is not None and self.janela_wizard.winfo_exists():
            self.janela_wizard.destroy()

        janela = ctk.CTkToplevel(self)
        janela.title("Cadastro de Propostas por Fornecedor")
        janela.configure(fg_color=COR_FUNDO_APP)
        janela.geometry(f"{min(1560, janela.winfo_screenwidth() - 40)}x"
                        f"{min(920, janela.winfo_screenheight() - 80)}")
        centralizar_janela(janela)
        janela.grab_set()
        self.janela_wizard = janela

        self.wizard_cabecalho = ctk.CTkLabel(janela, text="", font=("Segoe UI", 18, "bold"), text_color=COR_TEXTO)
        self.wizard_cabecalho.pack(anchor="w", padx=16, pady=(14, 4))

        self.wizard_conteudo = ctk.CTkFrame(janela, fg_color=COR_CARD, corner_radius=10)
        self.wizard_conteudo.pack(fill="both", expand=True, padx=16, pady=(0, 8))

        self.wizard_rodape = ctk.CTkFrame(janela, fg_color="transparent")
        self.wizard_rodape.pack(fill="x", padx=16, pady=(0, 14))

        self._renderizar_form_fornecedor()

    def _fornecedor_atual_nome(self):
        return self.lista_empresas_mapa[self.indice_fornecedor_atual]

    def _limpar_wizard(self):
        for w in self.wizard_conteudo.winfo_children():
            w.destroy()
        for w in self.wizard_rodape.winfo_children():
            w.destroy()
        for job_id in list(self._blink_jobs.values()):
            try:
                self.after_cancel(job_id)
            except Exception:
                pass
        self._blink_jobs = {}

    # ---------------- FORMULÁRIO (edição) ----------------
    def _renderizar_form_fornecedor(self):
        self._limpar_wizard()
        nome = self._fornecedor_atual_nome()
        total_fornecedores = len(self.lista_empresas_mapa)
        # Prioriza dados ainda não confirmados (Editar, vindo da prévia) sobre os já salvos
        if self._dados_form_atual is not None and self._dados_form_atual.get("_fornecedor") == nome:
            dados_existentes = self._dados_form_atual
        else:
            dados_existentes = self.mapa_fornecedores.get(nome)

        self.wizard_cabecalho.configure(
            text=f"Fornecedor {self.indice_fornecedor_atual + 1} de {total_fornecedores} — {nome}")

        topo = ctk.CTkFrame(self.wizard_conteudo, fg_color="transparent")
        topo.pack(fill="x", padx=14, pady=(12, 6))
        ctk.CTkLabel(topo, text="DATA DA PROPOSTA:", font=("Segoe UI", 15, "bold"),
                     text_color=COR_TEXTO).pack(side="left", padx=(0, 8))
        self.wizard_entry_data = ctk.CTkEntry(topo, font=("Segoe UI", 17), height=34, width=150,
                                              placeholder_text="DD/MM/AAAA")
        self.wizard_entry_data.pack(side="left")
        self.wizard_entry_data.bind("<KeyRelease>", self._on_data_proposta_digitada)
        if dados_existentes:
            self.wizard_entry_data.insert(0, dados_existentes.get("data_proposta", ""))

        self._historico_repeticao = []
        ctk.CTkButton(topo, text="Repetir dados do primeiro item", font=("Segoe UI", 15, "bold"),
                      command=self._repetir_dados_primeiro_item).pack(side="left", padx=(40, 10))
        self.wizard_botao_desfazer = ctk.CTkButton(topo, text="Desfazer repetição", font=("Segoe UI", 15, "bold"),
                      state="disabled", command=self._desfazer_repeticao)
        self.wizard_botao_desfazer.pack(side="left", padx=10)
        # Cabeçalhos e campos compartilham a mesma grade e as mesmas colunas.
        self.wizard_scroll = ctk.CTkScrollableFrame(self.wizard_conteudo, fg_color="#161616", height=380)
        self.wizard_scroll.pack(fill="both", expand=True, padx=14, pady=(6, 8))
        cabecalho_cols = self.wizard_scroll
        titulos = ["ITEM", "CÓDIGO", "DESCRIÇÃO", "UN", "QTD", "QUANTIDADE", "UNITÁRIO",
                   "TOTAL", "MARCA", "PRAZO"]
        larguras = [50, 90, 340, 50, 50, 115, 145, 145, 165, 75]
        for i, (t, w) in enumerate(zip(titulos, larguras)):
            self.wizard_scroll.grid_columnconfigure(i, minsize=w+4, weight=1 if i == 2 else 0)
            ctk.CTkLabel(cabecalho_cols, text=t, font=("Segoe UI", 15, "bold"),
                         text_color=COR_TEXTO, width=w).grid(row=0, column=i, padx=2, pady=4)

        itens_salvos = {it["codigo"]: it for it in dados_existentes["itens"]} if dados_existentes else {}
        self._entradas_item_atual = []
        self._item_proposta_ativo = None

        for i, item in enumerate(self.linhas_resultado):
            numero_item = item.get("ITEM")
            codigo = str(item.get("CODIGO") or "")
            descricao = str(item.get("DESCRICAO") or "")
            un = str(item.get("UN") or "")
            qtd_cotada = item.get("QTD")
            salvo = itens_salvos.get(codigo, {})

            linha = self.wizard_scroll
            linha_idx = 2*i+1
            fundo = "#161616"
            texto = "#ffffff"
            faixa = ctk.CTkFrame(linha, fg_color=fundo, corner_radius=0,
                                  border_width=2, border_color=fundo, height=1, width=1)
            faixa.grid(row=linha_idx, column=0, rowspan=2, columnspan=10, sticky="nsew")
            inicio_widgets = len(linha.winfo_children())

            ctk.CTkLabel(linha, text=formatar_quantidade_proposta(numero_item), font=("Segoe UI", 16), width=45
                         ).grid(row=linha_idx, pady=(4, 0), column=0, padx=4)
            ctk.CTkLabel(linha, text=codigo, font=("Segoe UI", 16), width=80).grid(row=linha_idx, pady=(4, 0), column=1, padx=2)
            ctk.CTkLabel(linha, text=descricao[:55], font=("Segoe UI", 16), width=340, wraplength=340,
                         anchor="w", justify="left").grid(row=linha_idx, pady=(4, 0), column=2, padx=2, sticky="w")
            ctk.CTkLabel(linha, text=un, font=("Segoe UI", 16), width=40).grid(row=linha_idx, pady=(4, 0), column=3, padx=2)
            ctk.CTkLabel(linha, text=formatar_quantidade_proposta(qtd_cotada), font=("Segoe UI", 16), width=40
                         ).grid(row=linha_idx, pady=(4, 0), column=4, padx=2)

            entry_qtd = ctk.CTkEntry(linha, width=115, height=34, font=("Segoe UI", 16))
            entry_qtd.grid(row=linha_idx, pady=(4, 0), column=5, padx=2)
            entry_qtd.insert(0, formatar_quantidade_proposta(salvo.get("quantidade", qtd_cotada)))

            entry_unit = ctk.CTkEntry(linha, width=145, height=34, font=("Segoe UI", 16),
                                      placeholder_text="0,0000")
            entry_unit.grid(row=linha_idx, pady=(4, 0), column=6, padx=2)
            if salvo.get("unitario") is not None:
                entry_unit.insert(0, formatar_unitario_proposta(salvo["unitario"]))

            label_total = ctk.CTkLabel(linha, text=formatar_moeda(salvo.get("total", 0)),
                                       font=("Segoe UI", 15, "bold"), width=145)
            label_total.grid(row=linha_idx, pady=(4, 0), column=7, padx=2)

            entry_marca = ctk.CTkEntry(linha, width=165, height=34, font=("Segoe UI", 16))
            entry_marca.grid(row=linha_idx, pady=(4, 0), column=8, padx=2)
            entry_marca.insert(0, salvo.get("marca", ""))

            entry_prazo = ctk.CTkEntry(linha, width=75, height=34, font=("Segoe UI", 16))
            entry_prazo.grid(row=linha_idx, pady=(4, 0), column=9, padx=2)
            entry_prazo.insert(0, formatar_quantidade_proposta(salvo.get("prazo", "")))

            label_alerta = ctk.CTkLabel(self.wizard_scroll, text="", height=18, font=("Segoe UI", 13, "bold"),
                                        text_color="#ff5c5c")
            label_alerta.grid(row=linha_idx+1, column=0, columnspan=10, sticky="w", padx=4, pady=(0, 2))

            registro = {
                "numero_item": numero_item, "codigo": codigo, "descricao": descricao, "un": un,
                "qtd_cotada": qtd_cotada,
                "entry_qtd": entry_qtd, "entry_unit": entry_unit, "label_total": label_total,
                "entry_marca": entry_marca, "entry_prazo": entry_prazo, "label_alerta": label_alerta,
                "excedido": False,
                "faixa": faixa, "fundo": fundo, "texto": texto,
                "labels": [w for w in linha.winfo_children()[inicio_widgets:]
                           if w not in (entry_qtd, entry_unit, entry_marca, entry_prazo)],
            }
            self._entradas_item_atual.append(registro)
            self._pintar_linha_proposta(i, False)
            for campo in ("entry_qtd", "entry_unit", "entry_marca", "entry_prazo"):
                registro[campo].bind("<FocusIn>", lambda e, idx=i: self._ativar_linha_proposta(idx))
                registro[campo].bind("<FocusOut>",
                    lambda e, idx=i, chave=campo: self._sair_linha_proposta(idx, chave))

            entry_qtd.bind("<KeyRelease>", lambda e, idx=i: self._atualizar_total_item(idx))
            entry_unit.bind("<KeyRelease>", lambda e, idx=i: self._atualizar_total_item(idx))
            entry_prazo.bind("<KeyRelease>", lambda e, idx=i: self._mascarar_prazo(idx))

            self._atualizar_total_item(i)

        # ---------------- RODAPÉ / NAVEGAÇÃO ----------------
        ctk.CTkButton(self.wizard_rodape, text="◀ Fornecedor anterior", font=("Segoe UI", 15, "bold"),
                      height=34, state=("normal" if self.indice_fornecedor_atual > 0 else "disabled"),
                      command=lambda: self._ir_fornecedor(self.indice_fornecedor_atual - 1)
                      ).pack(side="left", padx=4)
        ctk.CTkButton(self.wizard_rodape, text="Salvar / Pré-visualizar", font=("Segoe UI", 15, "bold"),
                      height=34, fg_color=COR_DESTAQUE,
                      command=self._salvar_formulario_atual).pack(side="left", padx=12)
        pode_avancar = self.indice_fornecedor_atual < total_fornecedores - 1
        ctk.CTkButton(self.wizard_rodape, text="Fornecedor seguinte ▶", font=("Segoe UI", 15, "bold"),
                      height=34, state=("normal" if pode_avancar else "disabled"),
                      command=lambda: self._ir_fornecedor(self.indice_fornecedor_atual + 1)
                      ).pack(side="left", padx=4)
        ctk.CTkButton(self.wizard_rodape, text="Fechar", font=("Segoe UI", 15, "bold"), height=34,
                      fg_color="#7a2626", command=self._fechar_wizard_e_gerar_relacao
                      ).pack(side="right", padx=4)

    def _pintar_linha_proposta(self, idx, ativo):
        reg = self._entradas_item_atual[idx]
        fundo = "#404040" if ativo else reg["fundo"]
        texto = reg["texto"]
        reg["faixa"].configure(fg_color=fundo, border_color="#8c8c8c" if ativo else fundo)
        for campo in ("entry_qtd", "entry_unit", "entry_marca", "entry_prazo"):
            reg[campo].configure(bg_color=fundo, text_color=texto,
                                 font=("Segoe UI", 16, "bold" if ativo else "normal"))
        for label in reg["labels"]:
            label.configure(fg_color=fundo)
            if label is not reg["label_alerta"]:
                label.configure(font=("Segoe UI", 16,
                    "bold" if ativo or label is reg["label_total"] else "normal"))
            if label is not reg["label_alerta"] and not (label is reg["label_total"] and reg["excedido"]):
                label.configure(text_color=texto)

    def _ativar_linha_proposta(self, idx):
        anterior = self._item_proposta_ativo
        if anterior is not None and anterior != idx:
            self._pintar_linha_proposta(anterior, False)
        self._item_proposta_ativo = idx
        self._pintar_linha_proposta(idx, True)

    def _sair_linha_proposta(self, idx, campo):
        if campo in ("entry_qtd", "entry_unit"):
            self._formatar_campo_proposta(idx, campo)
        self._pintar_linha_proposta(idx, False)
        self._item_proposta_ativo = None

    def _formatar_campo_proposta(self, idx, campo):
        entry = self._entradas_item_atual[idx][campo]
        funcao = formatar_unitario_proposta if campo == "entry_unit" else formatar_quantidade_proposta
        valor = funcao(entry.get())
        entry.delete(0, "end")
        entry.insert(0, valor)
        self._atualizar_total_item(idx)

    def _repetir_dados_primeiro_item(self):
        registros = self._entradas_item_atual
        if len(registros) < 2:
            return
        campos = [campo for campo in ("entry_qtd", "entry_unit", "entry_marca", "entry_prazo")
                  if registros[0][campo].get().strip()]
        if not campos:
            return
        self._historico_repeticao.append([{campo: reg[campo].get() for campo in campos} for reg in registros[1:]])
        for idx, reg in enumerate(registros[1:], 1):
            for campo in campos:
                valor = registros[0][campo].get()
                if campo == "entry_unit":
                    valor = formatar_unitario_proposta(valor)
                elif campo == "entry_qtd":
                    valor = formatar_quantidade_proposta(valor)
                reg[campo].delete(0, "end")
                reg[campo].insert(0, valor)
            self._atualizar_total_item(idx)
        self.wizard_botao_desfazer.configure(state="normal")

    def _desfazer_repeticao(self):
        if not self._historico_repeticao:
            return
        anterior = self._historico_repeticao.pop()
        for idx, (reg, valores) in enumerate(zip(self._entradas_item_atual[1:], anterior), 1):
            for campo, valor in valores.items():
                reg[campo].delete(0, "end")
                reg[campo].insert(0, valor)
            self._atualizar_total_item(idx)
        self.wizard_botao_desfazer.configure(state="normal" if self._historico_repeticao else "disabled")

    def _on_data_proposta_digitada(self, event=None):
        texto = formatar_data(self.wizard_entry_data.get())
        if texto != self.wizard_entry_data.get():
            self.wizard_entry_data.delete(0, "end")
            self.wizard_entry_data.insert(0, texto)
            self.wizard_entry_data.icursor("end")

    def _mascarar_prazo(self, idx):
        entry = self._entradas_item_atual[idx]["entry_prazo"]
        texto = formatar_apenas_digitos(entry.get(), 2)
        if texto != entry.get():
            entry.delete(0, "end")
            entry.insert(0, texto)
            entry.icursor("end")

    def _atualizar_total_item(self, idx):
        reg = self._entradas_item_atual[idx]
        quantidade = numero_proposta(reg["entry_qtd"].get())
        unitario = numero_proposta(reg["entry_unit"].get())
        total = (quantidade or 0) * (unitario or 0)
        reg["label_total"].configure(text=formatar_moeda(total))

        saldo = self.saldos_bd_atual.get(normalizar_texto(reg["codigo"]))
        excedido = saldo is not None and total > saldo and total > 0
        reg["excedido"] = excedido
        if excedido:
            reg["label_alerta"].configure(
                text=f"VALOR DISPONÍVEL EXCEDIDO (saldo BD: {formatar_moeda(saldo)})")
            self._iniciar_blink(idx)
        else:
            reg["label_alerta"].configure(text="")
            self._parar_blink(idx)

    def _iniciar_blink(self, idx):
        if idx in self._blink_jobs:
            return  # já está piscando

        def alternar(aceso=True):
            reg = self._entradas_item_atual[idx]
            if not reg["excedido"]:
                reg["label_total"].configure(text_color=reg["texto"])
                self._blink_jobs.pop(idx, None)
                return
            reg["label_total"].configure(text_color=("#ff2b2b" if aceso else "#7a0000"))
            job_id = self.after(400, lambda: alternar(not aceso))
            self._blink_jobs[idx] = job_id

        alternar(True)

    def _parar_blink(self, idx):
        job_id = self._blink_jobs.pop(idx, None)
        if job_id:
            try:
                self.after_cancel(job_id)
            except Exception:
                pass
        if idx < len(self._entradas_item_atual):
            reg = self._entradas_item_atual[idx]
            reg["label_total"].configure(text_color=reg["texto"])

    def _salvar_formulario_atual(self):
        data_proposta = self.wizard_entry_data.get().strip()
        if not data_proposta:
            messagebox.showwarning("Atenção", "Informe a DATA DA PROPOSTA.")
            return

        itens_calculados = []
        algum_excedido = False
        for idx, reg in enumerate(self._entradas_item_atual):
            for campo in ("entry_qtd", "entry_unit"):
                if reg[campo].get().strip() and numero_proposta(reg[campo].get()) is None:
                    messagebox.showwarning("Valor inválido", f"Confira a quantidade e o unitário do item {reg['numero_item']}.")
                    return
                self._formatar_campo_proposta(idx, campo)
            quantidade = numero_proposta(reg["entry_qtd"].get())
            unitario = numero_proposta(reg["entry_unit"].get())
            total = round((quantidade or 0) * (unitario or 0), 2)
            prazo_texto = reg["entry_prazo"].get().strip()
            itens_calculados.append({
                "numero_item": reg.get("numero_item"),
                "codigo": reg["codigo"], "descricao": reg["descricao"], "un": reg["un"],
                "qtd_cotada": reg["qtd_cotada"], "quantidade": quantidade, "unitario": unitario,
                "total": total, "marca": reg["entry_marca"].get().strip()[:20],
                "prazo": int(prazo_texto) if prazo_texto.isdigit() else None,
            })
            if reg["excedido"]:
                algum_excedido = True

        if algum_excedido:
            if not messagebox.askyesno(
                    "Valor disponível excedido",
                    "Um ou mais itens ultrapassam o saldo disponível na aba BD.\n"
                    "Deseja continuar mesmo assim?"):
                return

        self._dados_form_atual = {"_fornecedor": self._fornecedor_atual_nome(),
                                   "data_proposta": data_proposta, "itens": itens_calculados}
        self._renderizar_preview_fornecedor()

    # ---------------- PRÉVIA (somente leitura) ----------------
    def _renderizar_preview_fornecedor(self):
        for w in self.wizard_conteudo.winfo_children():
            w.destroy()
        for w in self.wizard_rodape.winfo_children():
            w.destroy()

        nome = self._fornecedor_atual_nome()
        dados = self._dados_form_atual

        ctk.CTkLabel(self.wizard_conteudo, text=f"Prévia — {nome} (Proposta de {dados['data_proposta']})",
                     font=FONTE_SECAO, text_color=COR_TEXTO).pack(anchor="w", padx=14, pady=(12, 6))

        # Somente itens com valor lançado (quantidade e unitário > 0) aparecem na prévia;
        # os demais continuam salvos internamente, só ficam ocultos aqui.
        itens_com_valor = [it for it in dados["itens"]
                            if (it.get("unitario") or 0) > 0 and (it.get("quantidade") or 0) > 0]
        itens_ocultos = len(dados["itens"]) - len(itens_com_valor)

        scroll = ctk.CTkScrollableFrame(self.wizard_conteudo, fg_color="#161616", height=410)
        scroll.pack(fill="both", expand=True, padx=14, pady=(0, 4))

        cabecalho = ctk.CTkFrame(scroll, fg_color="#1b1b1b")
        cabecalho.pack(fill="x", pady=(0, 4))
        titulos = ["ITEM", "CÓDIGO", "DESCRIÇÃO", "UN", "QTD", "QUANTIDADE", "UNITÁRIO", "TOTAL",
                   "MARCA", "PRAZO"]
        larguras = [45, 80, 280, 40, 40, 90, 90, 100, 130, 55]
        for i, (t, w) in enumerate(zip(titulos, larguras)):
            ctk.CTkLabel(cabecalho, text=t, font=("Segoe UI", 11, "bold"),
                         text_color=COR_TEXTO, width=w).grid(row=0, column=i, padx=2, pady=4)

        total_geral = 0
        for item in itens_com_valor:
            linha = ctk.CTkFrame(scroll, fg_color="transparent")
            linha.pack(fill="x", pady=1)
            valores = [formatar_quantidade_proposta(item.get("numero_item")), item["codigo"], item["descricao"][:55],
                       item["un"], formatar_quantidade_proposta(item["qtd_cotada"]), formatar_quantidade_proposta(item["quantidade"]),
                       formatar_unitario_proposta(item["unitario"] or 0), formatar_moeda(item["total"]),
                       item["marca"], str(item["prazo"] or "")]
            for i, (v, w) in enumerate(zip(valores, larguras)):
                ctk.CTkLabel(linha, text=v, font=FONTE_TABELA, width=w,
                             anchor="w" if i == 2 else "center").grid(row=0, column=i, padx=2)
            total_geral += item["total"] or 0

        if itens_ocultos:
            ctk.CTkLabel(self.wizard_conteudo,
                         text=f"({itens_ocultos} item(ns) sem valor lançado não exibido(s) aqui — "
                              "clique em Editar para preenchê-los)",
                         font=("Segoe UI", 11), text_color=COR_TEXTO_SEC
                         ).pack(anchor="w", padx=16, pady=(0, 2))

        ctk.CTkLabel(self.wizard_conteudo, text=f"TOTAL GERAL DA PROPOSTA: {formatar_moeda(total_geral)}",
                     font=("Segoe UI", 15, "bold"), text_color="#8ee08e"
                     ).pack(anchor="e", padx=16, pady=(4, 10))

        ctk.CTkButton(self.wizard_rodape, text="Editar", font=FONTE_BOTAO, height=34,
                      command=self._renderizar_form_fornecedor).pack(side="left", padx=4)
        ctk.CTkButton(self.wizard_rodape, text="Confirmar", font=FONTE_BOTAO, height=34,
                      fg_color=COR_DESTAQUE, command=self._confirmar_fornecedor_atual
                      ).pack(side="left", padx=12)
        ctk.CTkButton(self.wizard_rodape, text="Fechar", font=FONTE_BOTAO, height=34,
                      fg_color="#7a2626", command=self._fechar_wizard_e_gerar_relacao
                      ).pack(side="right", padx=4)

    def _confirmar_fornecedor_atual(self):
        nome = self._fornecedor_atual_nome()
        dados_finais = dict(self._dados_form_atual)
        dados_finais.pop("_fornecedor", None)
        self.mapa_fornecedores[nome] = dados_finais
        self._dados_form_atual = None
        self._salvar_mapa_em_arquivo()

        if self.indice_fornecedor_atual < len(self.lista_empresas_mapa) - 1:
            self.indice_fornecedor_atual += 1
            self._renderizar_form_fornecedor()
        else:
            self._finalizar_wizard()

    def _ir_fornecedor(self, indice):
        if indice < 0 or indice >= len(self.lista_empresas_mapa):
            return
        self._dados_form_atual = None
        self.indice_fornecedor_atual = indice
        self._renderizar_form_fornecedor()

    def _salvar_mapa_em_arquivo(self):
        caminho_json = self._caminho_arquivo_mapa_empresas(
            self._pasta_cotacao_atual, self._codigo_cotacao_atual)
        try:
            # Editar propostas não deve apagar as decisões salvas na adjudicação.
            conteudo = {}
            if os.path.exists(caminho_json):
                with open(caminho_json, "r", encoding="utf-8") as f:
                    conteudo = json.load(f)
            conteudo.update({
                "cotacao": self._codigo_cotacao_atual,
                "ano": self._ano_atual_wizard,
                "fornecedores": self.mapa_fornecedores,
            })
            temporario = caminho_json + ".tmp"
            with open(temporario, "w", encoding="utf-8") as f:
                json.dump(conteudo, f, ensure_ascii=False, indent=2)
            os.replace(temporario, caminho_json)
        except Exception as e:
            messagebox.showerror("Erro ao salvar", f"Não foi possível salvar o cadastro:\n{e}")

    def _finalizar_wizard(self):
        messagebox.showinfo(
            "Cadastro concluído",
            "Todos os fornecedores foram cadastrados e salvos.\n"
            f"Arquivo: {self._caminho_arquivo_mapa_empresas(self._pasta_cotacao_atual, self._codigo_cotacao_atual)}\n\n"
            "Gerando o MAPA COMPARATIVO PRÉVIO RESUMIDO...\n\n"
            "Para gerar o Mapa Comparativo (imagem/PDF), use o botão ADJUDICAÇÃO na tela principal.")
        self._fechar_wizard_e_gerar_relacao()

    def _fechar_wizard_e_gerar_relacao(self):
        """Usado pelo botão FECHAR do assistente: gera a Relação de Propostas
        por Menor Preço (tela + PDF) com o que já foi cadastrado até aqui, e
        então fecha a janela do assistente."""
        if self.mapa_fornecedores:
            self._gerar_relacao_menor_preco()
        if self.janela_wizard is not None and self.janela_wizard.winfo_exists():
            self.janela_wizard.destroy()

    # ==================================================================
    # PASSO 26 — Relação por menor preço (tela + PDF)
    # ==================================================================
    def _obter_fornecedores_texto_completo(self):
        """Mapa nome_normalizado -> 'NOME (codigo)' lido da aba Fornecedores,
        usado para formatar a coluna FORNECEDOR nos relatórios."""
        try:
            fornecedores = carregar_fornecedores(self._obter_caminho_arquivo())
        except Exception:
            return {}
        return {f["nome_normalizado"]: f["texto_completo"] for f in fornecedores}

    def _gerar_relacao_menor_preco(self):
        if not self.mapa_fornecedores:
            return
        fornecedores_texto = self._obter_fornecedores_texto_completo()
        caminho_json = self._caminho_arquivo_mapa_empresas(self._pasta_cotacao_atual, self._codigo_cotacao_atual)
        excluidos = {}
        if os.path.exists(caminho_json):
            try:
                with open(caminho_json, "r", encoding="utf-8") as arquivo:
                    decisoes = json.load(arquivo).get("decisoes_itens", {})
                excluidos = {d["codigo"]: d.get("fornecedores_retirados", [])
                             for d in decisoes.values() if "codigo" in d}
            except (OSError, ValueError) as e:
                messagebox.showerror("Erro ao ler classificação", str(e))
                return
        ranking = montar_ranking_itens(self.mapa_fornecedores, fornecedores_texto, excluidos)
        if not ranking:
            messagebox.showinfo("MAPA COMPARATIVO PRÉVIO RESUMIDO",
                                 "Nenhuma proposta com valor unitário informado foi encontrada.")
            return

        caminho_pdf = os.path.join(
            self._pasta_cotacao_atual,
            f"MAPA COMPARATIVO PRÉVIO RESUMIDO {self._codigo_cotacao_atual}.pdf")
        try:
            gerar_pdf_ranking(caminho_pdf, self._codigo_cotacao_atual, self._ano_atual_wizard, ranking)
        except Exception as e:
            messagebox.showerror("Erro ao gerar PDF", str(e))
            return

        self._mostrar_relacao_menor_preco(ranking, caminho_pdf)

    def _mostrar_relacao_menor_preco(self, ranking, caminho_pdf):
        janela = ctk.CTkToplevel(self)
        janela.title(f"MAPA COMPARATIVO PRÉVIO RESUMIDO - Cotação {self._codigo_cotacao_atual}")
        janela.configure(fg_color=COR_FUNDO_APP)
        janela.geometry("1100x600")
        centralizar_janela(janela)
        janela.grab_set()

        ctk.CTkLabel(janela, text="MAPA COMPARATIVO PRÉVIO RESUMIDO",
                     font=FONTE_SECAO, text_color=COR_TEXTO).pack(anchor="w", padx=16, pady=(14, 2))
        ctk.CTkLabel(janela, text=caminho_pdf, font=("Segoe UI", 11),
                     text_color=COR_TEXTO_SEC).pack(anchor="w", padx=16, pady=(0, 8))

        scroll = ctk.CTkScrollableFrame(janela, fg_color="#161616")
        scroll.pack(fill="both", expand=True, padx=16, pady=(0, 10))

        titulos = ["ITEM", "CÓDIGO", "FORNECEDOR", "QTD", "UNITÁRIO", "TOTAL", "PRAZO", "MARCA"]
        larguras = [45, 80, 260, 55, 90, 100, 60, 130]
        for i, (t, w) in enumerate(zip(titulos, larguras)):
            scroll.grid_columnconfigure(i, minsize=w+4, weight=1 if i == 2 else 0)
            ctk.CTkLabel(scroll, text=t, font=("Segoe UI", 11, "bold"), width=w
                         ).grid(row=0, column=i, padx=2, pady=4)
        linha_idx = 1
        for numero_item, (codigo, info) in enumerate(ranking, start=1):
            numero = formatar_quantidade_proposta(info.get("numero_item") if info.get("numero_item") is not None else numero_item)
            ctk.CTkLabel(scroll, text=numero, font=("Segoe UI", 13, "bold"), text_color="#FF5555").grid(row=linha_idx, column=0, padx=2, pady=(8, 2))
            ctk.CTkLabel(scroll, text=codigo, font=("Segoe UI", 13, "bold"), text_color="#FF5555").grid(row=linha_idx, column=1, padx=2, pady=(8, 2))
            descricao = " ".join(str(info.get("descricao", "")).split())
            ctk.CTkLabel(scroll, text=descricao[:130] + ("..." if len(descricao) > 130 else ""),
                font=("Segoe UI", 13, "bold"), text_color="#FF5555", anchor="w", wraplength=800
                ).grid(row=linha_idx, column=2, columnspan=6, padx=2, pady=(8, 2), sticky="w")
            linha_idx += 1
            for pos, p in enumerate(info["propostas"], start=1):
                destaque = pos <= 3
                cor = "#8ee08e" if pos == 1 else (COR_TEXTO if not destaque else "#e8e08e")
                fonte = ("Segoe UI", 13 if pos == 1 else 11, "bold") if destaque else FONTE_TABELA
                ctk.CTkLabel(scroll, text=f"{pos}º COLOCADO", font=fonte, text_color=cor
                             ).grid(row=linha_idx, column=0, columnspan=2, padx=2)
                valores = [p["fornecedor"], formatar_quantidade_proposta(p["quantidade"]),
                           formatar_moeda(p["unitario"]), formatar_moeda(p["total"]),
                           f'{p["prazo"]}d' if p["prazo"] else "", p["marca"] or ""]
                for i, (v, w) in enumerate(zip(valores, larguras[2:]), 2):
                    ctk.CTkLabel(scroll, text=v, font=fonte, text_color=cor, width=w,
                                 anchor="w" if i == 2 else "center", wraplength=w).grid(row=linha_idx, column=i, padx=2)
                linha_idx += 1

        ctk.CTkButton(janela, text="Fechar", font=FONTE_BOTAO, height=34,
                      command=janela.destroy).pack(pady=(0, 14))